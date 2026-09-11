# -*- coding: utf-8 -*-
"""O arquivo solto no dropzone do DCE Swap → grade de textos.

Aceita .xlsx/.xlsm (openpyxl), .xls (xlrd) e texto delimitado (.csv/.txt,
separador detectado na primeira linha). Toda aba entra, uma atrás da outra:
quem separa as duas tabelas é o cabeçalho de cada uma (`domain`), então tanto
faz se vieram em abas distintas ou coladas na mesma.

Célula → texto, de forma DETERMINÍSTICA (é o que vai para o JSON-dia e para
o arquivo da Intrag): data vira ISO `AAAA-MM-DD`; número inteiro sai sem
`.0`; booleano sai `TRUE`/`FALSE` como a planilha mostra; vazio é ''.
"""
import csv
import datetime as _dt
import io


def _cell_text(v):
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, _dt.datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0:
            return v.strftime('%Y-%m-%d')
        return v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, _dt.date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float):
        if v != v:                      # NaN
            return ''
        if v.is_integer() and abs(v) < 1e15:
            return str(int(v))
        return ('%.15g' % v)
    return str(v)


def grid_from_upload(filename, data):
    """(grid, sheets) — a grade concatenada de todas as abas e os nomes delas.
    Levanta ValueError com a razão quando o arquivo não é legível."""
    name = str(filename or '').lower()
    if name.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        grid, sheets = [], []
        for ws in wb.worksheets:
            sheets.append(ws.title)
            for row in ws.iter_rows(values_only=True):
                grid.append([_cell_text(c) for c in row])
            grid.append([])             # linha em branco entre abas
        return grid, sheets
    if name.endswith('.xls'):
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        grid, sheets = [], []
        for ws in wb.sheets():
            sheets.append(ws.name)
            for r in range(ws.nrows):
                cells = []
                for c in range(ws.ncols):
                    cell = ws.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        cells.append(_cell_text(xlrd.xldate_as_datetime(cell.value, wb.datemode)))
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        cells.append('TRUE' if cell.value else 'FALSE')
                    else:
                        cells.append(_cell_text(cell.value))
                grid.append(cells)
            grid.append([])
        return grid, sheets
    # Texto delimitado: ';' do extrato, '\t' de um copy/paste do Excel, ','.
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = data.decode('latin-1')
    first = next((ln for ln in text.splitlines() if ln.strip()), '')
    delim = max((';', '\t', ','), key=first.count)
    if first.count(delim) == 0:
        raise ValueError('no delimiter found on the first line')
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    return [[str(c) for c in row] for row in reader], ['text']
