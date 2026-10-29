# -*- coding: utf-8 -*-
"""RED Clients: a planilha (a aba pelo nome) e o status da última rodada."""
import io
import os

from apps.pages.features.red_clients import domain


def _R():
    """Busca ATRASADA no routes (plataforma; os testes trocam atributos lá)."""
    from apps.pages import routes
    return routes


class PlanilhaInvalida(ValueError):
    """A planilha não serve — `code` vai para a tela traduzido."""
    def __init__(self, code, **params):
        super().__init__(code)
        self.code = code
        self.params = params


def linhas_da_aba(raw):
    """As linhas da aba "JPM Entities-Open Settlements" (listas de valores),
    pelo CONTEÚDO do arquivo (xlsx ou xls de verdade, não pela extensão)."""
    fmt = _R()._latam_sniff_format(raw or b'')
    if fmt == 'xlsx':
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            nomes = list(wb.sheetnames)
            alvo = next((n for n in nomes if domain.same_sheet(n)), None)
            if not alvo:
                raise PlanilhaInvalida('red_sheet_missing', sheets=', '.join(nomes))
            ws = wb[alvo]
            try:
                ws.reset_dimensions()       # a dimensão gravada pode mentir
            except Exception:               # noqa: BLE001
                pass
            return [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    if fmt == 'xls':
        try:
            import xlrd
        except ImportError:
            raise PlanilhaInvalida('red_xls_reader')
        wb = xlrd.open_workbook(file_contents=raw)
        nomes = wb.sheet_names()
        alvo = next((n for n in nomes if domain.same_sheet(n)), None)
        if not alvo:
            raise PlanilhaInvalida('red_sheet_missing', sheets=', '.join(nomes))
        sh = wb.sheet_by_name(alvo)
        return [sh.row_values(i) for i in range(sh.nrows)]
    raise PlanilhaInvalida('red_not_excel')


def status_file():
    return os.path.join(_R()._DAILY_METRIC_DIR, 'red_clients_status.json')


def write_status(doc):
    try:
        _R()._atomic_write_json(status_file(), doc)
    except Exception:                       # noqa: BLE001
        _R().log.warning('[red-clients] status não gravado', exc_info=True)


def read_status():
    from apps.pages import data_store as _store
    try:
        d = _store.read(status_file())
    except Exception:                       # noqa: BLE001 — sem rodada ainda / ilegível
        return None
    return d if isinstance(d, dict) else None
