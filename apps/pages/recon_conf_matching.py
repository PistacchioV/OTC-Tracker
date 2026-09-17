# -*- coding: utf-8 -*-
"""Conf. Matching — FepWeb × Athena (a tradução do workflow Alteryx homônimo).

O FepWeb é onde a confirmação do termo de moeda nasce; a Athena é onde a
operação foi bookada. O batimento responde, para UM trade date (o D-1 ANBIMA
por padrão), três perguntas:

  * a operação da Athena tem contrato no FepWeb?      → senão `Missing FepWeb`
  * o contrato do FepWeb tem operação na Athena?      → senão `Missing Athena`
  * o contrato aparece mais de uma vez no FepWeb?     → `Duplicated`

e, para o que casou, qual é a situação da ASSINATURA: `Ok` (nada a cobrar) ou
`Pending` (assinatura a coletar), com o Pending Status ao lado.

O que mudou em relação ao workflow, e por quê:

  * **O lado FepWeb é o ANEXO do e-mail** `(REPORT) FEPWeb - Operacoes D-4`
    (box compartilhado › Inbox › Automatico), e não um arquivo que alguém salva
    numa pasta. A leitura do box é a MESMA função da Recon de CGD. Sem Outlook
    (dev, servidor que não é Windows) cai para o arquivo na pasta do workflow,
    avisando qual das duas fontes valeu.
  * **O lado Athena é a API de NDF do New Deals**, e não o `Athena.xlsx`.
  * **Cliente se identifica por CHAVE**: CPF/CNPJ no FepWeb, SPN na Athena, os
    dois contra o Reference Data. Por isso a limpeza de nome por RegEx do
    workflow (Tools 37–40) não existe aqui.
  * **Nenhuma lista de cliente no código** (§6 do CLAUDE.md). As contrapartes
    internas do Tool 11/71 (`LABAYSTR`, `FXECOM`…) saem dos MESMOS cadastros que
    o import do New Deals usa — `interbook-ndf`, `le-accronym` e o `ECONOMIC
    GROUP = INTERNAL` do Reference Data —, e a lista de exceção do Tool 84
    (`ADM`, `Cargill`…) virou o `SIGNATURE TYPE` do Reference Data.
  * **O Pending Status é o da casa** (`_pc_signature_pending_status`): é a
    função que o New Deals e o Pending Confirmation já chamam. Uma terceira
    cópia da regra daria uma terceira resposta para a mesma operação.
  * **`Pending Update.xlsx` (Tool 28) não é gravada**: o OTC Tracker já alimenta
    a planilha sozinho.

Os dois e-mails do workflow (Confirmações Manuais e MT300) não são desta tela.
"""

import logging
import os
import threading
from datetime import date, datetime, timedelta

from apps.config import Config
from apps.pages import data_store as _store
from apps.pages import recon_cgd as _cgd
from apps.pages.data_paths import data_write

_LOG = logging.getLogger(__name__)

_CACHE_DIR = data_write('cache', 'reconciliation', 'conf-matching')
_COMMENTS_PATH = data_write('recon-conf-matching-comments.json')
_COMMENTS_LOCK = threading.Lock()

# ── Entradas ────────────────────────────────────────────────────────────────
# O assunto casa por PEDAÇO e normalizado (ver `recon_cgd.baixar_fep_do_box`).
FEP_MAIL_SUBJECT = os.getenv('CONFMATCH_FEP_MAIL_SUBJECT',
                             '(REPORT) FEPWeb - Operacoes D-4')
# O plano B é a pasta de input do próprio workflow — é onde a mesa já salvava o
# anexo à mão. Pende do `Config.SHARED_DRIVE_ROOT`, nunca de um `I:\` escrito.
INPUT_ROOT = os.getenv(
    'CONFMATCH_INPUT_ROOT',
    os.path.join(Config.SHARED_DRIVE_ROOT, 'Confirmation', 'Derivativos', 'Movimento',
                 'Pending Confirmation', 'Recon Fep Web x Athena'))
FEP_XLSX = os.getenv('CONFMATCH_FEP_XLSX', 'FEPWeb - Operacoes D-4.xlsx')
# O relatório cobre uma JANELA de dias: o e-mail que serve a uma data é o que
# chegou DEPOIS dela e dentro da janela. Em dias corridos, com folga.
FEP_WINDOW_DAYS = 5

STATUS_FORA = 'CANCELADO'

# ── Status ──────────────────────────────────────────────────────────────────
ST_MISSING_FEP = 'Missing FepWeb'
ST_MISSING_ATH = 'Missing Athena'
ST_DUPLICATED = 'Duplicated'
ST_MANUAL = 'Manual Confirmation'
ST_PENDING = 'Pending'
ST_OK = 'Ok'
# A ordem é a da gravidade, e é a da tabela e a dos cards.
STATUS_ORDER = (ST_MISSING_FEP, ST_MISSING_ATH, ST_DUPLICATED, ST_MANUAL,
                ST_PENDING, ST_OK)
COUNT_KEYS = {ST_MISSING_FEP: 'missing_fepweb', ST_MISSING_ATH: 'missing_athena',
              ST_DUPLICATED: 'duplicated', ST_MANUAL: 'manual',
              ST_PENDING: 'pending', ST_OK: 'ok'}
# O Forward Start não tem confirmação no FepWeb: ela é confeccionada à mão e
# corre pela esteira (Confirmations Monitor) — é o ramo "Confirmações Manuais"
# do workflow. Prazo e assinatura não têm o que dizer sobre ele.
PS_MANUAL = 'Manual confirmation workflow'

# As colunas da tela, na ordem, e as chaves de cada linha: o Advanced Export
# por intervalo monta o arquivo de `columns` + `rows[coluna]` (o contrato da
# Recon FXO). `key` fica FORA: é a chave do comentário, não um output.
COMMENT_COLUMN = 'Comments'
COLUMNS = ('Status', COMMENT_COLUMN, 'Pending Status', 'Trade Date', 'FepWeb ID',
           'Athena ID', 'FepWeb Client', 'Athena Client', 'CNPJ', 'SPN',
           'Signature Type', 'Instrument Type', 'Settlement Date', 'Tenor',
           'Quantity Ccy', 'Quantity', 'Other Ccy', 'Other Quantity', 'Strike',
           'Publisher', 'Cetip ID', 'End Counterparty', 'FepWeb Type', 'FepWeb Count')


def _routes():
    """Busca ATRASADA: os testes trocam atributos do `routes`, e o import no
    topo fecharia o ciclo (é o `routes` que importa a vertical desta recon)."""
    from apps.pages import routes
    return routes


# ── Normalização ─────────────────────────────────────────────────────────────

def contract_key(v):
    """O contrato como CHAVE do batimento: o FepWeb e a Athena escrevem o mesmo
    Deal Name com caixa e separador diferentes (`_` × `-`, que é a troca que o
    import do New Deals também faz)."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v if v is not None else '').strip().upper().replace('_', '-')


def _tax_key(v):
    """CPF/CNPJ só em dígitos e SEM zero à esquerda: a planilha entrega o
    documento como NÚMERO (o zero some) e o Reference Data guarda mascarado."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return _cgd._digits(v).lstrip('0')


def _num(v):
    try:
        n = _routes()._fxo_num(v)
    except Exception:
        n = None
    return n


# ── Reference Data ───────────────────────────────────────────────────────────

def _refdata_by_tax():
    """{documento → registro} do Reference Data. Duas linhas com o mesmo
    documento: vence quem RESPONDE mais (o mesmo desempate do índice por SPN)."""
    R = _routes()
    out = {}
    for rec in (R._refdata_records() or []):
        k = _tax_key(rec.get('TAX ID', ''))
        if not k:
            continue
        atual = out.get(k)
        if atual is None or R._refdata_spn_peso(rec) > R._refdata_spn_peso(atual):
            out[k] = rec
    return out


# ── Lado FepWeb ──────────────────────────────────────────────────────────────

_FEP_COLS = {
    'CONTRATO': 'contrato',
    'DATA OPERACAO': 'data',
    'STATUS OPERACAO': 'status',
    'NOME CLIENTE': 'cliente',
    'TIPO OPERACAO': 'tipo',
    'CPF/CNPJ CLIENTE': 'cnpj',
}
_FEP_OBRIGATORIAS = ('contrato', 'data')


def _fep_col_idx(cabecalho):
    """`{campo: índice}` casando por NOME normalizado — por posição a leitura
    erraria calada no dia em que alguém mexesse na exportação."""
    out = {}
    for i, c in enumerate(cabecalho):
        k = _FEP_COLS.get(_cgd._norm(c))
        if k and k not in out:
            out[k] = i
    return out


def _fep_header(linhas):
    """(índice da linha de cabeçalho, mapa de colunas). O relatório pode vir
    com um título acima da tabela; o cabeçalho é a primeira linha que tem o
    `Contrato` e a `Data Operação`."""
    for n, l in enumerate(linhas[:15]):
        idx = _fep_col_idx(l)
        if all(k in idx for k in _FEP_OBRIGATORIAS):
            return n, idx
    return None, {}


def _aceita_email(ref):
    """O e-mail que cobre `ref`: chegou depois dela e dentro da janela."""
    def aceita(recebido):
        try:
            d = recebido.date() if hasattr(recebido, 'date') else recebido
            d = date(d.year, d.month, d.day)
        except Exception:
            return False
        return ref < d <= ref + timedelta(days=FEP_WINDOW_DAYS)
    return aceita


def ler_fep(ref, avisos, path=None):
    """As operações do FepWeb do dia `ref`, sem as canceladas.

    Devolve `(linhas, rotulo, info)`. Levanta `RuntimeError` quando o arquivo
    não pôde ser lido: sem um dos lados, o batimento inteiro viraria `Missing
    FepWeb` — uma recon que parece cheia de quebra e não rodou.
    """
    origem = ''
    if not path:
        try:
            path, origem = _cgd.baixar_fep_do_box(
                avisos, assunto=FEP_MAIL_SUBJECT, prefixo='fepweb-ops-',
                aceita=_aceita_email(ref))
        except EnvironmentError as e:
            avisos.append(str(e))
        if not path:
            path = os.path.join(INPUT_ROOT, FEP_XLSX)
            avisos.append('Usei o relatório em pasta ({}) em vez do anexo do e-mail.'
                          .format(path))
    rotulo = origem or path
    if not os.path.isfile(path):
        raise RuntimeError('Relatório do FepWeb não encontrado: nem o e-mail "{}" '
                           'no box, nem o arquivo {}.'.format(FEP_MAIL_SUBJECT, path))
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    n, idx = _fep_header(linhas)
    if n is None:
        raise RuntimeError('O relatório do FepWeb ({}) não tem as colunas "Contrato" e '
                           '"Data Operação" — nada foi lido dele.'.format(rotulo))
    if 'cnpj' not in idx:
        avisos.append('O relatório do FepWeb veio sem a coluna "CPF/CNPJ CLIENTE": '
                      'o cliente do lado FepWeb sai pelo nome do próprio relatório.')

    info = {'lidas': 0, 'canceladas': 0, 'outras_datas': 0}
    out = []
    for l in linhas[n + 1:]:
        def cel(k):
            i = idx.get(k)
            return l[i] if (i is not None and i < len(l)) else None
        contrato = contract_key(cel('contrato'))
        if not contrato:
            continue
        info['lidas'] += 1
        if _cgd._norm(cel('status')) == STATUS_FORA:
            info['canceladas'] += 1
            continue
        dia = _cgd._parse_date(cel('data'))
        if dia != ref:
            info['outras_datas'] += 1
            continue
        out.append({'contrato': contrato, 'data': dia,
                    'cliente': str(cel('cliente') or '').strip(),
                    'cnpj': cel('cnpj'),
                    'tipo': str(cel('tipo') or '').strip()})
    if info['lidas'] and not out:
        avisos.append('O relatório do FepWeb ({}) não traz nenhuma operação de {} — '
                      'a data está fora da janela dele?'.format(rotulo, _cgd._fmt_date(ref)))
    return out, rotulo, info


# ── Lado Athena ──────────────────────────────────────────────────────────────

def _is_fwd_start(instr):
    # O mesmo teste do roteamento do New Deals (`_ndf_deal_from_api`).
    return 'FORWARDSTART' in str(instr or '').upper().replace(' ', '')


def _interna(R, norm, end_cp, desc, spn):
    """A perna é interna? As MESMAS perguntas do import do New Deals e do
    Pending Confirmation — cadastro, nunca lista no código. O `NDF` no End
    Counterparty é padrão de NOME DE BOOK (`BR ON - LN LAWTON NDF`), o filtro
    dos Tools 18/31 do workflow."""
    flat = end_cp.upper().replace(' ', '_').replace('-', '_')
    if 'GLOBAL_HOLDING_BOOK' in flat:
        return True
    if 'NDF' in end_cp.upper():
        return True
    if R._ndf_is_interbook(norm):
        return True
    if R._ndf_le_from_accronym(end_cp):
        return True
    return bool(R._pc_is_internal_counterparty(desc, spn))


def buscar_athena(ref):
    """Os registros crus do getTrades de NDF do dia. Falha LEVANTA com o motivo:
    a tela precisa dizer se foi o SSO, o timeout ou a URL (§476)."""
    from apps.pages import athena_api
    if not athena_api.is_available():
        raise RuntimeError('A pilha HTTP da API da Athena não está instalada (requests).')
    payload = athena_api.fetch_ndf_trades(ref.strftime('%Y%m%d'))
    return athena_api.extract_records(payload)


def ler_athena(ref, avisos, records=None):
    """As operações de NDF da Athena com trade date `ref`, sem canceladas e sem
    pernas internas. `records` explícito é o caminho dos testes."""
    R = _routes()
    if records is None:
        records = buscar_athena(ref)
    info = {'lidas': 0, 'canceladas': 0, 'internas': 0, 'outras_datas': 0}
    out, vistos = [], set()
    for rec in (records or []):
        if not isinstance(rec, dict):
            continue
        norm = R._ndf_api_norm(rec)
        get = norm.get
        deal = contract_key(get('DEAL NAME'))
        if not deal:
            continue
        info['lidas'] += 1
        if R._api_rec_is_cancelled(norm):
            info['canceladas'] += 1
            continue
        end_cp = str(get('END COUNTERPARTY') or '').strip()
        desc = str(R._ndf_api_get(norm, 'END COUNTERPARTY DESCRIPTION',
                                  'END COUNTERPARTY DESC') or '').strip()
        spn = str(get('SPN') or '').strip()
        if _interna(R, norm, end_cp, desc, spn):
            info['internas'] += 1
            continue
        trade = _cgd._parse_date(R._fxo_date_dmy(get('TRADE DATE')))
        if trade != ref:
            info['outras_datas'] += 1
            continue
        if deal in vistos:          # a API repete o trade a cada versão viva
            continue
        vistos.add(deal)
        out.append({
            'deal': deal, 'end_cp': end_cp, 'desc': desc, 'spn': spn,
            'cetip_id': str(R._ndf_api_get(norm, 'CETIP ID', 'CETIPID', 'EXTERNAL ID') or '').strip(),
            'trade': trade,
            'settle': _cgd._parse_date(R._fxo_date_dmy(get('SETTLEMENT DATE'))),
            'instrument': str(get('INSTRUMENT TYPE') or '').strip(),
            'qty': _num(get('QUANTITY')),
            'qty_ccy': R._fxo_ccy(get('QUANTITY CURRENCY')),
            'other_qty': _num(get('OTHER QUANTITY')),
            'other_ccy': R._fxo_ccy(get('OTHER QUANTITY UNITS')),
            'publisher': str(get('PUBLISHER') or '').strip(),
            'strike': _num(get('STRIKE')),
        })
    if info['outras_datas']:
        avisos.append('{} operação(ões) da API vieram com Trade Date diferente de {} e '
                      'ficaram de fora.'.format(info['outras_datas'], _cgd._fmt_date(ref)))
    return out, info


# ── O batimento ──────────────────────────────────────────────────────────────

def _assinatura(R, ath, rec):
    """(Status, Pending Status) do lado da assinatura, para quem tem Athena."""
    if _is_fwd_start(ath['instrument']):
        return ST_MANUAL, PS_MANUAL
    ps = R._pc_signature_pending_status(rec, ath['trade'], ath['settle'])
    return (ST_OK if R._pc_is_ok_status(ps) else ST_PENDING), ps


def executar(ref=None, fep_path=None, athena_records=None):
    """Roda o batimento do trade date `ref` (padrão: o D-1 ANBIMA)."""
    R = _routes()
    avisos = []
    dia = ref or _cgd.dia_util_anterior()

    fep, fep_file, fep_info = ler_fep(dia, avisos, fep_path)
    ath, ath_info = ler_athena(dia, avisos, athena_records)

    by_spn = R._fxo_refdata_by_spn()
    by_tax = _refdata_by_tax()

    fep_por_chave = {}
    for f in fep:
        fep_por_chave.setdefault(f['contrato'], []).append(f)
    ath_por_chave = {a['deal']: a for a in ath}

    sem_cadastro = set()
    linhas = []

    def linha(f_list, a):
        f = f_list[0] if f_list else None
        rec_f = by_tax.get(_tax_key(f['cnpj'])) if (f and f.get('cnpj') not in (None, '')) else None
        rec_a = by_spn.get(R._norm_spn(a['spn'])) if (a and a['spn']) else None
        rec = rec_a or rec_f or {}
        if f and rec_f is None and _tax_key(f.get('cnpj')):
            sem_cadastro.add('CNPJ {}'.format(_cgd._digits(f['cnpj']) or f['cnpj']))
        if a and rec_a is None:
            sem_cadastro.add('SPN {}'.format(a['spn'] or '(vazio) — ' + a['end_cp']))

        status, pending = '', ''
        if a:
            status, pending = _assinatura(R, a, rec)
        if f and not a:
            status = ST_MISSING_ATH
        elif a and not f:
            status = ST_MISSING_FEP
        if f_list and len(f_list) > 1:
            status = ST_DUPLICATED

        tenor = (a['settle'] - a['trade']).days if (a and a['settle'] and a['trade']) else None
        tax = (rec_f or {}).get('TAX ID') or (f and _cgd._digits(f.get('cnpj'))) \
            or (rec_a or {}).get('TAX ID') or ''
        return {
            'key': (f['contrato'] if f else a['deal']),
            'Status': status,
            COMMENT_COLUMN: '',
            'Pending Status': pending,
            'Trade Date': _cgd._fmt_date(dia),
            'FepWeb ID': f['contrato'] if f else '',
            'Athena ID': a['deal'] if a else '',
            'FepWeb Client': ((rec_f or {}).get('COUNTERPARTY') or (f or {}).get('cliente') or ''),
            'Athena Client': ((rec_a or {}).get('COUNTERPARTY') or (a or {}).get('desc') or ''),
            'CNPJ': str(tax or ''),
            'SPN': (a['spn'] if a else '') or str((rec_f or {}).get('SPN', '') or ''),
            'Signature Type': str(rec.get('SIGNATURE TYPE', '') or '').strip(),
            'Instrument Type': a['instrument'] if a else '',
            'Settlement Date': _cgd._fmt_date(a['settle']) if a else '',
            'Tenor': tenor,
            'Quantity Ccy': a['qty_ccy'] if a else '',
            'Quantity': a['qty'] if a else None,
            'Other Ccy': a['other_ccy'] if a else '',
            'Other Quantity': a['other_qty'] if a else None,
            'Strike': a['strike'] if a else None,
            'Publisher': a['publisher'] if a else '',
            'Cetip ID': a['cetip_id'] if a else '',
            'End Counterparty': a['end_cp'] if a else '',
            'FepWeb Type': f['tipo'] if f else '',
            'FepWeb Count': len(f_list) if f_list else 0,
        }

    for chave, f_list in fep_por_chave.items():
        linhas.append(linha(f_list, ath_por_chave.get(chave)))
    for chave, a in ath_por_chave.items():
        if chave not in fep_por_chave:
            linhas.append(linha(None, a))

    if sem_cadastro:
        amostra = sorted(sem_cadastro)
        avisos.append('Sem cadastro no Reference Data ({}): {}{} — o cliente sai como veio '
                      'da fonte e a assinatura conta como não cadastrada.'.format(
                          len(amostra), '; '.join(amostra[:8]),
                          '…' if len(amostra) > 8 else ''))

    ordem = {s: i for i, s in enumerate(STATUS_ORDER)}
    linhas.sort(key=lambda r: (ordem.get(r['Status'], 99),
                               _cgd._norm(r['Athena Client'] or r['FepWeb Client']),
                               r['key']))
    aplicar_comentarios(linhas)

    return {
        'ref': dia.strftime('%Y-%m-%d'), 'ref_fmt': _cgd._fmt_date(dia),
        'generated_at': datetime.now().strftime('%d/%m/%Y %H:%M'),
        'fep_file': fep_file,
        'fep_count': len(fep), 'athena_count': len(ath),
        'fep_info': fep_info, 'athena_info': ath_info,
        'columns': list(COLUMNS),
        'rows': linhas, 'counts': contar(linhas),
        'warnings': avisos,
    }


def contar(linhas):
    c = {k: 0 for k in COUNT_KEYS.values()}
    for r in linhas:
        k = COUNT_KEYS.get(r.get('Status'))
        if k:
            c[k] += 1
    return c


# ── Comentários ──────────────────────────────────────────────────────────────
# O comentário é do TRADE, não da execução do dia (o desenho da Recon FXO): mora
# fora do cache por data e é reaplicado em toda leitura, então rodar de novo o
# mesmo dia não apaga o que a mesa escreveu.

def load_comments():
    try:
        data = _store.read(_COMMENTS_PATH)
    except Exception:
        return {}
    if not isinstance(data, dict):
        return {}
    return {str(k): str(v) for k, v in data.items() if str(v or '').strip()}


def save_comment(key, comment):
    k = contract_key(key)
    if not k:
        raise ValueError('Sem a chave da operação não há o que comentar.')
    txt = str(comment or '').strip()
    # Ler-alterar-gravar sob lock: dois usuários comentando linhas diferentes
    # gravariam o arquivo inteiro um por cima do outro.
    with _COMMENTS_LOCK:
        data = load_comments()
        if txt:
            data[k] = txt
        else:
            data.pop(k, None)
        os.makedirs(os.path.dirname(_COMMENTS_PATH), exist_ok=True)
        _routes()._atomic_write_json(_COMMENTS_PATH, data)
    return txt


def aplicar_comentarios(linhas, comments=None):
    if comments is None:
        comments = load_comments()
    for r in linhas:
        r[COMMENT_COLUMN] = comments.get(str(r.get('key', '')), '')
    return linhas


# ── Cache do dia ─────────────────────────────────────────────────────────────

def _cache_path(ref):
    # Sem data, o dia é o MESMO default do `executar` (o D-1), não hoje.
    d = _cgd._parse_date(ref) or _cgd.dia_util_anterior()
    return os.path.join(_CACHE_DIR, d.strftime('%Y'), d.strftime('%m'),
                        'conf-matching_{}.json'.format(d.strftime('%Y%m%d')))


def salvar(res):
    path = _cache_path(res.get('ref'))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    # Pelo FUNIL: a gravação vai para o BANCO do caminho (armazém, §434).
    _routes()._atomic_write_json(path, res)
    return path


def carregar(ref=None):
    """O resultado já rodado daquele dia (com os comentários de AGORA), ou
    `None`. Rodar é decisão de quem opera, não do carregamento da página."""
    try:
        res = _store.read(_cache_path(ref))
    except _store.BancoOcupado:
        raise           # ocupado/ilegível não é "ninguém rodou" (§434): vira 503
    except Exception:
        return None
    if not isinstance(res, dict):
        return None
    aplicar_comentarios(res.get('rows') or [])
    return res
