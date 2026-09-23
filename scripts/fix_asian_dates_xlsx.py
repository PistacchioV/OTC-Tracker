# -*- coding: utf-8 -*-
"""Corrige as datas da Média Asiática de uma planilha do Live Position Option.

    python scripts/fix_asian_dates_xlsx.py "C:\\Users\\<sid>\\Downloads\\Live Position Option OTC Tracker - Sistema de Gestão OTC.xlsx"
    python scripts/fix_asian_dates_xlsx.py <arquivo> --dry-run      # só o relatório
    python scripts/fix_asian_dates_xlsx.py <arquivo> --feriados ipe.xlsx

Cria uma aba NOVA no mesmo arquivo (`--aba`, padrão `Ajustado`) e deixa a aba
original como está. As colunas antes do bloco da Média Asiática (até a BH) são
copiadas iguais; o bloco (`Média Asiática (data) 1…N`, da BI em diante) é
reescrito linha a linha:

  · todas as datas no MESMO mês → só estão fora de ORDEM: são reordenadas, as
    mesmas datas e a mesma quantidade;
  · datas em DOIS (ou mais) meses → a JANELA está errada: vale o mês com MAIS
    datas, e a linha passa a ter todo dia útil desse mês, do primeiro ao
    último, no calendário `--calendario` (padrão IPE). A quantidade pode mudar
    (21 datas que viram os 22 dias úteis do mês), e o relatório diz quando muda;
  · EMPATE entre dois meses não se decide: a linha fica como está e sai no
    relatório para a mesa olhar.

As datas saem como DATA de verdade (`dd/mm/aaaa`), não como texto. Célula
mexida ganha fundo amarelo, e a aba `Log Ajuste Datas` diz, por linha, o que foi
feito e por quê.

**Os feriados são os do Holidays Calendar do app** (o mesmo arquivo que a tela
de feriados grava), lidos pelo armazém — rode de dentro do checkout da
instância. Fora dele, `--feriados` aponta para uma planilha (datas na coluna A)
ou um `.json` do Holidays. Calendário VAZIO, ou sem nenhum feriado num ano que
a planilha usa, PARA o script: sem feriado, "dia útil" vira "dia de semana", e
a janela sairia com o dia de bolsa fechada dentro, sem erro nenhum.

Antes de gravar, o original é copiado ao lado como `<nome> - original.xlsx`
(só na primeira vez). Rodar de novo substitui a aba `Ajustado`, e é sempre a aba
original que se lê.
"""
import argparse
import json
import os
import re
import shutil
import sys
import unicodedata
from collections import Counter
from copy import copy
from datetime import date, datetime, timedelta

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
sys.path.insert(0, ROOT)
os.environ['OTC_DISABLE_SCHEDULERS'] = '1'
# Fora do Windows a raiz do share é obrigatória (§9); aqui ela não é usada.
os.environ.setdefault('OTC_SHARED_DRIVE_ROOT', ROOT)

from openpyxl import load_workbook                      # noqa: E402
from openpyxl.styles import Font, PatternFill            # noqa: E402
from openpyxl.utils import get_column_letter             # noqa: E402

FORMATO = 'dd/mm/yyyy'
AMARELO = PatternFill('solid', fgColor='FFF2CC')
_EXCEL_ZERO = date(1899, 12, 30)


def _norm(texto):
    t = unicodedata.normalize('NFKD', str(texto or ''))
    return ''.join(c for c in t if not unicodedata.combining(c)).lower().strip()


def para_data(valor):
    """A data de uma célula: data do Excel, serial numérico ou texto
    `dd/mm/aaaa` / `aaaa-mm-dd`. Vazio → None; ilegível → ValueError."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)):
        if 20000 < valor < 80000:                     # serial do Excel (1954–2119)
            return _EXCEL_ZERO + timedelta(days=int(valor))
        raise ValueError(repr(valor))
    texto = str(valor).strip()
    for formato in ('%d/%m/%Y', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d/%m/%Y %H:%M:%S'):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError(repr(valor))


# ── feriados ────────────────────────────────────────────────────────────────

def feriados_do_app(nome):
    from apps.pages.precificador import calendario
    return set(calendario._feriados_do_arquivo(nome))


def feriados_do_arquivo(caminho):
    if caminho.lower().endswith('.json'):
        with open(caminho, encoding='utf-8') as fh:
            itens = json.load(fh)
        brutos = [i.get('date') if isinstance(i, dict) else i for i in itens]
    else:
        wb = load_workbook(caminho, read_only=True, data_only=True)
        brutos = [linha[0] for linha in wb.active.iter_rows(values_only=True) if linha]
    datas = set()
    for b in brutos:
        try:
            d = para_data(b)
        except ValueError:
            continue                                   # cabeçalho, texto solto
        if d:
            datas.add(d)
    return datas


def dias_uteis_do_mes(ano, mes, feriados):
    d, saida = date(ano, mes, 1), []
    while d.month == mes:
        if d.weekday() < 5 and d not in feriados:
            saida.append(d)
        d += timedelta(days=1)
    return saida


# ── a planilha ──────────────────────────────────────────────────────────────

def achar_bloco(ws):
    """(linha do cabeçalho, [colunas da Média Asiática]) — pelo RÓTULO, nas
    dez primeiras linhas."""
    for r in range(1, min(ws.max_row, 10) + 1):
        cols = [c for c in range(1, ws.max_column + 1)
                if _norm(ws.cell(r, c).value).startswith('media asiatica')]
        if cols:
            return r, cols
    return None, []


def colunas_de_id(ws, linha_cab):
    """As colunas que identificam o trade no log (contrato/identificador)."""
    saida = []
    for c in range(1, ws.max_column + 1):
        n = _norm(ws.cell(linha_cab, c).value)
        if ('contrato' in n or 'identificador' in n or n in ('b3 id', 'deal')) \
                and 'media asiatica' not in n:
            saida.append(c)
    return saida[:3] or [1]


def decidir(datas):
    """O que fazer com as datas de uma linha →
    (acao, meses {(a,m): n}, (a,m) escolhido ou None)."""
    meses = Counter((d.year, d.month) for d in datas)
    if len(meses) == 1:
        return 'reordenar', meses, next(iter(meses))
    ranking = meses.most_common()
    if ranking[0][1] == ranking[1][1]:
        return 'empate', meses, None
    return 'janela', meses, ranking[0][0]


def _fmt_meses(meses):
    return ', '.join('%02d/%d (%d)' % (m, a, n) for (a, m), n in sorted(meses.items()))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('arquivo')
    ap.add_argument('--aba-origem', help='aba a ler (padrão: a primeira)')
    ap.add_argument('--aba', default='Ajustado', help='nome da aba nova')
    ap.add_argument('--calendario', default='IPE', help='calendário do Holidays (padrão IPE)')
    ap.add_argument('--feriados', help='planilha (coluna A) ou .json com os feriados, no lugar do app')
    ap.add_argument('--dry-run', action='store_true', help='só o relatório, não grava')
    args = ap.parse_args(argv)

    wb = load_workbook(args.arquivo)
    origem = wb[args.aba_origem] if args.aba_origem else wb.worksheets[0]
    if origem.title in (args.aba, 'Log Ajuste Datas'):
        sys.exit('A aba de origem não pode ser a aba de saída (%s).' % origem.title)
    linha_cab, bloco = achar_bloco(origem)
    if not bloco:
        sys.exit('Nenhuma coluna "Média Asiática" nas dez primeiras linhas da aba %r.' % origem.title)
    print('Aba lida: %s · cabeçalho na linha %d · Média Asiática em %s:%s (%d colunas)'
          % (origem.title, linha_cab, get_column_letter(bloco[0]),
             get_column_letter(bloco[-1]), len(bloco)))
    if get_column_letter(bloco[0]) != 'BI':
        print('  aviso: o bloco começa na %s, não na BI' % get_column_letter(bloco[0]))
    ids = colunas_de_id(origem, linha_cab)

    # 1. ler e decidir
    planos, ilegiveis, anos = [], [], set()
    for r in range(linha_cab + 1, origem.max_row + 1):
        datas = []
        for c in bloco:
            try:
                d = para_data(origem.cell(r, c).value)
            except ValueError:
                ilegiveis.append('%s%d = %r' % (get_column_letter(c), r, origem.cell(r, c).value))
                continue
            if d:
                datas.append(d)
        if not datas:
            continue
        acao, meses, alvo = decidir(datas)
        planos.append((r, datas, acao, meses, alvo))
        if alvo:
            anos.add(alvo[0])
    if ilegiveis:
        print('\nCélulas que não são data (ficam como estão): %d' % len(ilegiveis))
        for x in ilegiveis[:10]:
            print('   ', x)

    # 2. feriados — calendário vazio ou ano sem feriado PARA (sem eles a
    #    janela sairia com dia de bolsa fechada dentro)
    precisa = sorted({p[4][0] for p in planos if p[2] == 'janela'})
    if args.feriados:
        feriados, fonte = feriados_do_arquivo(args.feriados), args.feriados
    else:
        feriados, fonte = feriados_do_app(args.calendario), 'Holidays %s do app' % args.calendario
    print('\nFeriados: %s — %d datas' % (fonte, len(feriados)))
    sem = [a for a in precisa if not any(f.year == a for f in feriados)]
    if precisa and sem:
        sys.exit('PARADO: o calendário não tem nenhum feriado em %s. Cadastre o %s na tela '
                 'de Holidays (ou passe --feriados) antes de ajustar as janelas.'
                 % (', '.join(map(str, sem)), args.calendario))
    for a in precisa:
        print('   %d: %s' % (a, ', '.join(f.strftime('%d/%m') for f in sorted(feriados) if f.year == a)))

    # 3. as novas datas de cada linha
    resultado = {}          # linha → (novas datas, texto do log, acao)
    cont = Counter()
    for r, datas, acao, meses, alvo in planos:
        if acao == 'reordenar':
            novas = sorted(datas)
            nota = 'Mesmo mês %02d/%d: reordenadas' % (alvo[1], alvo[0])
            if novas == datas:
                acao, nota = 'ok', 'Já em ordem, mesmo mês %02d/%d' % (alvo[1], alvo[0])
            dup = [d for d, n in Counter(datas).items() if n > 1]
            if dup:
                nota += ' · DATA REPETIDA: ' + ', '.join(d.strftime('%d/%m/%Y') for d in sorted(dup))
        elif acao == 'janela':
            novas = dias_uteis_do_mes(alvo[0], alvo[1], feriados)
            nota = 'Janela em %s → %02d/%d inteiro' % (_fmt_meses(meses), alvo[1], alvo[0])
            if len(novas) != len(datas):
                nota += ' · QUANTIDADE %d → %d' % (len(datas), len(novas))
        else:
            novas = datas
            nota = 'EMPATE entre %s: não ajustada' % _fmt_meses(meses)
        if len(novas) > len(bloco):
            nota += ' · %d dias úteis e só %d colunas: NÃO AJUSTADA' % (len(novas), len(bloco))
            novas, acao = datas, 'estouro'
        cont[acao] += 1
        resultado[r] = (novas, nota, acao)

    print('\nLinhas com data: %d' % len(planos))
    for k, rot in (('reordenar', 'reordenadas'), ('janela', 'janela ajustada'),
                   ('ok', 'já estavam certas'), ('empate', 'EMPATE (não ajustadas)'),
                   ('estouro', 'mais dias úteis que colunas (não ajustadas)')):
        if cont[k]:
            print('   %-45s %d' % (rot, cont[k]))
    mudou_qtd = [r for r, (n, nota, a) in resultado.items() if 'QUANTIDADE' in nota]
    if mudou_qtd:
        print('   %-45s %d' % ('com quantidade de datas diferente', len(mudou_qtd)))
    for r, (n, nota, a) in resultado.items():
        if a in ('empate', 'estouro'):
            print('   linha %d: %s' % (r, nota))

    if args.dry_run:
        print('\n--dry-run: nada gravado.')
        return 0

    # 4. a aba nova: cópia da origem + o bloco reescrito
    for nome in (args.aba, 'Log Ajuste Datas'):
        if nome in wb.sheetnames:
            del wb[nome]
    ws = wb.create_sheet(args.aba, index=wb.sheetnames.index(origem.title) + 1)
    for linha in origem.iter_rows():
        for cel in linha:
            novo = ws.cell(cel.row, cel.column, cel.value)
            if cel.has_style:
                novo.font, novo.border, novo.fill = copy(cel.font), copy(cel.border), copy(cel.fill)
                novo.number_format, novo.alignment = cel.number_format, copy(cel.alignment)
    for letra, dim in origem.column_dimensions.items():
        ws.column_dimensions[letra].width = dim.width
    ws.freeze_panes = origem.freeze_panes

    for r in range(linha_cab + 1, origem.max_row + 1):
        novas, nota, acao = resultado.get(r, (None, '', ''))
        for i, c in enumerate(bloco):
            antes = origem.cell(r, c).value
            cel = ws.cell(r, c)
            if novas is not None:
                depois = novas[i] if i < len(novas) else None
                cel.value = datetime(depois.year, depois.month, depois.day) if depois else None
                try:
                    igual = para_data(antes) == depois
                except ValueError:
                    igual = False
                if not igual and acao not in ('empate', 'estouro'):
                    cel.fill = AMARELO
            else:
                try:                               # linha sem ajuste: a data vira DATA
                    d = para_data(antes)
                    cel.value = datetime(d.year, d.month, d.day) if d else antes
                except ValueError:
                    pass
            if isinstance(cel.value, datetime):
                cel.number_format = FORMATO
    for c in bloco:
        ws.column_dimensions[get_column_letter(c)].width = max(
            ws.column_dimensions[get_column_letter(c)].width or 0, 12)

    log = wb.create_sheet('Log Ajuste Datas', index=wb.sheetnames.index(args.aba) + 1)
    cab = (['Linha'] + [str(origem.cell(linha_cab, c).value) for c in ids]
           + ['Ação', 'Datas antes', 'Datas depois', 'Detalhe'])
    log.append(cab)
    for c in log[1]:
        c.font = Font(bold=True)
    rotulo = {'reordenar': 'Reordenada', 'janela': 'Janela ajustada', 'ok': 'Sem alteração',
              'empate': 'Não ajustada (empate)', 'estouro': 'Não ajustada'}
    for r, datas, acao, meses, alvo in planos:
        novas, nota, acao = resultado[r]
        log.append([r] + [origem.cell(r, c).value for c in ids]
                   + [rotulo[acao], len(datas), len(novas), nota])
    for i, w in enumerate([8] + [22] * len(ids) + [22, 12, 12, 70], start=1):
        log.column_dimensions[get_column_letter(i)].width = w
    log.freeze_panes = 'A2'

    base, ext = os.path.splitext(args.arquivo)
    copia = base + ' - original' + ext
    if not os.path.exists(copia):
        shutil.copy2(args.arquivo, copia)
        print('\nOriginal copiado para: %s' % copia)
    wb.save(args.arquivo)
    print('Gravado: abas "%s" e "Log Ajuste Datas" em %s' % (args.aba, args.arquivo))
    return 0


if __name__ == '__main__':
    sys.exit(main())
