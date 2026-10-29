# -*- coding: utf-8 -*-
"""Onde foram parar as operações da planilha diária que sumiram da tela.

    python scripts/diag_pending_missing.py                       # a planilha do share
    python scripts/diag_pending_missing.py "C:\\caminho\\PENDING - Outstanding Confirmation OTC.xlsx"
    python scripts/diag_pending_missing.py --out sumidas.xlsx

A planilha `PENDING - Outstanding Confirmation OTC.xlsx` (a rotina das 10:45)
é a fila pendente NAQUELE momento; a tela do Pending Confirmation é o banco
AGORA. Quando uma operação está na planilha e não está na tela, há cinco
respostas diferentes, e cada uma se trata de um jeito — este script diz qual,
Trade Number por Trade Number:

  · NA FILA         — está no banco pending e segue pendente: a tela a mostra
                      (se não mostrou, é filtro/chip da tela, não dado);
  · RESOLVIDA       — o Pending Status AGORA é de resolvida (Concluded, Signed
                      Digitally, Exception…): saiu da fila por alguém ter
                      mudado o status (tela, Pending Update, esteira);
  · BACKLOG         — Trade Date passou de 12 meses;
  · FORA DO LUGAR   — segue pendente, mas mora no ok/backlog: a manutenção das
                      11:30 a devolve ao pending;
  · AUSENTE         — não está em NENHUM dos três bancos: foi APAGADA (Delete
                      da tela, Delete da recompra, script de reimportação).

E diz se ela estava na FOTO das 11:30 de hoje (o balde pending logo depois da
manutenção) — separa "sumiu antes das 11:30" de "sumiu depois".

Só LÊ, roda com o app de pé (abre pela camada `duckdb_read`), e banco que não
abre PARA com o motivo — lido como vazio, toda operação sairia AUSENTE.
"""
import argparse
import os
import sys
from collections import Counter
from datetime import date, datetime

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ['OTC_DISABLE_SCHEDULERS'] = '1'
os.environ.setdefault('OTC_SHARED_DRIVE_ROOT', ROOT)

import diag_pending_confirmation_db as D                      # noqa: E402

PLANILHA = 'PENDING - Outstanding Confirmation OTC.xlsx'


def _planilha_padrao():
    """A mesma pasta em que a rotina grava (`features/pcx/infra/persistence.DIR`)."""
    from apps.config import Config
    pasta = os.getenv('PCX_SPREADSHEET_DIR', os.path.join(
        Config.SHARED_DRIVE_ROOT, 'Confirmation', 'Derivativos', 'Movimento',
        'Pending Confirmation'))
    return os.path.join(pasta, PLANILHA)


def _txt(v):
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def ler_planilha(caminho):
    """{Trade Number: linha} da aba CONFIRMATIONS (cabeçalho na linha 1)."""
    import openpyxl
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb['CONFIRMATIONS'] if 'CONFIRMATIONS' in wb.sheetnames else wb.active
    it = ws.iter_rows(values_only=True)
    cab = [_txt(c) for c in next(it)]
    if 'Trade Number' not in cab:
        raise SystemExit('a planilha não tem a coluna "Trade Number" na linha 1: %s' % caminho)
    out = {}
    for linha in it:
        d = dict(zip(cab, (_txt(c) for c in linha)))
        tn = d.get('Trade Number', '')
        if tn:
            out[tn] = d
    return out


def ler_bancos(pasta):
    """{categoria: {Trade Number: linha}} dos três — PARA se um não abre."""
    out = {}
    for cat, arq in D.DBS.items():
        caminho = os.path.join(pasta, arq)
        if not os.path.isfile(caminho):
            raise SystemExit('banco %s não existe: %s' % (cat, caminho))
        try:
            cols, linhas = D.ler(caminho)
        except Exception as e:                                   # noqa: BLE001
            raise SystemExit('NÃO DEU PARA LER o banco %s (%s): %s: %s'
                             % (cat, caminho, type(e).__name__, e))
        idx = cols.index('Trade Number')
        out[cat] = {}
        for l in linhas:
            tn = _txt(l[idx])
            if tn:
                out[cat][tn] = dict(zip(cols, (_txt(v) for v in l)))
    return out


def foto_de_hoje():
    """Os Trade Numbers da foto das 11:30 de hoje — None se não há foto."""
    try:
        from apps.pages.platform.pending_confirmation import _PC_SNAPSHOT_DIR
        from apps.pages import duck_read
        hoje = date.today()
        caminho = os.path.join(_PC_SNAPSHOT_DIR, hoje.strftime('%Y'), hoje.strftime('%m'),
                               hoje.strftime('%d'),
                               'pending-confirmation_%s.json' % hoje.strftime('%Y%m%d'))
        dados = duck_read.day_records(caminho)
    except Exception as e:                                       # noqa: BLE001
        print('foto das 11:30 de hoje: não lida (%s: %s)' % (type(e).__name__, e))
        return None
    return {_txt(r.get('Trade Number')) for r in (dados or []) if isinstance(r, dict)}


def veredito(tn, bancos):
    """(diagnóstico, onde, Pending Status agora)."""
    from apps.pages.platform.pending_confirmation import _pc_is_ok_status, _pc_cutoff_date
    onde = [c for c in ('pending', 'ok', 'backlog') if tn in bancos[c]]
    if not onde:
        return 'AUSENTE', '', ''
    linha = bancos[onde[0]][tn]
    status = linha.get('Pending Status', '')
    td = D._parse_data(linha.get('Trade Date', ''))
    if _pc_is_ok_status(status):
        diag = 'RESOLVIDA'
    elif td and td < _pc_cutoff_date():
        diag = 'BACKLOG'
    elif onde[0] == 'pending':
        diag = 'NA FILA'
    else:
        diag = 'FORA DO LUGAR'
    return diag, '+'.join(onde), status


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('planilha', nargs='?', default=None,
                    help='a planilha da rotina (padrão: a do share)')
    ap.add_argument('--db-dir', default=None, help='outra pasta de bancos (o padrão é a do app)')
    ap.add_argument('--out', default=None, help='planilha de saída (.xlsx)')
    args = ap.parse_args()

    caminho = args.planilha or _planilha_padrao()
    print('planilha: %s' % caminho)
    print('  modificada em %s' % datetime.fromtimestamp(
        os.path.getmtime(caminho)).strftime('%d/%m/%Y %H:%M'))
    planilha = ler_planilha(caminho)
    pasta = args.db_dir or D._db_dir()
    print('bancos em: %s' % pasta)
    bancos = ler_bancos(pasta)
    print('  pending=%d  ok=%d  backlog=%d' % tuple(len(bancos[c]) for c in ('pending', 'ok', 'backlog')))
    foto = foto_de_hoje()
    print('planilha: %d operações\n' % len(planilha))

    linhas, cont, status_resolvida = [], Counter(), Counter()
    for tn, d in planilha.items():
        diag, onde, agora = veredito(tn, bancos)
        cont[diag] += 1
        if diag == 'RESOLVIDA':
            status_resolvida[agora] += 1
        na_foto = '' if foto is None else ('sim' if tn in foto else 'não')
        linhas.append([tn, d.get('Client', ''), d.get('Trade Date', ''),
                       d.get('Pending Status', ''), diag, onde, agora, na_foto])

    for k in ('NA FILA', 'RESOLVIDA', 'BACKLOG', 'FORA DO LUGAR', 'AUSENTE'):
        print('  %-14s %6d' % (k, cont[k]))
    if status_resolvida:
        print('\n  RESOLVIDA — o Pending Status de agora:')
        for s, n in status_resolvida.most_common():
            print('    %-34s %6d' % (s or '(vazio)', n))
    if foto is not None:
        aus = [l for l in linhas if l[4] == 'AUSENTE']
        if aus:
            print('\n  AUSENTE e na foto das 11:30: %d (sumiram DEPOIS das 11:30)'
                  % sum(1 for l in aus if l[7] == 'sim'))
        print('  foto das 11:30: %d operações' % len(foto))

    if args.out or any(l[4] != 'NA FILA' for l in linhas):
        import openpyxl
        destino = args.out or 'pending-sumidas-%s.xlsx' % datetime.now().strftime('%Y%m%d-%H%M')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Diagnostico'
        ws.append(['Trade Number', 'Client', 'Trade Date', 'Pending Status (planilha)',
                   'Diagnóstico', 'Banco agora', 'Pending Status agora', 'Na foto das 11:30'])
        for l in sorted(linhas, key=lambda x: (x[4] == 'NA FILA', x[4], x[1])):
            ws.append(l)
        for row in ws.iter_rows(min_row=2):       # Trade Number como TEXTO (§477)
            for c in row:
                c.number_format = '@'
        wb.save(destino)
        print('\nXLSX: %s' % os.path.abspath(destino))
    return 0


if __name__ == '__main__':
    sys.exit(main())
