# -*- coding: utf-8 -*-
"""Regressão da página New Deals › Swap › Cashflow (o swap da CEM, 21/09/2026).

O que este script prende:

1. o CONTRATO com o catálogo: as colunas do `domain` são as do
   `new_deals/catalog.py` (feature não importa feature — cópia com guarda),
   o cronograma é o `_SCHEDULE`, e as rotas que o `new_deals-product.html`
   chama existem (e o stub 501 não sobrou para o cashflow);
2. o Deal Ticket é o MESMO do Bullet (lido pela horizontal
   `platform/swap_deal_ticket`): a aba do DT com a tabela Cash Flow dá o MESMO
   deal que o Bullet daria, mais o cronograma, o tipo de amortização, a LOB e
   a paridade do DT da CEM; o PDF chega ao mesmo cronograma; o arquivo é lido
   pelo CONTEÚDO (xlsx renomeado, e-mail com o DT anexo);
3. o registro 0301 de Fluxo Não Constante (4.2.7 v00003) tem 2021 caracteres
   e leva, CAMPO A CAMPO, os mesmos valores do 0301 de Pagamento Final do
   Bullet nas posições deste layout (as regras das pernas são uma só);
4. o 0034 (o cronograma): header 38, registro 62, uma linha de 194 por fluxo,
   com o Meu Número do contrato, a ponta pela conta menor e o código do
   cadastro `swap-amortizacao`;
5. as lacunas são ESTRUTURADAS (código + params) e todo código tem texto nas
   três línguas do `_TRANS` da página;
6. import (dry-run → batch com `_replace` → Amend), edit (cronograma, 4-olhos,
   B3 ID = mapeamento → Pending Confirmation / Intrag), confirm, delete, send
   (cp1252, lote com lacuna recusa tudo) e o sino com o rótulo do catálogo;
7. os vizinhos: o Monitor soma a linha no card da LOB, a segregação das
   confirmações lê o cashflow, o backfill tem família para a pasta, e o
   e-mail do Economic Affirmation leva a tabela Cash Flow.

Roda em tmp: o cache, a pasta Conecta e a Intrag apontam para temporários, e
o sino, o Pending Confirmation e a esteira são espiões.
"""
import io
import json
import os
import re
import sys
import tempfile
import zipfile
from datetime import datetime, timedelta, timezone

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
sys.path.insert(0, ROOT)
from apps.config import Config as _Cfg                      # noqa: E402
_Cfg.FOUR_EYES = True   # prova a trava da PROD; na dev o maker/checker e desligado (§553)
os.environ.setdefault('OTC_DISABLE_SCHEDULERS', '1')
if os.name != 'nt' and not os.environ.get('OTC_SHARED_DRIVE_ROOT'):
    os.environ['OTC_SHARED_DRIVE_ROOT'] = tempfile.mkdtemp(prefix='otc-share-')

FALHAS = []


def check(nome, cond, extra=''):
    print(('  ok  ' if cond else '  FAIL ') + nome + (('  — ' + str(extra)) if (not cond and extra != '') else ''))
    if not cond:
        FALHAS.append(nome)


FLOWS = [('2026-09-15', '2026-12-15', 0.25, 63, 91), ('2026-12-15', '2027-03-15', 0.25, 58, 90),
         ('2027-03-15', '2027-06-07', 0.5, 58, 84)]


def _dt_sheet(ws, client, spn, vcp_holder, van_holder, payer, categoria, cashflow=True):
    """O DT do Swap Bullet (o mesmo do `check_swap_bullet`) + o que o de
    cashflow da CEM traz a mais (SUPOSIÇÃO de layout, a do `domain`)."""
    ws['A1'] = 'DEAL TICKET SWAP VCP'
    ws['A4'] = 'Swap Com Opção de Arrependimento - Call'
    ws['A7'] = 'Cliente'; ws['B7'] = client
    ws['A8'] = 'SPN'; ws['B8'] = spn
    ws['A9'] = 'Inicio'; ws['B9'] = datetime(2026, 9, 15); ws['D9'] = 'Funcionalidades'; ws['E9'] = 'Opção de Arrependimento'; ws['G9'] = 'Agenda de premios'; ws['H9'] = 'Sim'
    ws['A10'] = 'Vencimento'; ws['B10'] = datetime(2027, 6, 7); ws['D10'] = 'Ativo VCP'; ws['E10'] = vcp_holder; ws['G10'] = 'Data de Pagamento do Premio'; ws['H10'] = datetime(2026, 9, 16)
    ws['A11'] = 'Valor Base'; ws['B11'] = 'BRL 346.000,00'; ws['D11'] = 'Ativo Curva Vanilla'; ws['E11'] = van_holder; ws['G11'] = 'Pagador do Premio'; ws['H11'] = payer
    ws['G12'] = 'Premio'; ws['H12'] = 23355
    ws['A15'] = 'Curva VCP'; ws['D15'] = 'Curva Vanilla'; ws['G15'] = 'Informações Curvas VCP'
    ws['A16'] = 'Percentual'; ws['B16'] = 1.0; ws['B16'].number_format = '0.00%'; ws['D16'] = 'Percentual'; ws['E16'] = 1.0; ws['E16'].number_format = '0.00%'
    ws['G16'] = 'Preço Inicial(Cupom Limpo)'; ws['H16'] = '100.00% Spot'
    ws['A17'] = 'Categoria'; ws['B17'] = 'VCP'; ws['D17'] = 'Categoria'; ws['E17'] = 'JUROS'; ws['G17'] = 'Fonte de Informação'; ws['H17'] = 'Bloomberg'
    ws['A18'] = 'Curva'; ws['B18'] = 'MSFT US'; ws['D18'] = 'Curva'; ws['E18'] = 'PRÉ FIXADO BRL'; ws['G18'] = 'Data de Cotação'; ws['H18'] = datetime(2027, 6, 4)
    ws['A20'] = 'Sinal +/-'; ws['B20'] = '+'; ws['D20'] = 'Sinal +/-'; ws['E20'] = '+'; ws['G20'] = 'Denominação'; ws['H20'] = 'Quanto (ausencia variação cambial)'
    ws['A21'] = 'Juros'; ws['B21'] = 0; ws['B21'].number_format = '0.00%'; ws['D21'] = 'Juros'; ws['E21'] = 0; ws['E21'].number_format = '0.00%'; ws['H21'] = 'Cupom Limpo inicial em percentual'
    ws['A22'] = 'Lim Superior'; ws['B22'] = '117% Spot'; ws['H22'] = 'Preco in ativo Close 15-Sep-26'
    ws['A23'] = 'Lim Inferior'
    ws['A25'] = 'OBS: Caso aplicável, colocar também no campo denominação as informações dos limitadores'
    ws['A27'] = 'Informações do Indicador'
    ws['A29'] = 'Categoria da Curva VCP'; ws['B29'] = categoria
    ws['A30'] = 'Código'; ws['B30'] = 10615
    ws['A31'] = 'Curva VCP'; ws['B31'] = 'MSFT US'
    ws['A32'] = 'Descrição da Curva VCP'; ws['B32'] = 'Microsoft Corporation'
    ws['A34'] = 'Informações necessárias para registro'; ws['B34'] = 'Consultar o Manual de Operações de Swap'
    ws['A35'] = 'Link para o Manual de Operações'; ws['B35'] = 'www.b3.com.br'
    if not cashflow:
        return
    ws['A13'] = 'Tipo de Amortização'; ws['B13'] = 'Sobre Valor Base Original'
    ws['D13'] = 'LOB'; ws['E13'] = 'CEM'
    ws['A36'] = 'Paridade Inicial'; ws['B36'] = '5,4321'
    ws['A38'] = 'Cash Flow'
    for col, h in zip('ABCDEF', ('Data Início', 'Data Pagamento', 'Amortização (%)', 'Dias Úteis (BD/252)',
                                 'Dias Corridos (Act/360)', 'Data Fixing')):
        ws[col + '39'] = h
    for i, (ini, pg, am, du, dc) in enumerate(FLOWS):
        r = str(40 + i)
        ws['A' + r] = datetime.strptime(ini, '%Y-%m-%d'); ws['B' + r] = datetime.strptime(pg, '%Y-%m-%d')
        ws['C' + r] = am; ws['C' + r].number_format = '0.00%'; ws['D' + r] = du; ws['E' + r] = dc


def build_xlsx(cashflow=True):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active; ws.title = 'Deal Ticket Cliente (CETIP)'
    _dt_sheet(ws, 'Safra', 281808, 'Cliente', 'Banco JP Morgan', 'Cliente', 'ACOES INTERNACIONAIS', cashflow)
    ws2 = wb.create_sheet('B2B - Atacama')
    _dt_sheet(ws2, 'Atacama', 9632845, 'Banco JP Morgan', 'Atacama', 'Banco JP Morgan', 'INDICES INTERNACIONAIS', cashflow)
    ws3 = wb.create_sheet('Recap'); ws3['A1'] = 'Recap'; ws3['A2'] = 'nada aqui'
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


PDF_DT = (
    "DEAL TICKET SWAP VCP\n"
    "Swap Com Opção de Arrependimento - Call\n"
    "Cliente Safra\n"
    "SPN 281808\n"
    "Inicio 15-Sep-26 Funcionalidades Opção de Arrependimento Agenda de premios Sim\n"
    "Vencimento 7-Jun-27 Ativo VCP Cliente Data de Pagamento do Premio 16-Sep-26\n"
    "Valor Base BRL 346.000,00 Ativo Curva Vanilla Banco JP Morgan Pagador do Premio Cliente\n"
    "Premio 23.355,00\n"
    "Curva VCP Curva Vanilla Informações Curvas VCP\n"
    "Percentual 100,00% Percentual 100,00% Preço Inicial(Cupom Limpo) 100.00% Spot\n"
    "Categoria VCP Categoria JUROS Fonte de Informação Bloomberg\n"
    "Curva MSFT US Curva PRÉ FIXADO BRL Data de Cotação 04/jun/27\n"
    "Sinal +/- + Sinal +/- + Denominação Quanto (ausencia variação cambial)\n"
    "Juros 0,00% Juros 0,00% Cupom Limpo inicial em percentual\n"
    "Lim Superior 117% Spot Preco in ativo Close 15-Sep-26\n"
    "Lim Inferior\n"
    "OBS: Caso aplicável, colocar também no campo denominação as informações dos limitadores\n"
    "Informações do Indicador\n"
    "Categoria da Curva VCP ACOES INTERNACIONAIS\n"
    "Código 10615\n"
    "Curva VCP MSFT US\n"
    "Descrição da Curva VCP Microsoft Corporation\n"
    "Informações necessárias para registro Consultar o Manual\n"
)
PDF_CF = (
    "Cash Flow\n"
    "Data Início Data Pagamento Amortização (%) Dias Úteis Dias Corridos\n"
    "15/09/2026 15/12/2026 25,00% 63 91\n"
    "15/12/2026 15/03/2027 25,00% 58 90\n"
    "15/03/2027 07/06/2027 50,00% 58 84\n"
)


def _pos(line, a, b):
    return line[a - 1:b]


def main():
    from run import app
    from apps.pages import routes as R
    from apps.pages.features.new_deals import catalog
    from apps.pages.features.swap_bullet import domain as bullet
    from apps.pages.features.swap_bullet import queries as bullet_q
    from apps.pages.features.swap_cashflow import commands, domain, entrypoint, queries
    from apps.pages.features.swap_cashflow.infra import persistence
    from apps.pages.platform import swap_deal_ticket as dtk
    from apps.pages.platform import swap_new_deals as sw

    tmp = tempfile.mkdtemp(prefix='otc-swc-')
    persistence.cache_dir = lambda: os.path.join(tmp, 'cache')
    R.CONECTA_NEW_PATH = os.path.join(tmp, 'conecta')
    SINO = []
    _notif_orig = R._create_notification
    R._create_notification = lambda *a, **k: SINO.append(a)
    REF = [{'SPN': '281808', 'COUNTERPARTY': 'BANCO SAFRA S.A.', 'TAX ID': '58.160.789/0001-28',
            'B3 ACCOUNT': '74220.00-5'}]
    _ref_orig = R._refdata_records
    R._refdata_records = lambda: REF
    _rows_orig = R._mapping_rows

    def _rows(key):
        if key == 'le-spn':
            return [{'LE': 'ATACAMA', 'NAME': 'ATACAMA MULTIMERCADO FI', 'SPN': '9632845', 'NOTES': ''}]
        return _rows_orig(key)
    R._mapping_rows = _rows
    try:
        _main(R, app, catalog, bullet, bullet_q, commands, domain, entrypoint, queries, persistence, dtk, sw,
              tmp, SINO)
    finally:
        R._create_notification = _notif_orig
        R._refdata_records = _ref_orig
        R._mapping_rows = _rows_orig
    print('\n%s' % ('TUDO OK' if not FALHAS else '%d FALHA(S): %s' % (len(FALHAS), FALHAS)))
    return 1 if FALHAS else 0


def _main(R, app, catalog, bullet, bullet_q, commands, domain, entrypoint, queries, persistence, dtk, sw,
          tmp, SINO):
    print('== 1. o contrato com o catálogo e com a tela ==')
    pagina = catalog.PAGES['/new_deals-swap-cashflow']
    check('as colunas do domain são as do catálogo (campo e rótulo, na ordem)',
          [tuple(c[:2]) for c in pagina['columns']] == list(domain.SWC_COLUMNS),
          [c for c in pagina['columns'] if tuple(c[:2]) not in domain.SWC_COLUMNS])
    check('o cronograma é o _SCHEDULE do catálogo', [c[0] for c in pagina['schedule']] == list(domain.SCHEDULE_FIELDS)
          and pagina['schedule_field'] == domain.SCHEDULE_FIELD)
    check('o rótulo do sino é o `label` do catálogo', entrypoint.PAGE == pagina['label'])
    check('o rótulo tem destino nos três mapas', R._notif_page_url(entrypoint.PAGE) == '/new_deals-swap-cashflow')
    regras = {str(r): r for r in app.url_map.iter_rules()}
    tpl = io.open(os.path.join(ROOT, 'apps', 'templates', 'pages', 'new_deals-product.html'), encoding='utf-8').read()
    acoes = sorted(set(re.findall(r"PAGE\.api \+ '/([a-z0-9\-/]+)", tpl)))
    # O Mapping B3 ID só é chamado pela página que o declara no catálogo
    # (`mapping`, a Cetes — §650); o Cashflow não tem o botão.
    if not pagina.get('mapping'):
        acoes = [x for x in acoes if x != 'mapping-b3']
    check('toda ação que a tela chama tem rota ESTÁTICA', all(('/api/new-deals/swap-cashflow/' + a) in regras for a in acoes),
          [a for a in acoes if ('/api/new-deals/swap-cashflow/' + a) not in regras])
    check('as ações são as esperadas', set(acoes) >= {'import-file', 'cache/batch', 'cache/search', 'economic-affirmation',
                                                      'preview', 'delete', 'confirm', 'send-conecta', 'edit', 'add'}, acoes)
    check('o stub 501 não atende mais o cashflow',
          regras['/api/new-deals/swap-cashflow/<path:resto>'].endpoint.endswith('api_swap_cashflow_unknown'))

    print('== 2. o Deal Ticket (o do Bullet + a tabela Cash Flow) ==')
    xlsx = build_xlsx()
    sheets = sw.sheets_from_xlsx(xlsx)
    grade, flows = domain.split_cashflow_grid(sheets[0][1])
    check('três fluxos lidos da tabela', len(flows) == 3, flows)
    check('fluxo 1: datas ISO, 25% (formato de porcentagem do Excel), DU/DC inteiros',
          flows[0] == {'StartDate': '2026-09-15', 'PaymentDate': '2026-12-15', 'AmortizationPct': '25',
                       'BusinessDays': '63', 'CalendarDays': '91', 'FixingDate': ''}, flows[0])
    check('a tabela sai da grade (o parser do DT não a lê como rótulo)',
          not any('Data Pagamento' in str(c) or 'Cash Flow' == str(c) for r in grade for c in r))
    raw = dtk.parse_dt_grid(grade)
    base = dtk.deal_from_raw(dtk.parse_dt_grid(sw.sheets_from_xlsx(build_xlsx(cashflow=False))[0][1]), '2026-09-16')
    cf = domain.deal_from_raw(raw, flows, '2026-09-16', domain.cem_extras_grid(grade))
    comuns = [k for k in bullet.SWB_FIELDS if k not in ('Maker', 'Checker', 'Type', 'LOB', 'FXStart')]
    check('os campos do Bullet saem IGUAIS aos do DT sem a tabela (um parser só)',
          [k for k in comuns if cf.get(k) != base.get(k)] == [], [k for k in comuns if cf.get(k) != base.get(k)])
    check('tipo Fluxo de Caixa, chave SWC-, Deal e B3 ID em branco',
          cf['Type'] == 'Fluxo de Caixa' and cf['_id'].startswith('SWC-') and cf['Deal'] == '' and cf['B3ID'] == '')
    check('cronograma no deal e a contagem', len(cf['CashFlows']) == 3 and cf['Flows'] == '3')
    check('o que o DT da CEM traz a mais: LOB, tipo de amortização, paridade inicial',
          cf['LOB'] == 'CEM' and cf['AmortizationType'] == 'Sobre Valor Base Original' and cf['FXStart'] == '5.4321',
          (cf['LOB'], cf['AmortizationType'], cf['FXStart']))
    sem_lob = domain.deal_from_raw(dtk.parse_dt_grid(sw.sheets_from_xlsx(build_xlsx(cashflow=False))[0][1]), [], '2026-09-16')
    check('DT sem LOB: a LOB fica em BRANCO (não se chuta EDG)', sem_lob['LOB'] == '')
    check('linha do topo com datas NÃO é tabela (Inicio | data | Vencimento | data)',
          domain.split_cashflow_grid([['Inicio', '2026-09-15', 'Vencimento', '2027-06-07'],
                                      ['Data Pagamento', '2026-12-15']])[1] == [])
    # PDF: o mesmo cronograma, e o resto do texto chega ao MESMO deal do Bullet.
    texto, fl_pdf = domain.split_cashflow_text(PDF_DT + PDF_CF)
    check('PDF: os três fluxos pela ordem do cabeçalho', [(f['StartDate'], f['PaymentDate'], f['AmortizationPct'], f['BusinessDays'])
                                                          for f in fl_pdf] == [(a, b, str(int(c * 100)), str(d)) for a, b, c, d, _e in FLOWS], fl_pdf)
    check('PDF: sem a tabela o texto dá o MESMO raw do DT', dtk.parse_dt_text(texto) == dtk.parse_dt_text(PDF_DT))
    # Pelo CONTEÚDO: xlsx com nome de .txt, e o DT anexo num .eml.
    itens = sw.read_upload('dt.txt', xlsx)
    check('xlsx renomeado para .txt é lido pelo conteúdo', [t for _k, t, _p in itens][:2] == ['Deal Ticket Cliente (CETIP)', 'B2B - Atacama'])
    from email.message import EmailMessage
    em = EmailMessage(); em['Subject'] = 'DT swap'; em['From'] = 'a@b.com'; em['To'] = 'c@d.com'
    em.set_content('segue'); em.add_attachment(xlsx, maintype='application',
                                               subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                                               filename='DT CEM.xlsx')
    itens_eml = sw.read_upload('mail.eml', em.as_bytes())
    check('.eml: o DT anexo vira as abas', len(itens_eml) == 3 and itens_eml[0][1].startswith('DT CEM.xlsx › '))
    try:
        sw.read_upload('x.bin', b'\x00\x01 nada')
        check('arquivo que não é DT recusa', False)
    except ValueError:
        check('arquivo que não é DT recusa', True)

    print('== 3. o 0301 de Fluxo Não Constante = as regras do Bullet nas posições do 4.2.7 ==')
    accounts = queries.own_accounts()
    check('contas próprias do b3-accounts', accounts.get('JPM') == '73760009' and accounts.get('ATACAMA') == '85398005', accounts)
    commands.enrich(cf)
    check('enrich: a contraparte do Reference Data pela SPN (conta própria, sem CNPJ)',
          cf['Client'] == 'BANCO SAFRA S.A.' and cf['ClientAccount'] == '74220005' and cf['ClientRefData'] == 'ok')
    codes = queries.codes_for(cf)
    check('código da amortização pelo cadastro swap-amortizacao', codes['amortization'] == '00', codes.get('amortization'))
    faltas = domain.missing_for_send(cf, codes, accounts, codes['amortization'])
    check('sem lacunas', faltas == [], domain.lacuna_texts(faltas))
    MY = '8985710698'
    for view, deal in (('client', cf),):
        vals = domain.contract_record_values(deal, view, accounts, codes, MY)
        linha = commands._record(commands.CONTRACT_FI_KEY, vals, view, deal)
        check('%s: 2021 caracteres' % view, len(linha) == domain.CONTRACT_RECORD_LENGTH, len(linha))
        bv = bullet.swap_record_values(dict(deal), view, accounts, bullet_q.codes_for(deal), MY)
        pares = [(str(s), str(s)) for s in range(1, 19)] + [(str(s), str(s)) for s in range(26, 40)] + \
            [('47', '41'), ('48', '42'), ('49', '43'), ('50', '44'), ('51', '45'), ('52', '46'),
             ('53', '47'), ('54', '48'), ('55', '49'), ('56', '50'), ('116', '98')]
        # Sinal/Juros da curva VCP (os 36/37 da Contraparte aqui) o cashflow
        # leva e o Bullet não (equipe, 26/10/2026)
        vcp_sj = {str(b + o) for b, sd in ((26, 'A'), (33, 'B'))
                  if domain.dtk._curve(deal, sd)['category'] == 'VCP' for o in (3, 4)}
        dif = [(b, c) for b, c in pares if b not in ('15', '17') and c not in vcp_sj and bv[b] != vals[c]]
        check('%s: os campos que os dois layouts têm saem IGUAIS aos do Bullet' % view, dif == [], dif)
        check('%s: a curva VCP leva Sinal e Juros (%% aa)' % view,
              all(vals[c].strip() for c in vcp_sj) and bool(vcp_sj), {c: vals[c] for c in vcp_sj})
        check('%s: posições do manual (tipo, Meu Número, contas, datas, valor base, curva, LOB)' % view,
              _pos(linha, 1, 10) == 'SWAP 10301' and _pos(linha, 11, 20) == MY and _pos(linha, 21, 28) == '73760009'
              and _pos(linha, 53, 60) == '74220005' and _pos(linha, 85, 92) == '20260915' and _pos(linha, 93, 100) == '20270607'
              and _pos(linha, 103, 118) == '0000000034600000' and _pos(linha, 121, 122) == '00'
              and _pos(linha, 634, 636) == 'C99' and _pos(linha, 691, 693) == 'C00'
              and _pos(linha, 1125, 1146) == '0000000000000100000000' and _pos(linha, 1884, 1897).strip() == 'CEM')
    hdr = R._fi_build_line(commands.CONTRACT_FI_KEY, 'header', domain.contract_header_values('JPMORGANBM', '20260916'))
    check('header do 0301 v00003', hdr == 'SWAP 00301JPMORGANBM          2026091600003', hdr)

    print('== 4. o 0034 (o cronograma) ==')
    h, reg, linhas = domain.flow_values(cf, 'client', accounts, MY, '1234567890', codes['amortization'], 'JPMORGANBM', '20260916')
    lh = R._fi_build_line(commands.FLOW_FI_KEY, 'header', h)
    lr = R._fi_build_line(commands.FLOW_FI_KEY, 'registro', reg)
    lf = [R._fi_build_line(commands.FLOW_FI_KEY, 'fluxo', x) for x in linhas]
    check('header 38 = SWAP 0 0034 participante data', lh == 'SWAP 00034JPMORGANBM          20260916', lh)
    check('registro 62: 3 eventos, MN do contrato, MN próprio, contas, papel (conta menor = JPM → 00), amortização 00',
          lr == 'SWAP 10034' + '0003' + MY + '1234567890' + '73760009' + '74220005' + '00' + '00' + ' ' * 8, lr)
    check('uma linha de 194 por fluxo', len(lf) == 3 and all(len(x) == domain.FLOW_LINE_LENGTH for x in lf))
    check('fluxo: data de pagamento nas duas pontas e a taxa 9(03)v9(05)',
          _pos(lf[0], 1, 8) == '20261215' and _pos(lf[0], 50, 57) == '02500000' and _pos(lf[0], 58, 65) == '20261215'
          and _pos(lf[0], 107, 114) == '02500000' and _pos(lf[2], 50, 57) == '05000000', lf[0][:120])
    # as taxas de juros das duas pontas (equipe, 26/10/2026): as da curva de
    # cada uma no 0301 — P1 = a ponta da menor conta (aqui a Parte, a JPM)
    vcf = domain.contract_record_values(cf, 'client', accounts, codes, MY)
    _h, _r, lfc = domain.flow_values(cf, 'client', accounts, MY, '1234567890', codes['amortization'],
                                     'JPMORGANBM', '20260916', codes)
    lfc = [R._fi_build_line(commands.FLOW_FI_KEY, 'fluxo', x) for x in lfc]
    check('fluxo: Sinal e Taxa de Juros P1 (9-17) = os da Parte no 0301; P2 (66-74) = os da Contraparte',
          all(_pos(x, 9, 10) == vcf['29'] and _pos(x, 11, 17) == vcf['30']
              and _pos(x, 66, 67) == vcf['36'] and _pos(x, 68, 74) == vcf['37'] for x in lfc)
          and all(len(x) == domain.FLOW_LINE_LENGTH for x in lfc),
          (lfc[0][8:17], lfc[0][65:74], vcf['29'], vcf['30'], vcf['36'], vcf['37']))
    # o exemplo da equipe: P1 0,95% e P2 4,8824% (USD, TAXAS DE CAMBIO)
    cfx = dict(cf, CurveACategory='JUROS', CurveARate='0.9500', CurveASign='+',
               CurveBCategory='TAXAS DE CAMBIO', CurveBRate='4.8824', CurveBSign='+')
    _h, _r, lfx = domain.flow_values(cfx, 'client', accounts, MY, '1', codes['amortization'],
                                     'JPMORGANBM', '20260916', dict(codes, signA='00', signB='00'))
    lfx = [R._fi_build_line(commands.FLOW_FI_KEY, 'fluxo', x) for x in lfx]
    check('   no exemplo: 00 0009500 na P1 e 00 0048824 na P2; a VCP também leva a dela',
          all(_pos(x, 9, 17) == '000009500' and _pos(x, 66, 74) == '000048824' for x in lfx)
          and lfc[0][65:74].strip(), (lfx[0][8:17], lfx[0][65:74], lfc[0][65:74]))
    ha, rega, _l = domain.flow_values(cf, 'atacama', accounts, MY, '1', '00', 'INTRAGATACAMAFDO', '20260916')
    check('na visão da Atacama a ponta da parte é 01 (85398005 > 73760009)', rega['9'] == '01')

    print('== 5. lacunas estruturadas e o _TRANS ==')
    ruim = dict(cf, CashFlows=[{'PaymentDate': '2026-12-15'}, {'PaymentDate': '2026-11-01'}], AmortizationType='',
                LOB='', CurveBCategory='JUROS INTERNACIONAIS')
    lac = domain.missing_for_send(ruim, codes, accounts, '')
    cods = [x['code'] for x in lac]
    check('ordem, vencimento, amortização, LOB e curva internacional viram lacuna com código',
          {'swc_flow_order', 'swc_last_flow_maturity', 'swc_no_amortization', 'swc_no_lob', 'swc_intl_curve'} <= set(cods), cods)
    check('lacuna é {code, params, text}', all(set(x) == {'code', 'params', 'text'} for x in lac))
    srcs = ''.join(io.open(os.path.join(ROOT, 'apps', 'pages', 'features', 'swap_cashflow', f), encoding='utf-8').read()
                   for f in ('domain.py', 'commands.py', 'entrypoint.py'))
    usados = set(re.findall(r"'((?:swc|nd)_[a-z_]+)'", srcs)) - {'swc_'}
    for lang in ('en', 'br', 'es'):
        bloco = re.search(r'\n    ' + lang + r': \{(.*?)\n    \}', tpl, re.S)
        chaves = set(re.findall(r'\b([ew]_[a-z_]+)\s*:', bloco.group(1))) if bloco else set()
        faltam = sorted(c for c in usados if ('e_' + c) not in chaves and ('w_' + c) not in chaves)
        check('_TRANS.%s tem texto para todo código do servidor' % lang, not faltam, faltam)

    print('== 6. import → batch → edit → confirm → send → delete ==')
    res = commands.import_upload('dt.xlsx', xlsx, datetime(2026, 9, 16), sid='A111111', dry_run=True)
    check('dry-run: dois deals, Recap ignorada, nada gravado', len(res['deals']) == 2 and res['ignored'] == ['Recap']
          and res['imported'] == 0 and not os.path.isdir(os.path.join(tmp, 'cache')))
    cli, b2b = res['deals']
    check('B2B: par JPM x ATACAMA pela SPN do le-spn', b2b['Pair'] == 'JPM x ATACAMA' and b2b['ClientAccount'] == '85398005')
    n = commands.persist_deals([dict(cli), dict(b2b)], sid='A111111')
    fp, lst, idx = queries.find(cli['_id'], '2026-09-16')
    check('batch grava no arquivo-dia da Trade Date, na pasta Swap/Cashflow', n == 2 and idx is not None
          and fp.endswith(os.path.join('2026', '09', '20260916_swapcashflow.json')))
    e = lst[idx]
    check('nasce New com os seis Meu Número', e['Status'] == 'New' and all(len(e[k]) == 10 for k in commands._MY_NUMBERS))
    nums = dict((k, e[k]) for k in commands._MY_NUMBERS)
    d, cod = commands.set_status(cli['_id'], '2026-09-16', 'Approved', sid='A111111')
    check('Confirm: New → Approved', d is not None and d['Status'] == 'Approved')
    commands.persist_deals([dict(cli, _replace=True)], sid='B222222')
    _f, l2, i2 = queries.find(cli['_id'], '2026-09-16')
    check('batch com _replace: Approved vira Amend, Meu Número preservado', l2[i2]['Status'] == 'Amend'
          and all(l2[i2][k] == nums[k] for k in commands._MY_NUMBERS))
    d, cod = commands.set_status(cli['_id'], '2026-09-16', 'Approved', sid='C333333')
    # Amend → Approved direto, como o New (mesa, 23/09/2026): Pending é só do Edit.
    check('Amend → Approved (maker = quem confirmou)',
          d['Status'] == 'Approved' and d['Maker'] == 'C333333' and d['Checker'] == '')
    commands.set_status(cli['_id'], '2026-09-16', 'Pending', sid='C333333')
    d2, cod2 = commands.set_status(cli['_id'], '2026-09-16', 'Approved', sid='C333333')
    check('Pending: o maker não aprova o próprio', d2 is None and cod2 == 'swc_maker_checker')
    novo = [dict(f) for f in cli['CashFlows']] + [{'StartDate': '2027-06-07', 'PaymentDate': '2027-06-07'}]
    d3, mapped = commands.edit(cli['_id'], '2026-09-16', {'CashFlows': novo[:3] + [{'PaymentDate': '', 'StartDate': ''}]},
                               sid='D444444')
    check('edit do cronograma: linha em branco cai, contagem refeita, Pending com maker', d3['Flows'] == '3'
          and d3['Status'] == 'Pending' and d3['Maker'] == 'D444444' and not mapped)
    d4, _c = commands.set_status(cli['_id'], '2026-09-16', 'Approved', sid='E555555')
    check('Pending → Approved por outro usuário', d4['Status'] == 'Approved' and d4['Checker'] == 'E555555')
    files = commands.preview(queries.find(b2b['_id'], '2026-09-16')[1][queries.find(b2b['_id'], '2026-09-16')[2]])
    check('preview do B2B: contrato + cronograma + prêmio × Banco e Atacama',
          [(f['kind'], f['view']) for f in files] == [('swap', 'bank'), ('flow', 'bank'), ('premium', 'bank'),
                                                      ('swap', 'atacama'), ('flow', 'atacama'), ('premium', 'atacama')])
    check('preview campo a campo com o rótulo do template', any(x['field'] == 'Quantidade de eventos de Fluxo'
                                                                 and x['value'] == '0003' for x in files[1]['fields']))
    # Lote com lacuna recusa tudo, dizendo qual (estruturado).
    _f, lb, ib = queries.find(b2b['_id'], '2026-09-16'); lb[ib]['AmortizationType'] = ''
    persistence.write_day(_f, lb)
    try:
        commands.send([{'deal_id': cli['_id'], 'trade_date': '2026-09-16'}, {'deal_id': b2b['_id'], 'trade_date': '2026-09-16'}])
        check('lote com lacuna recusa tudo', False)
    except commands.EnvioRecusado as exc:
        check('lote com lacuna recusa tudo, dizendo o deal e o código',
              any(x['deal_id'] == b2b['_id'] and x['code'] == 'swc_no_amortization' for x in exc.itens))
    check('nada foi escrito', not os.path.isdir(R.CONECTA_NEW_PATH) or not os.listdir(R.CONECTA_NEW_PATH))
    lb[ib]['AmortizationType'] = 'Sobre Valor Base Original'; persistence.write_day(_f, lb)
    out = commands.send([{'deal_id': cli['_id'], 'trade_date': '2026-09-16'}, {'deal_id': b2b['_id'], 'trade_date': '2026-09-16'}],
                        sid='F666666')
    nomes = sorted(x['filename'] for x in out['files'])
    check('nove arquivos: contrato, cronograma e prêmio × Cliente/Banco/Atacama, com CF e a LOB no nome',
          nomes == sorted(['73760009.SWAP_CF_CEM_CLIENTE.txt', '73760009.SWAP_CF_CEM_BANCO.txt',
                           '85398005.SWAP_CF_CEM_ATACAMA.txt', '73760009.FLUXO_SWAP_CF_CEM_CLIENTE.txt',
                           '73760009.FLUXO_SWAP_CF_CEM_BANCO.txt', '85398005.FLUXO_SWAP_CF_CEM_ATACAMA.txt',
                           '73760009.PREMIO_SWAP_CF_CEM_CLIENTE.txt', '73760009.PREMIO_SWAP_CF_CEM_BANCO.txt',
                           '85398005.PREMIO_SWAP_CF_CEM_ATACAMA.txt']), nomes)
    brut = io.open(os.path.join(R.CONECTA_NEW_PATH, '73760009.SWAP_CF_CEM_BANCO.txt'), 'rb').read().split(b'\n')
    check('contrato do Banco: 2021 BYTES em cp1252 (o travessão da denominação VCP em um byte)',
          len(brut) == 2 and len(brut[1]) == 2021 and b'\x96' in brut[1] and b'\xe2\x80\x93' not in brut[1])
    fl_txt = io.open(os.path.join(R.CONECTA_NEW_PATH, '85398005.FLUXO_SWAP_CF_CEM_ATACAMA.txt'), encoding='cp1252').read().split('\n')
    check('cronograma da Atacama: header INTRAGATACAMAFDO + registro + 3 fluxos', len(fl_txt) == 5
          and 'INTRAGATACAMAFDO' in fl_txt[0] and len(fl_txt[1]) == 62)
    _f, l5, i5 = queries.find(cli['_id'], '2026-09-16')
    check('vira Sent com os arquivos anotados', l5[i5]['Status'] == 'Sent' and '73760009.SWAP_CF_CEM_CLIENTE.txt' in l5[i5]['SentFiles'])

    print('== 6b. B3 ID = mapeamento → Pending Confirmation (cliente) / Intrag Swap (B2B) ==')
    calls, intrag = [], []
    _pc, _ie = R._pc_save_from_deal, R._intrag_engine

    class _Intrag(object):
        @staticmethod
        def _save_intrag_swap_entry(entry, start):
            intrag.append((entry, start))
    R._pc_save_from_deal = lambda deal, pt, pending_status=None, trade_number=None, source=None: \
        calls.append((deal, pt, pending_status, trade_number, source))
    R._intrag_engine = lambda: _Intrag
    try:
        d6, m6 = commands.edit(cli['_id'], '2026-09-16', {'B3ID': '26F04329911'}, sid='G777777')
        d7, m7 = commands.edit(b2b['_id'], '2026-09-16', {'B3ID': '26F04329912'}, sid='G777777')
    finally:
        R._pc_save_from_deal, R._intrag_engine = _pc, _ie
    check('cliente: Success, e a porta do Pending Confirmation com SWAP (Opção de Arrependimento), Pending OTC, chave B3 ID, LOB CEM',
          m6 and d6['Status'] == 'Success' and len(calls) == 1 and calls[0][1:] == ('SWAP', 'Pending OTC', '26F04329911', 'SWAP')
          and calls[0][0]['LOB'] == 'CEM' and calls[0][0]['TaxID'] == '58160789000128')
    check('B2B: a linha da Intrag Swap (visão da Atacama), no dia do TRADE DATE (§656)', m7 and len(intrag) == 1
          and intrag[0][0]['b3_id'] == '26F04329912' and intrag[0][1] == datetime(2026, 9, 16))
    confs = R._swap_cashflow_engine().confirmation_deals(datetime(2026, 9, 16))
    check('a segregação das confirmações recebe o cashflow contra cliente (o B2B fica fora)',
          [c['B3_ID'] for c in confs] == ['26F04329911'])
    from apps.pages.platform import confirmations as PCF
    check('e o `_conf_load_swap` da platform lê as DUAS páginas',
          '26F04329911' in [c.get('B3_ID') for c in PCF._conf_load_swap(datetime(2026, 9, 16))])

    print('== 7. os vizinhos: Monitor, backfill, e-mail ==')
    from apps.pages.features.deals_monitor import domain as ndm
    card = [c for c in ndm._NDM_CARDS if c['key'] == 'swap-cem'][0]
    check('Monitor: a linha CEM da pasta Swap/Cashflow cai no card Swap CEM, que perdeu o selo soon',
          ndm._ndm_bucket('Swap/Cashflow', {'LOB': 'CEM'}) in ndm.card_buckets(card) and not card.get('soon'))
    check('persistência na pasta que o Monitor varre', 'Swap/Cashflow' in ndm._NDM_SWAP_DIRS
          and persistence.SUFFIX == '_swapcashflow.json')
    bf = io.open(os.path.join(ROOT, 'scripts', 'backfill_manual_confirmations.py'), encoding='utf-8').read()
    check('backfill: famílias SWAP e SWAP CORPORATE para a pasta Swap/Cashflow',
          bf.count("'dir': ('Swap', 'Cashflow')") == 2)
    from apps.pages import otc_emails
    fi = dict(l5[i5], TradeDate=datetime.now().strftime('%Y-%m-%d'))
    drafts = otc_emails.build_swap_bullet_affirmation_emails([fi])
    check('Economic Affirmation (IF): o Deal Ticket do e-mail leva a tabela Cash Flow',
          len(drafts) == 1 and 'Cash Flow' in drafts[0]['html'] and 'Sobre Valor Base Original' in drafts[0]['html']
          and '25%' in drafts[0]['html'])
    apagados, nao = commands.delete([{'deal_id': cli['_id'], 'trade_date': '2026-09-16'}])
    check('delete apaga do arquivo', apagados == 1 and not nao and queries.find(cli['_id'], '2026-09-16')[2] is None)

    print('== 8. as APIs (test client) ==')
    cl = app.test_client()
    with cl.session_transaction() as ss:
        ss['authenticated'] = True
        ss['user_sid'] = 'T000000'
        ss['user_name'] = 'Teste'
        ss['session_expires_at'] = (datetime.now(tz=timezone.utc) + timedelta(hours=8)).isoformat()
    api = pagina['api']
    del SINO[:]
    r = cl.post(api + '/import-file?dry_run=1', data={'file': (io.BytesIO(xlsx), 'dt.xlsx'), 'trade_date': '17/09/2026'},
                content_type='multipart/form-data')
    j = r.get_json()
    check('import dry-run: 200, dois deals, lacunas por id, sem sino', r.status_code == 200 and len(j['deals']) == 2
          and isinstance(j['missing'], dict) and not SINO, (r.status_code, j and j.get('code')))
    r = cl.post(api + '/cache/batch', json={'deals': j['deals']})
    check('batch: 200 e o sino com o rótulo do catálogo', r.status_code == 200 and r.get_json()['imported'] == 2
          and SINO and SINO[-1][3] == 'Swap Cashflow')
    busca = cl.post(api + '/cache/search', json={'filters': [{'field': 'TradeDate', 'type': 'date', 'value': '17/09/2026',
                                                              'mode': 'exact'}]}).get_json()
    check('cache/search: os deals do dia, com o contrato', busca['success'] and len(busca['deals']) == 2
          and busca['fields'] == list(domain.SWC_FIELDS))
    alvo = busca['deals'][0]
    r = cl.get(api + '/preview?deal_id=%s&trade_date=%s' % (alvo['_id'], alvo['TradeDate']))
    check('preview: 200 com os arquivos', r.status_code == 200 and r.get_json()['files'])
    r = cl.post(api + '/edit', json={'deal_id': alvo['_id'], 'trade_date': alvo['TradeDate'], 'changes': {'AmortizationType': ''}})
    check('edit: 200, Pending', r.status_code == 200 and r.get_json()['entry']['Status'] == 'Pending')
    r = cl.get(api + '/preview?deal_id=%s&trade_date=%s' % (alvo['_id'], alvo['TradeDate']))
    jj = r.get_json()
    check('preview com lacuna: 422 com código e as lacunas estruturadas', r.status_code == 422 and jj['code'] == 'swc_lacunas'
          and any(x['code'] == 'swc_no_amortization' for x in jj['lacunas']))
    r = cl.post(api + '/confirm', json={'deal_id': alvo['_id'], 'trade_date': alvo['TradeDate']})
    check('confirm do próprio maker: 403 com código', r.status_code == 403 and r.get_json()['code'] == 'swc_maker_checker')
    r = cl.post(api + '/send-conecta', json={'items': [{'deal_id': alvo['_id'], 'trade_date': alvo['TradeDate']}]})
    check('send de Pending: 400 recusado com a lacuna de status', r.status_code == 400
          and r.get_json()['code'] == 'swc_send_refused' and r.get_json()['lacunas'][0]['code'] == 'swc_status')
    r = cl.post(api + '/add', json={'trade_date': '2026-09-17', 'fields': {'SPN': '281808', 'LOB': 'EDG', 'StartDate': '2026-09-17',
                                                                           'MaturityDate': '2027-09-17', 'Notional': '1000000',
                                                                           'CashFlows': [{'PaymentDate': '2027-09-17', 'AmortizationPct': '100'}]}})
    ja = r.get_json()
    check('add: 200, New, EDG, um fluxo, contraparte pela SPN', r.status_code == 200 and ja['entry']['Status'] == 'New'
          and ja['entry']['LOB'] == 'EDG' and ja['entry']['Flows'] == '1' and ja['entry']['Client'] == 'BANCO SAFRA S.A.')
    r = cl.post(api + '/delete', json={'items': [{'deal_id': ja['entry']['_id'], 'trade_date': '2026-09-17'}]})
    check('delete: 200 e 1 apagado', r.status_code == 200 and r.get_json()['deleted'] == 1)
    r = cl.post(api + '/nada')
    check('ação desconhecida: 404 JSON com código', r.status_code == 404 and r.get_json()['code'] == 'nd_unknown_action')
    r = cl.post(api + '/import-file', data={'file': (io.BytesIO(b'\x00lixo'), 'x.xlsx')}, content_type='multipart/form-data')
    check('arquivo ilegível: 400 com código e o motivo', r.status_code == 400 and r.get_json()['code'] == 'swc_unreadable'
          and r.get_json()['params'].get('reason'))
    check('sem sessão: 401', app.test_client().post(api + '/cache/search').status_code == 401)
    outra = cl.post('/api/new-deals/opt-edg/cache/search', json={'filters': []}).get_json()
    check('o opt-edg segue no stub (vazio com o contrato)', outra.get('success') and outra.get('backend') is False)

    print('== 9. o Internal Trade Recap no CORPO do e-mail (§555): só a Onshore Swap ==')
    import email.message
    def _p(txt):
        return "<p class=MsoNormal>%s</p>" % txt
    def _lbl(rot, val):
        return _p("%s<span style='mso-tab-count:1'>&nbsp;&nbsp;&nbsp; </span>%s" % (rot, val))
    corpo = ''.join([
        "<html><head><style>p.MsoNormal{margin:0}</style></head><body><div class=WordSection1>",
        _p('Hi team,'), _p('We have just executed the below transaction. Booking details below.'),
        "<ul><li>XCCY onshore booked</li><li>XCCY offshore booked</li><li>TRS - could you please proceed with booking?</li></ul>",
        "<table class=MsoNormalTable><tr><td>Gross</td><td>1,195,000.00</td></tr><tr><td>Net AFR</td><td>801,968.00</td></tr></table>",
        _p('<u>Please note that there is a JPM LE to each swap below</u>:'),
        _p("<b><u><span lang=EN-US>Vibra Energia SA SPN : </span></u></b><u>1962701</u>"), _p('&nbsp;'),
        _p('<b><u>I) Onshore Cross-Currency Swap</u></b>'),
        _lbl('Party A', 'Banco J.P.Morgan'), _lbl('Party B', 'Vibra Energia S.A.'), _lbl('Agreement', 'CGD'),
        _lbl('Effective', '11 Sep 2026 (T+5)'), _lbl('Maturity', '11 Sep 2029 (3yrs from Effective)'),
        _lbl('Notional', 'USD 30,000,000'), _lbl('Initial FX', '[5.1150]'), _lbl('Notional', 'BRL [153,450,000]'),
        _lbl('Vibra Recs', 'USD [5.69]% semiannual act/360 (equiv to S+152)'),
        _lbl('Vibra Pays', 'CDI + [0.15]% semiannual bd/252'), _lbl('Amortization', 'At maturity'), _p('&nbsp;'),
        _p('<b><u>Inicio</u></b>&nbsp;&nbsp;&nbsp;&nbsp; <b><u>Datas de Fluxo</u></b>'),
        _p('&nbsp;&nbsp;11-Sep-26&nbsp;&nbsp;&nbsp; 11-Mar-27'), _p('&nbsp;&nbsp;11-Mar-27&nbsp;&nbsp;&nbsp; 13-Sep-27'),
        _p('&nbsp;&nbsp;13-Sep-27&nbsp;&nbsp;&nbsp; 13-Mar-28'), _p('&nbsp;&nbsp;13-Mar-28&nbsp;&nbsp;&nbsp; 11-Sep-28'),
        _p('&nbsp;&nbsp;11-Sep-28&nbsp;&nbsp;&nbsp; 12-Mar-29'), _p('&nbsp;&nbsp;12-Mar-29&nbsp;&nbsp;&nbsp; 11-Sep-29'), _p('&nbsp;'),
        _p('<b><u>II) Brazil Jan-30 LTN TRS</u></b> - <a href="mailto:x">@LATAM NY MO TEAM</a> – could you please proceed with the booking?'),
        _lbl('Party A', 'J.P.Morgan Chase Bank N.A.'), _lbl('Party B', 'Vibra Energia S.A.'), _lbl('Agreement', 'ISDA'),
        _lbl('Effective Date', '11 Sep 2026 (T+5)'), _lbl('Maturity', '01 Jul 2029'), _lbl('Notional', 'USD 30,000,000'),
        _lbl('Bonds', 'Brazil Jul-29 LTN'), _lbl('Qty Bonds', '[220,057]'), _lbl('Bond PU', '[695.274406177]'),
        _lbl('ISIN', 'BRSTNCLTN8K6'), _lbl('Vibra Recs', 'Total Return of Brazil LTN Bonds'),
        _lbl('Vibra Pays', 'USD [5.79]% bullet act/360'), _lbl('Amortization', 'At maturity'), _p('&nbsp;'),
        _p('<b><u>III) Brazil Offshore Cross-Currency Swap</u></b>'),
        _lbl('Party A', 'J.P.Morgan Chase Bank N.A.'), _lbl('Party B', 'Vibra Energia S.A.'), _lbl('Agreement', 'ISDA'),
        _lbl('Effective Date', '11 Sep 2026'), _lbl('Maturity', '01 Jul 2029'), _lbl('Notional', 'USD 30,000,000'),
        _lbl('Initial FX', '[5.1150]'), _lbl('Notional', 'BRL 153,450,000'),
        _lbl('Vibra Recs', 'USD [5.79]% bullet act/360'), _lbl('Vibra Pays', 'BRL [14.00]% bd/252'),
        _lbl('Amortization', 'At maturity'),
        "</div></body></html>"])
    em = email.message.EmailMessage()
    em['Subject'] = 'Internal Trade Recap - $30M 3yrs USD EPP + Convertibility + Swap - 03 Sep 2026'
    em['From'] = 'desk@jpmorgan.com'; em['To'] = 'ops@jpmorgan.com'
    em.set_content('plain'); em.add_alternative(corpo, subtype='html')
    eml = em.as_bytes()
    itens = sw.read_upload('recap.eml', eml)
    check('o corpo vira TRÊS seções (a tabela do AFR e a lista do "Hi team" não são seção)',
          [k for k, _t, _p2 in itens] == ['recap'] * 3
          and [t for _k, t, _p2 in itens] == ['I) Onshore Cross-Currency Swap',
                                              'II) Brazil Jan-30 LTN TRS - @LATAM NY MO TEAM – could you please proceed with the booking?',
                                              'III) Brazil Offshore Cross-Currency Swap'],
          [t for _k, t, _p2 in itens])
    sec_trs = itens[1][2]
    check('as linhas do TRS ficam NA seção dele (não se somam ao swap de cima)',
          itens[0][2]['notionals'] == [('USD', 30000000.0), ('BRL', 153450000.0)] and len(itens[0][2]['legs']) == 2
          and sec_trs['fields'].get('Agreement') == 'ISDA' and sec_trs['fields'].get('Effective') == '11 Sep 2026 (T+5)')
    res9 = commands.import_upload('recap.eml', eml, datetime(2026, 9, 3), sid='A111111', dry_run=True)
    check('só a Onshore Swap é importada; o TRS e a offshore voltam como ignoradas',
          len(res9['deals']) == 1 and len(res9['ignored']) == 2 and 'TRS' in res9['ignored'][0]
          and 'Offshore' in res9['ignored'][1])
    d9 = res9['deals'][0]
    check('contraparte pela SPN do cabeçalho, datas do Effective/Maturity, valor base em BRL',
          d9['SPN'] == '1962701' and d9['StartDate'] == '2026-09-11' and d9['MaturityDate'] == '2029-09-11'
          and d9['Notional'] == '153450000.00' and d9['Currency'] == 'BRL' and d9['Type'] == 'Fluxo de Caixa', d9)
    check('Notional (Foreign Ccy) = USD 30 mi e FX Start = 5.115', d9['NotionalFC'] == '30000000.00' and d9['FXStart'] == '5.115')
    check('Curva A = a perna que o BANCO recebe (a que a Vibra paga): CDI + 0,15%, JUROS, bd/252',
          d9['CurveA'] == 'CDI' and d9['CurveACategory'] == 'JUROS' and d9['CurveASign'] == '+'
          and float(d9['CurveARate']) == 0.15 and float(d9['CurveAPct']) == 100 and d9['CurveADCC'] == 'BD/252', d9)
    check('Curva B = a perna em dólar, 5,69% act/360; sem a curva no cadastro, a categoria fica EM BRANCO',
          d9['CurveB'] == 'USD' and d9['CurveBCategory'] == '' and float(d9['CurveBRate']) == 5.69
          and d9['CurveBDCC'] == 'ACT/360' and not d9['VcpText'], d9)
    fl9 = d9['CashFlows']
    check('o cronograma: seis fluxos, datas do e-mail, 100% de amortização só no último (At maturity)',
          len(fl9) == 6 and fl9[0]['StartDate'] == '2026-09-11' and fl9[0]['PaymentDate'] == '2027-03-11'
          and fl9[-1]['PaymentDate'] == '2029-09-11' and fl9[-1]['AmortizationPct'] == '100'
          and all(f['AmortizationPct'] == '0' for f in fl9[:-1]) and d9['Flows'] == '6', fl9)
    check('Adesão CGD (o Agreement) e LOB em BRANCO (o recap não diz; não se chuta)',
          d9['Adhesion'] == 'CGD' and d9['LOB'] == '')
    lac9 = domain.missing_for_send(d9, sw.codes_for(d9), sw.own_accounts(), '')
    cods9 = {x['code'] for x in lac9}
    check('lacunas estruturadas: LOB e a categoria da Curva B (a do dólar)',
          'swc_no_lob' in cods9 and any(x['code'] == 'swc_no_category' and x['params']['side'] == 'B' for x in lac9)
          and not any(x['code'] == 'swc_no_category' and x['params']['side'] == 'A' for x in lac9), sorted(cods9))
    # Com a curva USD cadastrada (a mesa: USD é TAXAS DE CAMBIO, código 220 no
    # Swap Index), a categoria sai dos cadastros e a perna de moeda vai ao 0301.
    _mr = R._mapping_rows
    def _mr_usd(key, *a, **k):
        rows = list(_mr(key, *a, **k) or [])
        if key == 'swap-bullet-curve':
            rows.append({'DT CURVE': 'USD', 'MATCH': 'Exact', 'B3 CODE': '220', 'NOTES': 'DOLAR DOS EUA'})
        if key == 'swap-index' and not any(str(r.get('Codigo Referencia Externa')) == '220' for r in rows):
            rows.append({'Codigo Referencia Externa': '220', 'Nome Curva': 'DOLAR DOS EUA',
                         'Nome Categoria': 'TAXAS DE CAMBIO', 'STATUS': 'ACTIVE'})
        if key == 'swap-index' and not any(str(r.get('Codigo Referencia Externa')) == 'C03' for r in rows):
            rows.append({'Codigo Referencia Externa': 'C03', 'Nome Curva': 'DI', 'Nome Categoria': 'JUROS',
                         'STATUS': 'ACTIVE'})
        return rows
    R._mapping_rows = _mr_usd
    try:
        d9b = commands.import_upload('recap.eml', eml, datetime(2026, 9, 3), sid='A111111', dry_run=True)['deals'][0]
        check('com USD → 220 cadastrado: Curva B = TAXAS DE CAMBIO (pelo Swap Index), Curva A = JUROS',
              d9b['CurveBCategory'] == 'TAXAS DE CAMBIO' and d9b['CurveACategory'] == 'JUROS', (d9b['CurveACategory'], d9b['CurveBCategory']))
        check('   e o Clean Coupon da perna de moeda nasce do Initial FX (5.115)', d9b['CurveBCleanCoupon'] == '5.115')
        lb = domain.missing_for_send(d9b, commands.queries.codes_for(d9b), sw.own_accounts(), '')
        cb = {x['code'] for x in lb}
        check('   a Data de Cotação da perna de moeda é LACUNA (o recap não a diz; nada de D-1 presumido)',
              'swc_fx_quote' in cb and 'swc_no_category' not in cb and 'swc_fx_coupon' not in cb, sorted(cb))
        d9b['CurveBQuote'] = 'D-1'
        lb2 = {x['code'] for x in domain.missing_for_send(d9b, commands.queries.codes_for(d9b), sw.own_accounts(), '')}
        check('   preenchida (D-1), a lacuna some', 'swc_fx_quote' not in lb2, sorted(lb2))
        acc = {'JPM': '73760009'}
        d9b['ClientAccount'] = '12345678'
        vals = domain.contract_record_values(d9b, 'client', acc, commands.queries.codes_for(d9b), '1234567890')
        check('0301: a Contraparte (Curva B, moeda) leva código 220, taxa 5,69 e sinal +, sem bloco VCP',
              vals['34'] == '220' and vals['37'] == '0056900' and vals['36'] == '00' and vals['44'].strip() == ''
              and vals['46'].strip() == '', (vals['34'], vals['36'], vals['37']))
        check('0301: Cupom Limpo da Contraparte = 5,1150000 (9(08)v9(07)) e Data de Cotação = 01 (D-1)',
              vals['49'] == '000000051150000' and vals['50'] == '01', (vals['49'], vals['50']))
        check('0301: a Parte (Curva A, CDI) fica sem cupom limpo', vals['47'].strip() == '' and vals['48'].strip() == '')
        check('   e Curva A = C03 (DI + 0,15%)', vals['27'] == 'C03' and vals['30'] == '0001500', (vals['27'], vals['30']))
        # Com Atualização de Notional (a termo): Cupom Limpo em branco, mas a
        # Data de Cotação VAI — sem a coluna, o D-n da fixação ao início.
        dt9 = dict(d9b, CurveBQuote='', NotionalIndex='PTAX-V', NotionalObsDate='2026-09-01')
        ct9 = dict(commands.queries.codes_for(d9b), termo_index='15', termo_quote='01')
        vt9 = domain.contract_record_values(dt9, 'client', acc, ct9, '1234567890')
        check('a termo: Cupom Limpo da moeda em branco e a Data de Cotação = o D-n do termo (01)',
              not vt9['49'].strip() and vt9['50'] == '01', (vt9['49'], vt9['50']))
        lt9 = {x['code'] for x in domain.missing_for_send(dt9, ct9, sw.own_accounts(), '')}
        check('   e a Data de Cotação não é lacuna', 'swc_fx_quote' not in lt9, sorted(lt9))
        vt9b = domain.contract_record_values(dict(dt9, CurveBQuote='D-3'), 'client', acc, ct9, '1234567890')
        check('   a coluna Curve B Quote preenchida vence', vt9b['50'] == '03', vt9b['50'])
    finally:
        R._mapping_rows = _mr
    check('fx_quote_code: D-1/1/01 → 01; D0 → 00; D-7 e texto → vazio',
          [domain.fx_quote_code({'CurveBQuote': v}, 'B') for v in ('D-1', '1', '01', 'D0', 'D-7', 'PTAX')]
          == ['01', '01', '01', '00', '', ''])
    r = cl.post(api + '/import-file?dry_run=1', data={'file': (io.BytesIO(eml), 'Internal Trade Recap.msg'),
                                                      'trade_date': '03/09/2026'}, content_type='multipart/form-data')
    j9 = r.get_json()
    check('a API: o e-mail renomeado entra pelo CONTEÚDO e responde o deal da Onshore Swap',
          r.status_code == 200 and len(j9['deals']) == 1 and j9['deals'][0]['SPN'] == '1962701', r.status_code)
    so_trs = eml.replace(b'I) Onshore Cross-Currency Swap', b'I) Onshore Cross-Currency Option')
    r = cl.post(api + '/import-file?dry_run=1', data={'file': (io.BytesIO(so_trs), 'r.eml')}, content_type='multipart/form-data')
    check('recap sem Onshore Swap: 400 com código e as seções ignoradas',
          r.status_code == 400 and r.get_json()['code'] == 'swc_no_deal_ticket' and len(r.get_json()['ignored']) == 3)

    print('\n== 10. o recap da Boticario (§697): seção por letra, CNPJ, FX de booking, Term SOFR como VCP, a termo ==')
    # Seção `B)`, sem a linha de SPN (a contraparte vem pelo CNPJ do Party B),
    # notional "USD Notional: …", o FX de booking que faz o Valor Base, o
    # Initial FX que é uma FIXAÇÃO futura (a Atualização de Notional), a perna
    # Term SOFR com multiplicador e a tabela do cronograma uma célula por linha.
    def _tab(cells):
        return ''.join('<td><p class=MsoNormal>%s</p></td>' % c for c in cells)
    corpo10 = ''.join([
        "<html><body><div class=WordSection1>", _p('Hi team,'),
        _p('<b><u>B) USD/BRL Onshore Cross-currency Swap</u></b>'), _p('&nbsp;'),
        _p('<b>Party A:</b> Banco J.P. Morgan'),
        _p('<b>Party B:</b> Boticario Produtos de Beleza LTDA (CNPJ: 11.137.051/0001-86)'),
        _p('<b>Effective Date:</b> 10-Nov-26'), _p('<b>Maturity Date:</b> 13-Nov-29'),
        _p('<b>USD Notional</b>: <span style="color:red">40,000,000</span>'),
        _p('<b>BRL Notional: 216,000,000 <a href="mailto:x">@Brazil OTC Ops</a></b>'),
        _p('<b>FX: 5.40 ( just to booking)</b>'), _p('<b>Initial FX:</b> PTAX-V 06-Nov-2026'),
        _p('<b>BRL Notional:</b> USD Notional * Initial FX'),
        _p('<b>Boticario Receives:</b> (6M TERM SFR + 0.60%)* 1.17647 (A/360)'),
        _p('<b>Boticario Pays:</b> CDI+0.30% (bd/252)'),
        _p('<b>Interest Payment Frequency:</b> Semiannual'),
        _p('<b>Amortization:</b> At maturity'), _p('<b>Fixings:</b> PTAX-Ask T-1'), _p('&nbsp;'),
        '<table class=MsoNormalTable><tr>', _tab(['<b>Start Date</b>', '<b>End Date</b>', '<b>Amortization</b>']), '</tr>']
        + ['<tr>%s</tr>' % _tab(r) for r in (('10-Nov-26', '10-May-27', '0%'), ('10-May-27', '10-Nov-27', '0%'),
                                             ('10-Nov-27', '10-May-28', '0%'), ('10-May-28', '10-Nov-28', '0%'),
                                             ('10-Nov-28', '10-May-29', '0%'), ('10-May-29', '13-Nov-29', '100%'))]
        + ['</table></div></body></html>'])
    em10 = email.message.EmailMessage()
    em10['Subject'] = 'Internal Trade Recap - Boticario USD/BRL Onshore XCCY'
    em10['From'] = 'desk@jpmorgan.com'; em10['To'] = 'ops@jpmorgan.com'
    em10.set_content('plain'); em10.add_alternative(corpo10, subtype='html')
    eml10 = em10.as_bytes()
    R._refdata_records().append({'SPN': '7777777', 'COUNTERPARTY': 'BOTICARIO PRODUTOS DE BELEZA LTDA',
                                 'TAX ID': '11.137.051/0001-86', 'B3 ACCOUNT': ''})
    _mr10 = R._mapping_rows
    termo_seed = R._MAPPING_DEFS['swap-termo-index']['seed']

    def _mr_10(key, *a, **k):
        rows = list(_mr10(key, *a, **k) or [])
        if key == 'swap-bullet-curve':
            rows = R._swap_bullet_curve_upgrade([r for r in rows if 'TERM' not in str(r.get('DT CURVE')).upper()])
        if key == 'swap-index':
            for code, nome, cat in (('C00', 'VCP', 'VCP'), ('C03', 'DI', 'JUROS')):
                if not any(str(r.get('Codigo Referencia Externa')) == code for r in rows):
                    rows.append({'Codigo Referencia Externa': code, 'Nome Curva': nome, 'Nome Categoria': cat,
                                 'STATUS': 'ACTIVE'})
        if key == 'swap-termo-index':
            rows = [dict(r) for r in termo_seed]
        if key == 'swap-amortizacao' and not any('MATURITY' in str(r).upper() for r in rows):
            rows.append({'LABEL': 'At maturity', 'CODE': '03', 'NOTES': 'Na Data de Vencimento'})
        return rows
    R._mapping_rows = _mr_10
    try:
        sec10 = [p2 for k, _t, p2 in sw.read_upload('recap.eml', eml10) if k == 'recap']
        check('a seção "B)" é lida sem a linha de SPN', len(sec10) == 1 and sec10[0]['title'].startswith('B)'),
              [t for _k, t, _p2 in sw.read_upload('recap.eml', eml10)])
        res10 = commands.import_upload('recap.eml', eml10, datetime(2026, 11, 6), sid='A111111', dry_run=True)
        check('vira UM deal', len(res10['deals']) == 1, res10.get('ignored'))
        d10 = res10['deals'][0]
        check('contraparte pelo CNPJ do Party B (Reference Data)',
              d10['SPN'] == '7777777' and d10['Client'] == 'BOTICARIO PRODUTOS DE BELEZA LTDA', (d10['SPN'], d10['Client']))
        check('Valor Base = FX de booking 5,40 × USD 40 mi; FX Start 5,4; Notional FC 40 mi',
              d10['Notional'] == '216000000.00' and d10['FXStart'] == '5.4' and d10['NotionalFC'] == '40000000.00',
              (d10['Notional'], d10['FXStart'], d10['NotionalFC']))
        check('o Initial FX "PTAX-V 06-Nov-2026" é a Atualização de Notional (índice e data)',
              d10['NotionalIndex'] == 'PTAX-V' and d10['NotionalObsDate'] == '2026-11-06',
              (d10['NotionalIndex'], d10['NotionalObsDate']))
        check('Curva A = CDI + 0,30% (o banco recebe); Curva B = 6M TERM SFR → VCP 17922, D-1 (Fixings T-1)',
              d10['CurveA'] == 'CDI' and d10['CurveACategory'] == 'JUROS' and d10['CurveB'] == '6M TERM SFR'
              and d10['CurveBCategory'] == 'VCP' and d10['VcpCode'] == '17922' and d10['QuoteDateCode'] == '01'
              and d10['VcpText'] == '(6M TERM SFR + 0.60%)* 1.17647 (A/360)',
              {k: d10.get(k) for k in ('CurveA', 'CurveACategory', 'CurveB', 'CurveBCategory', 'VcpCode', 'VcpText')})
        f10 = d10['CashFlows']
        check('o cronograma da tabela uma célula por linha: 6 fluxos, 100% no último',
              len(f10) == 6 and f10[0]['StartDate'] == '2026-11-10' and f10[-1]['PaymentDate'] == '2029-11-13'
              and f10[-1]['AmortizationPct'] == '100' and all(x['AmortizationPct'] == '0' for x in f10[:-1]), f10)
        d10['LOB'] = 'CEM'
        cod10 = commands.queries.codes_for(d10)
        check('Índice Termo 15 (DOLAR, pelo cadastro) e D-2 (06/11 → 10/11, dias úteis ANBIMA)',
              (cod10.get('termo_index'), cod10.get('termo_quote')) == ('15', '02'), cod10)
        lac10 = {x['code'] for x in domain.missing_for_send(d10, cod10, sw.own_accounts(), cod10.get('amortization'))}
        check('a termo o Preço Inicial do VCP não é pedido; sem lacuna', not lac10, sorted(lac10))
        v10 = domain.contract_record_values(d10, 'client', {'JPM': '73760009'}, cod10, '1234567890')
        check('0301: 14 = 216.000.000,00; 19 = Data da operação; 20 = 15; 21 = 100%; 22 = o FX inicial 5,4; '
              '23-24 em branco; 99 = 02',
              v10['14'] == '0000021600000000' and v10['19'] == '20261106' and v10['20'] == '15'
              and v10['21'] == '10000' and v10['22'] == '0000000000000540000000'
              and not (v10['23'] + v10['24']).strip() and v10['99'] == '02',
              {k: v10[k] for k in ('14', '19', '20', '21', '22', '99')})
        check('0301: a curva VCP do cliente leva Sinal e Juros (36/37)',
              v10['36'] == '00' and v10['37'].isdigit() and len(v10['37']) == 7, (v10['36'], v10['37']))
        v10q = domain.contract_record_values(dict(d10, QuoteDateCode='', CurveBQuote=''), 'client',
                                             {'JPM': '73760009'}, cod10, '1234567890')
        check('0301: sem Quote Date D-n, a Data de Cotação da VCP a termo = o D-n do termo (02)',
              v10q['50'] == '02' and not v10q['49'].strip(), (v10q['49'], v10q['50']))
        l10f = {x['code'] for x in domain.missing_for_send(dict(d10, FXStart=''), cod10, sw.own_accounts(),
                                                            cod10.get('amortization'))}
        check('   a termo sem FX Start é lacuna (o PU Inicial)', 'swc_fx_coupon' in l10f, sorted(l10f))
        check('0301: com FWD Start e Atualização de Notional, o Cupom Limpo (47/49) vai em BRANCO',
              not v10['47'].strip() and not v10['49'].strip())
        check('0301: a perna VCP do cliente (guarda-chuva) leva PU 1, 17922, a Denominação e D-1',
              v10['34'] == 'C00' and v10['44'] == '0000000000000100000000' and v10['45'] == '17922'
              and v10['46'].strip() == '(6M TERM SFR + 0.60%)* 1.17647 (A/360)' and v10['50'] == '01'
              and not v10['43'].strip(), {k: v10[k].strip() for k in ('34', '44', '45', '46', '50')})
        # A fixação no PRÓPRIO dia do início: o índice vira VCP (00), com PU,
        # Tipo/Classe e Denominação, e sem o D-n (a regra da macro da mesa).
        d10b = dict(d10, NotionalObsDate='2026-11-10')
        cb = domain.termo_codes(d10b, termo_seed, lambda a, b: 0)
        check('fixação no dia do início → Índice Termo 00 (VCP); sem Tipo/Classe é lacuna',
              cb['termo_index'] == '00' and 'swc_termo_class' in {
                  x['code'] for x in domain.missing_for_send(d10b, dict(cod10, **cb), sw.own_accounts(), '03')}, cb)
        vb = domain.contract_record_values(d10b, 'client', {'JPM': '73760009'},
                                           dict(cod10, termo_index='00', termo_class='12345', termo_quote=''), '1')
        check('   e o 0301 leva PU = FX 5,4, Tipo/Classe e a Denominação, sem o 99',
              vb['20'] == '00' and vb['22'] == '0000000000000540000000' and vb['23'] == '12345'
              and vb['24'].startswith('Valor base convertido pelo FX Inicial: 5.4000000 e sera atualizado em '
                                      '10/11/2026 pelo indice PTAX-V') and not vb['99'].strip(), (vb['22'], vb['24'][:90]))
        sem = dict(d10, NotionalIndex='', NotionalObsDate='')
        vs = domain.contract_record_values(sem, 'client', {'JPM': '73760009'}, commands.queries.codes_for(sem), '1')
        check('sem Atualização de Notional, 19-24 e 99 seguem em branco (o swap em D0)',
              not (vs['19'] + vs['20'] + vs['21'] + vs['99']).strip())
        check('o upgrade do cadastro de curvas acrescenta o Term SOFR só onde falta',
              len(R._swap_bullet_curve_upgrade([{'DT CURVE': '6M TERM SFR', 'B3 CODE': 'X'}]))
              == len(R._SWAP_TERM_SOFR_CURVES))
    finally:
        R._mapping_rows = _mr10

    print('== 11. o B2B × Lawton (CEM) e a guarda-chuva CLIENT 1 (§698) ==')
    _f, la, ia = queries.find(b2b['_id'], '2026-09-16')
    law = dtk.set_pair(dict(la[ia], LE='LAWTON', Client='Lawton'))
    check('set_pair: LE e par do Lawton', (law['LE'], law['Pair']) == ('LAWTON', 'JPM x LAWTON'))
    check('o fundo, as visões e os pares', dtk.b2b_fund(law) == 'LAWTON' and dtk.views_of(law) == ['bank', 'lawton']
          and dtk.le_pair('bank', law) == 'JPM x LAWTON' and dtk.le_pair('lawton', law) == 'LAWTON x JPM'
          and dtk.le_pair('bank') == 'JPM x ATACAMA')
    check('a Atacama segue igual', dtk.views_of(la[ia]) == ['bank', 'atacama'] and dtk.b2b_fund(la[ia]) == 'ATACAMA')
    check('pelo Client, sem LE de fundo', dtk.set_pair({'LE': 'JPM', 'Client': 'LAWTON FIM'})['LE'] == 'LAWTON'
          and dtk.set_pair({'LE': '', 'Client': 'BANCO SAFRA'})['Pair'] == 'JPM x CLI')
    contas = sw.own_accounts()
    check('a conta própria do Lawton vem do b3-accounts', contas.get('LAWTON') == '00041007', contas.get('LAWTON'))
    fl = commands.preview(law)
    check('preview: contrato + cronograma + prêmio × Banco e Lawton',
          [(f['kind'], f['view']) for f in fl] == [('swap', 'bank'), ('flow', 'bank'), ('premium', 'bank'),
                                                   ('swap', 'lawton'), ('flow', 'lawton'), ('premium', 'lawton')],
          [(f['kind'], f['view']) for f in fl])
    nomes = [f['file_name'] for f in fl]
    check('arquivos do Lawton com a conta dele na frente',
          '00041007.SWAP_CF_CEM_LAWTON.txt' in nomes and '73760009.SWAP_CF_CEM_BANCO.txt' in nomes, nomes)
    rb, rl = fl[0]['records'][0], fl[3]['records'][0]
    check('0301: Banco × Lawton na visão do Banco, Lawton × Banco na dele',
          (rb[20:28], rb[52:60], rl[20:28], rl[52:60]) == ('73760009', '00041007', '00041007', '73760009'),
          (rb[20:28], rb[52:60], rl[20:28], rl[52:60]))
    check('o espelho leva o Meu Número do espelho', rl[10:20] == law['MyNumberMirror'])
    check('header do Lawton com o participante dele', fl[3]['header'][10:30].strip() == sw.participant('LAWTON'),
          fl[3]['header'])
    ent = dtk.intrag_swap_entry(dict(law, B3ID='26X1'), sw.codes_for(law))
    ent_a = dtk.intrag_swap_entry(dict(la[ia], B3ID='26X2'), sw.codes_for(la[ia]))
    check('Intrag Swap na carteira do fundo do deal', (ent['carteira'], ent['parte'], ent_a['carteira'], ent_a['parte'])
          == ('INTRAGJP552', 'Lawton', 'INTRAGJP633', 'Atacama'), (ent['carteira'], ent['parte']))
    check('sem a conta do Lawton, lacuna dele', any('LAWTON own account' in str(x) for x in
          dtk.missing_for_send(law, sw.codes_for(law), {'JPM': '73760009'})))
    R._refdata_records().append({'SPN': '3871464', 'COUNTERPARTY': 'BOTICARIO PRODUTOS DE BELEZA LTDA',
                'TAX ID': '11.137.051/0001-86', 'B3 ACCOUNT': ''})
    bot = sw.enrich({'SPN': '3871464', 'LE': 'JPM', 'Client': 'Boticario'})
    check('cliente sem conta própria: a guarda-chuva CLIENT 1 (73760102) + o CNPJ',
          (bot['ClientAccount'], bot['ClientAccountNote'], bot['ClientTaxId'])
          == ('73760102', 'omnibus', '11137051000186'), (bot.get('ClientAccount'), bot.get('ClientTaxId')))
    velho = dict(bot, ClientAccount='73760205')
    check('o deal com a guarda-chuva antiga gravada sai com a de hoje (preview/Send)',
          sw.with_current_omnibus(velho)['ClientAccount'] == '73760102' and velho['ClientAccount'] == '73760205')
    check('   e a conta própria do cliente não é tocada',
          sw.with_current_omnibus({'ClientAccount': '74220005'})['ClientAccount'] == '74220005')

    print('== 12. a tela: Add Deal, a curva com o spread, o Type traduzido e o Create B2B ==')
    import shutil
    import subprocess
    js = os.path.join(ROOT, 'apps', 'static', 'js', 'swap-deal.js')
    node = shutil.which('node')
    if not node:
        check('node disponível', False)
    else:
        prog = r'''
global.window = global; require(process.argv[1]);
var cli = {LE:'JPM', LOB:'CEM', Client:'BOTICARIO', SPN:'3871464', Deal:'D1', B3ID:'26X', _id:'SWC-1',
  ClientAccount:'73760102', CurveA:'CDI', CurveACategory:'JUROS', CurveAPct:'100.00', CurveARate:'3.0000', CurveASign:'+',
  CurveB:'6M TERM SFR', CurveBCategory:'VCP', CurveBPct:'110.00', CurveBRate:'0.5000', CurveBSign:'-', CurveBQuote:'01',
  PremiumPayer:'BOTICARIO', CashFlows:[{PaymentDate:'2027-01-01'}], Notional:'216000000.00'};
var b = otcSwapB2B(cli);
console.log(JSON.stringify({
  a: otcSwapCurveText(cli,'A'), b: otcSwapCurveText(cli,'B'),
  zero: otcSwapCurveText({CurveA:'CDI', CurveARate:'0.0000', CurveAPct:'100'},'A'),
  fino: otcSwapCurveText({CurveA:'IPCA', CurveARate:'5.1234'},'A'),
  tipo: [otcSwapTypeLabel('Fluxo de Caixa'), otcSwapTypeLabel('Pagamento Final'), otcSwapTypeLabel('X')],
  b2b: b, edg: otcSwapB2B({LOB:'EDG', PremiumPayer:'Banco JP Morgan'}), semlob: otcSwapB2B({LOB:''}),
  isb2b: [otcSwapIsB2B(b), otcSwapIsB2B(cli)]}));
'''
        out = subprocess.run([node, '-e', prog, js], capture_output=True, text=True, timeout=30)
        try:
            r = json.loads(out.stdout.strip().splitlines()[-1])
        except Exception:                                            # noqa: BLE001
            r = {}
        check('curva com spread: "CDI + 3.00%", e o percentual fora dos 100%',
              (r.get('a'), r.get('b')) == ('CDI + 3.00%', '110.00% 6M TERM SFR - 0.50%'), (r, out.stderr[-300:]))
        check('   sem spread só o nome; quatro casas quando há', (r.get('zero'), r.get('fino')) == ('CDI', 'IPCA + 5.1234%'),
              (r.get('zero'), r.get('fino')))
        check('Type no idioma da tela (fallback em inglês)', r.get('tipo') == ['Cash Flow', 'Bullet', 'X'], r.get('tipo'))
        b2 = r.get('b2b') or {}
        check('B2B da CEM: × Lawton, pernas trocadas campo a campo',
              (b2.get('LE'), b2.get('Client'), b2.get('CurveA'), b2.get('CurveB'), b2.get('CurveARate'), b2.get('CurveBRate'),
               b2.get('CurveASign'), b2.get('CurveBPct'), b2.get('CurveAQuote'), b2.get('CurveBQuote'))
              == ('LAWTON', 'Lawton', '6M TERM SFR', 'CDI', '0.5000', '3.0000', '-', '100.00', '01', ''), b2)
        check('   sem o que identifica a operação do cliente, com o cronograma',
              not any(b2.get(k) for k in ('Deal', 'B3ID', '_id', 'SPN', 'ClientAccount'))
              and b2.get('CashFlows') == [{'PaymentDate': '2027-01-01'}] and b2.get('Notional') == '216000000.00')
        check('   o prêmio inverte: o cliente pagava → o Banco paga ao fundo; o Banco pagava → o fundo paga',
              b2.get('PremiumPayer') == 'Banco' and (r.get('edg') or {}).get('PremiumPayer') == 'Atacama'
              and (r.get('edg') or {}).get('LE') == 'ATACAMA')
        check('   sem LOB não há B2B (não se chuta); o B2B não oferece outro',
              r.get('semlob') is None and r.get('isb2b') == [True, False])
    pag = catalog.PAGES['/new_deals-swap-cashflow']
    check('catálogo: Add Deal nos swaps, Create B2B só no Cashflow, a curva como "curve"',
          pag['add_label'] == 'Add Deal' and pag['b2b'] is True
          and catalog.PAGES['/new_deals-swap-cetes']['add_label'] == 'Add Deal'
          and catalog.PAGES['/new_deals-swap-cetes']['b2b'] is False
          and catalog.PAGES['/new_deals-opt-edg']['add_label'] == 'Add Row'
          and ['CurveA', 'Curve A', 'curve'] in pag['cards']['facts'] and 'LAWTON' in pag['selects']['LE']
          and 'LAWTON' not in catalog.PAGES['/new_deals-opt-edg']['selects']['LE'])
    swb = io.open(os.path.join(ROOT, 'apps', 'templates', 'pages', 'new_deals-swap-bullet.html'), encoding='utf-8').read()
    cards = io.open(os.path.join(ROOT, 'apps', 'static', 'js', 'deal-cards.js'), encoding='utf-8').read()
    for nome, h in (('catálogo', tpl), ('Bullet', swb)):
        check('%s: carrega o swap-deal.js e abre o modal de Add com o B2B' % nome,
              "js/swap-deal.js?v={{ asset_v('js/swap-deal.js') }}" in h and 'openModal(null, d)' in h
              and 'onB2B' in h and "t('nd-add-deal', 'Add Deal')" in h)
    check('deal-cards: o quinto botão (Create B2B) só com onB2B', 'btn-row-b2b' in cards and 'o.onB2B &&' in cards)
    for lg in ('en', 'br', 'es'):
        trj = json.load(io.open(os.path.join(ROOT, 'apps', 'static', 'data', 'translations', lg + '.json'), encoding='utf-8'))
        faltam = [k for k in ('nd-add-deal', 'nd-add-deal-tip', 'nd-card-b2b', 'w_nd_b2b_no_lob',
                              'nd-swap-type-cashflow', 'nd-swap-type-bullet') if not trj.get(k)]
        check('tradução %s das chaves novas' % lg, not faltam, faltam)


if __name__ == '__main__':
    sys.exit(main())
