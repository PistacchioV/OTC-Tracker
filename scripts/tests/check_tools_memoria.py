# -*- coding: utf-8 -*-
"""check_tools_memoria.py — Swap Calculator > Export: a memoria de calculo em
.xlsx (HANDOFF §455).

O arquivo nao e um retrato do resultado: e a CONTA. Cada celula derivada e uma
formula de Excel encadeada ate as entradas — o fixing do DI de cada dia, a taxa
contratada, as datas e o notional —, e e isso que faz a mesa poder mandar o
arquivo para o cliente que contesta. Um export que so copia numeros passaria
neste teste por engano, entao o que se prende aqui e o RECALCULO:

  1. um avaliador minimo le o workbook, resolve as referencias (inclusive as
     que atravessam abas) e CALCULA as formulas como o Excel faria. O ajuste
     liquido assim obtido tem de bater com o do motor. Se alguem trocar uma
     formula por um valor pronto — ou escrever uma formula errada —, o numero
     deixa de fechar e este teste cai;
  2. os dias do indice vao inteiros numa aba propria, com o fator de CADA dia
     escrito como conta sobre a taxa daquele dia e o acumulado como produto do
     anterior. A aba principal REFERENCIA a celula do acumulado: mexer num dia
     refaz a liquidacao;
  3. o fator diario do DI respeita o arredondamento na 8a casa quando a tela
     pede o padrao B3/CETIP — e e o mesmo `ROUND(...,8)` que o motor aplica;
  4. o documento e do BANCO: o timbre entra no arquivo e nao ha uma linha
     sobre o sistema que o gerou, nem no conteudo nem nas propriedades;
  5. o nome do arquivo e `Memoria de Calculo - CETIP ID - contraparte - data`,
     com o que falta SUMINDO em vez de virar um traco solto, e sem os
     caracteres que o Windows recusa;
  6. a rota e POST (o que se exporta e o formulario), exige sessao, e conta que
     nao fecha volta para a TELA com a mensagem em vez de baixar um arquivo
     quebrado — em JSON quando quem pede e o fetch do botao, que e quem
     desliga o spinner. Por isso o `<form>` declara `action` explicito;
  7. memoria de LIQUIDACAO nao tem valor FUTURO: cada ponta fecha em UMA
     linha, a do que ela liquida naquele fluxo — juros no intermediario,
     valor da ponta no vencimento. Mostrar os dois lados transformaria a
     memoria de uma liquidacao num comparativo de bases, com metade dos
     numeros descrevendo o que nao aconteceu.

Nada sai da maquina: a serie do CDI e um stub e os fixings de moeda sao
digitados.
"""
import datetime as _dt
import io
import os
import re
import sys
import zipfile
from datetime import date, datetime

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
sys.path.insert(0, ROOT)
os.chdir(ROOT)
os.environ.setdefault('OTC_SHARED_DRIVE_ROOT', ROOT)
os.environ['OTC_DISABLE_SCHEDULERS'] = '1'

falhas = []

# O app sobe ANTES de tudo, e nao no fim: importar um modulo de feature primeiro
# faz o `create_app` seguinte encontrar o blueprint ja montado e recusar o
# registro ("overwriting an existing endpoint"). E tambem a ordem real da
# subida — o `routes.py` e quem importa as verticais.
from apps import create_app                                             # noqa: E402
from apps.config import DebugConfig                                     # noqa: E402

app = create_app(DebugConfig)


def check(rotulo, obtido, esperado):
    ok = obtido == esperado
    print(('  ok  ' if ok else ' FAIL ') + rotulo +
          ('' if ok else '\n        got=%r\n        exp=%r' % (obtido, esperado)))
    if not ok:
        falhas.append(rotulo)


def perto(rotulo, obtido, esperado, tol=1e-6):
    ok = obtido is not None and abs(obtido - esperado) <= tol
    print(('  ok  ' if ok else ' FAIL ') + rotulo +
          ('' if ok else '\n        got=%r\n        exp=%r' % (obtido, esperado)))
    if not ok:
        falhas.append(rotulo)


# ─────────────────────────────────────────────────────────────────────────────
# O avaliador: o Excel que este teste precisa, e nada mais.
#
# Uma dependencia nova so para conferir quatro operadores nao se paga, e uma
# biblioteca de calculo tambem teria de ser confiada. O que as formulas deste
# arquivo usam esta todo aqui: as quatro operacoes, a potencia, IF/AND/MIN/
# ABS/ROUND, comparacoes e referencias (com ou sem aba). Data vira SERIAL, como
# no Excel — e por isso que `fim - inicio` da os dias corridos.
# ─────────────────────────────────────────────────────────────────────────────
REF = re.compile(r"(?:'([^']+)'!)?(\$?[A-Z]{1,2}\$?[0-9]{1,6})")
EPOCA = datetime(1899, 12, 30)


def _serial(v):
    if isinstance(v, datetime):
        return (v - EPOCA).days
    if isinstance(v, date):
        return (datetime(v.year, v.month, v.day) - EPOCA).days
    return v


def celula(wb, aba, endereco, pilha=()):
    chave = (aba, endereco.replace('$', ''))
    if chave in pilha:
        raise RuntimeError('referencia circular em %r' % (chave,))
    v = wb[aba][chave[1]].value
    if isinstance(v, str) and v.startswith('='):
        return avaliar(wb, aba, v[1:], pilha + (chave,))
    return _serial(v)


def avaliar(wb, aba, expressao, pilha=()):
    def troca(m):
        folha = m.group(1) or aba
        return repr(celula(wb, folha, m.group(2), pilha))

    py = REF.sub(troca, expressao).replace('^', '**')
    return eval(py, {'__builtins__': {}},                              # noqa: S307
                {'IF': lambda c, a, b: a if c else b, 'AND': lambda *a: all(a),
                 'MIN': min, 'ABS': abs, 'ROUND': round, 'TRUE': True, 'FALSE': False})


def por_rotulo(ws, rotulo):
    """O endereco do VALOR (coluna B) da linha cujo rotulo e este."""
    for ln in range(1, ws.max_row + 1):
        if ws.cell(row=ln, column=1).value == rotulo:
            return 'B%d' % ln
    raise AssertionError('rotulo ausente na memoria: %r' % rotulo)


# ─────────────────────────────────────────────────────────────────────────────
print('\n== 1. o nome do arquivo ==')
from apps.pages.features.tools import domain                            # noqa: E402

check('nome completo, com a data da liquidacao em dd-mm-aaaa',
      domain.nome_memoria('26G53382860', 'FUNDO ABC', date(2026, 6, 30)),
      'Memória de Cálculo - 26G53382860 - FUNDO ABC - 30-06-2026.xlsx')
# O `/` da razao social e da data e ilegal no Windows, e um `\n` colado de uma
# celula partiria o cabecalho HTTP do download.
check('o que o Windows recusa vira espaco',
      domain.nome_memoria('A/B', 'X:Y*Z?\nW', date(2026, 1, 2)),
      'Memória de Cálculo - A B - X Y Z W - 02-01-2026.xlsx')
check('segmento vazio SOME, em vez de virar um traco solto',
      domain.nome_memoria('', '   ', date(2026, 1, 2)),
      'Memória de Cálculo - 02-01-2026.xlsx')

# ─────────────────────────────────────────────────────────────────────────────
print('\n== 2. a conta: o Excel recalcula o que o motor calculou ==')
from apps.pages.features.tools import queries                           # noqa: E402
from apps.pages.precificador import cdi, contagem, liquidacao           # noqa: E402

# A serie do BCB e um STUB: nada sai da maquina, e o numero fica previsivel.
_serie_real = cdi.serie
cdi.serie = lambda d0, d1, *a, **k: [
    cdi.FixingCDI(d, 0.1065 + (i % 3) * 0.0005)
    for i, d in enumerate(_dias_uteis_do_periodo(d0, d1))]


def _dias_uteis_do_periodo(d0, d1):
    from apps.pages.precificador.calendario import calendario_anbima
    cal, saida, d = calendario_anbima(), [], d0
    while d <= d1:
        if cal.eh_dia_util(d):
            saida.append(d)
        d += _dt.timedelta(days=1)
    return saida


FORM = {
    'b3_id': '26G53382860', 'counterparty': 'FUNDO DE INVESTIMENTO ABC',
    'data_operacao': '2025-09-01', 'inicio': '2025-09-01', 'fim': '2026-03-02',
    'vencimento': '2027-09-01', 'base_ajuste': liquidacao.BASE_AUTOMATICA,
    'nocional': '10,000,000.00', 'nocional_original': '12,000,000.00',
    'amortizacao': '10', 'base_amortizacao': liquidacao.SOBRE_ORIGINAL,
    'calendario': 'ANBIMA', 'reter_ir': '1', 'arredondar_di': '1',
    'ativa_indexador': liquidacao.PRE, 'ativa_taxa': '14',
    'ativa_convencao': contagem.DU_252, 'ativa_regime': contagem.COMPOSTO,
    'passiva_indexador': liquidacao.CDI, 'passiva_percentual': '104',
    'passiva_taxa': '0.5', 'passiva_convencao': contagem.DU_252,
    'passiva_regime': contagem.COMPOSTO,
}

conteudo, nome = queries.memoria_de_calculo(FORM)
r = queries.liquidar(FORM)['r']
check('o nome sai do CETIP ID, da contraparte e da liquidacao', nome,
      'Memória de Cálculo - 26G53382860 - FUNDO DE INVESTIMENTO ABC - 02-03-2026.xlsx')

import openpyxl                                                         # noqa: E402
wb = openpyxl.load_workbook(io.BytesIO(conteudo))
from apps.pages.features.tools.infra import memoria_xlsx as mx          # noqa: E402

check('as abas: a memoria e a apuracao diaria da ponta com indice',
      wb.sheetnames, [mx.ABA, mx.ABA_DIARIA[liquidacao.PASSIVA]])
ws = wb[mx.ABA]
diaria = wb[mx.ABA_DIARIA[liquidacao.PASSIVA]]

# ── o coracao: o ajuste liquido do Excel tem de ser o do motor ──────────────
liquido = avaliar(wb, mx.ABA, por_rotulo(ws, 'Ajuste líquido'))
perto('o ajuste LIQUIDO recalculado bate com o motor', liquido, r.ajuste_liquido, 1e-6)
perto('   e o bruto tambem', avaliar(wb, mx.ABA, por_rotulo(ws, 'Ajuste bruto')),
      r.ajuste_bruto, 1e-6)
perto('   e o IR retido', avaliar(wb, mx.ABA, por_rotulo(ws, 'IR retido')), r.ir, 1e-6)
perto('   com a aliquota da tabela regressiva num SE aninhado',
      avaliar(wb, mx.ABA, por_rotulo(ws, 'Alíquota de IR')), r.aliquota_ir, 1e-12)
check('o prazo do IR e contado das DATAS, nao copiado',
      avaliar(wb, mx.ABA, por_rotulo(ws, 'Prazo desde a data da operação (dias)')),
      r.dias_da_operacao)
check('e os dias corridos do fluxo tambem',
      avaliar(wb, mx.ABA, por_rotulo(ws, 'Dias corridos do fluxo')), r.dias_corridos)

# ── os fatores de cada ponta, um a um ──────────────────────────────────────
perto('o fator do indice da ponta PRE sai da taxa e do tau',
      avaliar(wb, mx.ABA, por_rotulo(ws, 'Fator do índice')), r.ativa.fator_do_indice, 1e-12)
perto('o tau e dias / base, calculado na planilha',
      avaliar(wb, mx.ABA, por_rotulo(ws, 'τ — fração de ano')), r.ativa.fracao_de_ano, 1e-12)
# A segunda ocorrencia de cada rotulo e a ponta passiva.
linhas_fator = [ln for ln in range(1, ws.max_row + 1)
                if ws.cell(row=ln, column=1).value == 'Fator acumulado da ponta']
check('as duas pontas tem o fator acumulado', len(linhas_fator), 2)
perto('fator acumulado da ponta ATIVA', avaliar(wb, mx.ABA, 'B%d' % linhas_fator[0]),
      r.ativa.fator, 1e-12)
perto('fator acumulado da ponta PASSIVA', avaliar(wb, mx.ABA, 'B%d' % linhas_fator[1]),
      r.passiva.fator, 1e-12)
linhas_juros = [ln for ln in range(1, ws.max_row + 1)
                if ws.cell(row=ln, column=1).value == 'Juros do período (R$)']
perto('juros da ponta ativa', avaliar(wb, mx.ABA, 'B%d' % linhas_juros[0]),
      r.ativa.juros, 1e-6)
perto('juros da ponta passiva', avaliar(wb, mx.ABA, 'B%d' % linhas_juros[1]),
      r.passiva.juros, 1e-6)

# Este fluxo termina ANTES do vencimento: liquida o diferencial de juros, e o
# principal segue para o periodo seguinte. O valor futuro das pontas nao
# aconteceu — nao e "outra forma de ver", e uma projecao, e memoria de
# liquidacao nao projeta.
check('fluxo intermediario: a memoria nao fala em valor futuro', r.so_juros, True)
rotulos = [ws.cell(row=ln, column=1).value for ln in range(1, ws.max_row + 1)]
check('   nenhuma linha de valor futuro',
      [x for x in rotulos if x and 'futuro' in str(x).lower()], [])
check('   e cada ponta fecha em UMA linha, a do que ela liquida',
      (len(linhas_juros), 'Valor da ponta na liquidação (R$)' in rotulos), (2, False))
# Um diferencial so: as duas linhas de ponta vao direto ao Ajuste bruto. A
# linha 'Diferencial de ...' era a segunda base aparecendo de novo.
check('   com um diferencial so — o ajuste bruto',
      [x for x in rotulos if x and str(x).startswith('Diferencial')], [])
perto('o valor amortizado sai do MIN, nao copiado',
      avaliar(wb, mx.ABA, por_rotulo(ws, 'Valor amortizado')), r.valor_amortizado, 1e-6)
perto('e o saldo do fluxo seguinte',
      avaliar(wb, mx.ABA, por_rotulo(ws, 'Saldo do fluxo seguinte')), r.saldo_seguinte, 1e-6)

# ── os dias do DI, evidenciados ────────────────────────────────────────────
print('\n== 3. os dias do indice, dia a dia ==')
check('um dia util publicado por linha', diaria.max_row - 6, len(r.passiva.fixings))
check('a taxa do primeiro dia e a do fixing', diaria['B5'].value, r.passiva.fixings[0].taxa)
check('o fator do dia e FORMULA, nao numero',
      str(diaria['C5'].value).startswith('=') and str(diaria['D6'].value).startswith('='), True)
check('com o arredondamento na 8a casa do padrao B3/CETIP',
      'ROUND(' in diaria['C5'].value, True)
perto('o fator acumulado do ultimo dia e o do motor',
      avaliar(wb, mx.ABA_DIARIA[liquidacao.PASSIVA],
              'D%d' % (4 + len(r.passiva.fixings))),
      r.passiva.fixings[-1].fator_acumulado, 1e-12)
# A aba principal REFERENCIA a celula da diaria: mexer num dia refaz a conta.
alvo = ws[por_rotulo(ws, 'Fator acumulado do CDI')].value
check('a memoria referencia a aba diaria, nao o valor copiado',
      isinstance(alvo, str) and mx.ABA_DIARIA[liquidacao.PASSIVA] in alvo, True)
check('e o percentual do CDI entra na conta do dia pela celula do topo',
      diaria['B2'].value, 1.04)
# Repetido em dois lugares, alguem corrige um e a planilha passa a exibir um
# percentual que nao e o que multiplicou os fatores diarios.
check('   e a memoria REFERENCIA essa celula, em vez de repetir o numero',
      str(ws[por_rotulo(ws, 'Percentual do CDI')].value).startswith('='), True)
perto('   recalculando no percentual contratado',
      avaliar(wb, mx.ABA, por_rotulo(ws, 'Percentual do CDI')), 1.04, 1e-12)

# Sem arredondamento a formula do dia e a outra — e o numero muda.
cru = dict(FORM); cru.pop('arredondar_di')
conteudo_cru, _ = queries.memoria_de_calculo(cru)
wb2 = openpyxl.load_workbook(io.BytesIO(conteudo_cru))
check('desligado o arredondamento, o ROUND some da formula do dia',
      'ROUND(' in wb2[mx.ABA_DIARIA[liquidacao.PASSIVA]]['C5'].value, False)
r2 = queries.liquidar(cru)['r']
perto('e o liquido recalculado segue batendo com o motor',
      avaliar(wb2, mx.ABA, por_rotulo(wb2[mx.ABA], 'Ajuste líquido')),
      r2.ajuste_liquido, 1e-6)

# ─────────────────────────────────────────────────────────────────────────────
print('\n== 4. a perna com moeda: o fator cambial e uma DIVISAO ==')
cambial = dict(FORM)
cambial.update({'ativa_indexador': liquidacao.CAMBIO, 'ativa_taxa': '3.25',
                'ativa_convencao': contagem.ACT_360, 'ativa_regime': contagem.SIMPLES,
                'ativa_moeda': 'USD', 'ativa_ptax_inicial': '5.1234',
                'ativa_ptax_final': '5.4321'})
conteudo_fx, _ = queries.memoria_de_calculo(cambial)
wb3 = openpyxl.load_workbook(io.BytesIO(conteudo_fx))
ws3 = wb3[mx.ABA]
r3 = queries.liquidar(cambial)['r']
perto('fixing final / fixing inicial', avaliar(wb3, mx.ABA, por_rotulo(ws3, 'Fator cambial')),
      r3.ativa.fator_cambial, 1e-12)
perto('e o liquido da perna cambial bate',
      avaliar(wb3, mx.ABA, por_rotulo(ws3, 'Ajuste líquido')), r3.ajuste_liquido, 1e-6)
# Os juros NAO carregam o principal corrigido pela moeda: e o que separa o
# fluxo intermediario do final.
perto('os juros da perna cambial excluem o efeito no principal',
      avaliar(wb3, mx.ABA, 'B%d' % [ln for ln in range(1, ws3.max_row + 1)
                                    if ws3.cell(row=ln, column=1).value
                                    == 'Juros do período (R$)'][0]),
      r3.ativa.juros, 1e-6)

# ─────────────────────────────────────────────────────────────────────────────
print('\n== 4b. SOFR: o n sai das DATAS, e a sexta remunera tres dias ==')
from apps.pages.precificador import sofr                                # noqa: E402
from apps.pages.precificador.calendario import calendario_sofr          # noqa: E402

_serie_sofr_real = sofr.serie_sofr
_cal_us = calendario_sofr()


def _uteis_us(d0, d1):
    saida, d = [], d0
    while d <= d1:
        if _cal_us.eh_dia_util(d):
            saida.append(d)
        d += _dt.timedelta(days=1)
    return saida


sofr.serie_sofr = lambda d0, d1, *a, **k: [sofr.FixingSOFR(d, 0.0432)
                                           for d in _uteis_us(d0, d1)]
usd = dict(FORM)
usd.update({'ativa_indexador': liquidacao.SOFR, 'ativa_taxa': '1.5',
            'ativa_convencao': contagem.ACT_360, 'ativa_regime': contagem.SIMPLES,
            'ativa_moeda': 'USD', 'ativa_ptax_inicial': '5.10',
            'ativa_ptax_final': '5.42', 'ativa_lookback': '0', 'ativa_shift': '0'})
conteudo_us, _ = queries.memoria_de_calculo(usd)
wb4 = openpyxl.load_workbook(io.BytesIO(conteudo_us))
ws4, r4 = wb4[mx.ABA], queries.liquidar(usd)['r']
diaria4 = wb4[mx.ABA_DIARIA[liquidacao.ATIVA]]
check('a aba da ponta ATIVA aparece quando o indice e dela',
      mx.ABA_DIARIA[liquidacao.ATIVA] in wb4.sheetnames, True)
# n e a distancia ate a proxima linha; o ultimo fecha na janela de observacao,
# que esta no topo da aba. Contar de novo aqui seria uma segunda regra para
# divergir da primeira.
ns = [avaliar(wb4, mx.ABA_DIARIA[liquidacao.ATIVA], 'D%d' % (5 + i))
      for i in range(len(r4.ativa.fixings))]
check('o n do Excel e o n do motor, dia a dia',
      ns[:6], [(b.data - a.data).days for a, b in
               zip(r4.ativa.fixings, r4.ativa.fixings[1:])][:6])
check('   e a sexta-feira remunera tres dias', 3 in ns, True)
perto('   o fator composto do ultimo dia bate com o motor',
      avaliar(wb4, mx.ABA_DIARIA[liquidacao.ATIVA], 'F%d' % (4 + len(r4.ativa.fixings))),
      r4.ativa.fixings[-1].fator_acumulado, 1e-12)
perto('e o liquido da perna SOFR recalcula igual ao motor',
      avaliar(wb4, mx.ABA, por_rotulo(ws4, 'Ajuste líquido')), r4.ajuste_liquido, 1e-6)
sofr.serie_sofr = _serie_sofr_real

print('\n== 4c. IPCA e equity: correcao no principal, preco no indice ==')
infl = dict(FORM)
infl.update({'ativa_indexador': liquidacao.IPCA, 'ativa_taxa': '6.5',
             'ativa_ni_inicial': '7000.00', 'ativa_ni_final': '7350.00'})
conteudo_infl = queries.memoria_de_calculo(infl)[0]
wb5 = openpyxl.load_workbook(io.BytesIO(conteudo_infl))
ws5, r5 = wb5[mx.ABA], queries.liquidar(infl)['r']
perto('a correcao e o numero-indice final ÷ o inicial',
      avaliar(wb5, mx.ABA, por_rotulo(ws5, 'Fator de correção monetária')),
      r5.ativa.fator_correcao, 1e-12)
# So o CUPOM liquida no fluxo intermediario: a correcao fica no principal, que
# segue corrigido para o fluxo seguinte (a planilha da mesa faz assim).
perto('   e os juros do IPCA sao so o cupom sobre o principal corrigido',
      avaliar(wb5, mx.ABA, 'B%d' % [ln for ln in range(1, ws5.max_row + 1)
                                    if ws5.cell(row=ln, column=1).value
                                    == 'Juros do período (R$)'][0]),
      r5.ativa.juros, 1e-6)
perto('   e o liquido fecha', avaliar(wb5, mx.ABA, por_rotulo(ws5, 'Ajuste líquido')),
      r5.ajuste_liquido, 1e-6)

acao = dict(FORM)
acao.update({'ativa_indexador': liquidacao.EQUITY, 'ativa_taxa': '0',
             'ativa_ativo': 'FLRY3', 'ativa_preco_inicial': '18.50',
             'ativa_preco_final': '21.20'})
conteudo_acao = queries.memoria_de_calculo(acao)[0]
wb6 = openpyxl.load_workbook(io.BytesIO(conteudo_acao))
ws6, r6 = wb6[mx.ABA], queries.liquidar(acao)['r']
perto('o fator da perna de equity e a razao dos precos',
      avaliar(wb6, mx.ABA, por_rotulo(ws6, 'Fator do índice')), r6.ativa.fator_do_indice, 1e-12)
perto('   e o liquido fecha', avaliar(wb6, mx.ABA, por_rotulo(ws6, 'Ajuste líquido')),
      r6.ajuste_liquido, 1e-6)

print('\n== 4d. a liquidacao FINAL: a ponta inteira, e nada de valor futuro ==')
# No vencimento o principal liquida, entao o que a ponta entrega E o valor
# dela. O rotulo muda junto: "valor futuro" descreveria o mesmo numero como
# uma projecao de um fluxo que ainda nao aconteceu — e aqui ele aconteceu.
final = dict(FORM)
final['fim'] = final['vencimento'] = '2026-03-02'
wb7 = openpyxl.load_workbook(io.BytesIO(queries.memoria_de_calculo(final)[0]))
ws7, r7 = wb7[mx.ABA], queries.liquidar(final)['r']
check('o fluxo que fecha no vencimento liquida as duas pontas', r7.so_juros, False)
rotulos7 = [ws7.cell(row=ln, column=1).value for ln in range(1, ws7.max_row + 1)]
check('   e nem aqui a palavra "futuro" aparece',
      [x for x in rotulos7 if x and 'futuro' in str(x).lower()], [])
check('   a ponta fecha na linha do valor liquidado',
      (len([x for x in rotulos7 if x == 'Valor da ponta na liquidação (R$)']),
       'Juros do período (R$)' in rotulos7), (2, False))
linhas7 = [ln for ln in range(1, ws7.max_row + 1)
           if ws7.cell(row=ln, column=1).value == 'Valor da ponta na liquidação (R$)']
perto('   valor da ponta ativa', avaliar(wb7, mx.ABA, 'B%d' % linhas7[0]),
      r7.ativa.valor, 1e-6)
perto('   e o liquido fecha', avaliar(wb7, mx.ABA, por_rotulo(ws7, 'Ajuste líquido')),
      r7.ajuste_liquido, 1e-6)

# ─────────────────────────────────────────────────────────────────────────────
print('\n== 5. o documento e do banco ==')
z = zipfile.ZipFile(io.BytesIO(conteudo))
nomes = z.namelist()
check('o timbre vai dentro do arquivo', any(n.startswith('xl/media/') for n in nomes), True)
bruto = b''.join(z.read(n) for n in nomes if n.endswith(('.xml', '.rels'))).decode(
    'utf-8', 'ignore').lower()
for proibido in ('otc tracker', 'otc-tracker', 'otctracker', 'swap calculator', 'openpyxl'):
    check('nenhuma referencia a %r no arquivo' % proibido, proibido in bruto, False)
check('as propriedades dizem o banco', z.read('docProps/core.xml').decode('utf-8').count(
    'J.P. Morgan') >= 1, True)
check('o titulo da aba principal e a memoria', mx.ABA, 'Memória de Cálculo')

# ─────────────────────────────────────────────────────────────────────────────
print('\n== 6. a rota e a tela ==')
regras = {str(r_) for r_ in app.url_map.iter_rules()}
check('/tools/swap-calculator/extract registrada',
      '/tools/swap-calculator/extract' in regras, True)
metodos = {m for r_ in app.url_map.iter_rules()
           if str(r_) == '/tools/swap-calculator/extract' for m in r_.methods}
check('e so por POST (o que se exporta e o FORMULARIO)',
      ('POST' in metodos, 'GET' in metodos), (True, False))

c = app.test_client()
sem_sessao = c.post('/tools/swap-calculator/extract', data=FORM)
check('sem sessao redireciona para o login', sem_sessao.status_code, 302)
with c.session_transaction() as s:
    s.update(authenticated=True, user_sid='X1', user_name='t', user_role='ADMIN',
             session_expires_at=(_dt.datetime.now(_dt.timezone.utc).replace(tzinfo=None)
                                 + _dt.timedelta(hours=1)).isoformat())
resp = c.post('/tools/swap-calculator/extract', data=FORM)
check('com sessao, o download vem', resp.status_code, 200)
check('   com o tipo de .xlsx',
      'spreadsheetml.sheet' in resp.headers.get('Content-Type', ''), True)
disp = resp.headers.get('Content-Disposition', '')
check('   como anexo, e com o nome acentuado em RFC 5987',
      disp.startswith('attachment;') and 'Mem%C3%B3ria%20de%20C%C3%A1lculo' in disp, True)
check('   e o corpo e um .xlsx de verdade', resp.data[:2], b'PK')

# Conta que nao fecha nao baixa arquivo quebrado. Sem JavaScript volta a TELA
# com o motivo; pelo fetch do botao volta em JSON — e e o JSON que permite
# dizer o que faltou SEM recarregar a pagina e perder o formulario.
quebrado = dict(FORM); quebrado['fim'] = '2025-08-01'      # fim antes do inicio
erro = c.post('/tools/swap-calculator/extract', data=quebrado)
check('conta que nao fecha volta a tela, nao um arquivo',
      (erro.status_code, erro.headers.get('Content-Type', '').startswith('text/html')),
      (200, True))
erro_js = c.post('/tools/swap-calculator/extract', data=quebrado,
                 headers={'X-Requested-With': 'XMLHttpRequest'})
check('   e pelo fetch volta em JSON, com o motivo',
      (erro_js.status_code, erro_js.get_json().get('success'),
       bool(erro_js.get_json().get('error'))), (422, False, True))

html = io.open(os.path.join(ROOT, 'apps/templates/pages/tools-swap-calculator.html'),
               encoding='utf-8').read()
check('o botao Export sai do mesmo formulario, por formaction',
      'tools_swap_calculator_extract' in html and 'formaction=' in html, True)
# Sem o `action` explicito a tela passaria a postar na rota do arquivo depois
# do primeiro Export, e o Calculate seguinte baixaria uma planilha.
check('e o formulario declara action explicito',
      "action=\"{{ url_for('pages_blueprint.tools_swap_calculator') }}\"" in html, True)
check('o rotulo nasce em ingles e traduz por data-lang',
      'data-lang="tl-export">Export<' in html, True)
# `submit` de verdade: o download funciona com o JS fora do ar, e o fetch so
# acrescenta o fim do spinner.
check('   e o botao continua sendo um submit, nao um type=button',
      'type="submit" id="tl-export"' in html, True)
import json                                                             # noqa: E402
for lang in ('en', 'br', 'es'):
    d = json.load(io.open(os.path.join(ROOT, 'apps/static/data/translations/%s.json' % lang),
                          encoding='utf-8'))
    check('   %s traduz o tl-export' % lang, bool(d.get('tl-export')), True)
    check('   e o tl-extract antigo saiu de %s' % lang, 'tl-extract' in d, False)

print('\n== 7. o spinner: quem liga desliga ==')
js = io.open(os.path.join(ROOT, 'apps/static/js/pages/tools.js'), encoding='utf-8').read()
check('o clique troca o botao por um spinner',
      "spinner-border spinner-border-sm" in js and "t('exporting')" in js, True)
# Navegacao que baixa arquivo nao emite evento nenhum: sem o corpo na mao o
# spinner giraria para sempre. Por isso o POST vai por fetch.
check('e o fim vem do fetch — sucesso, erro e falha de rede, os tres',
      js.count('terminar()') >= 3, True)
check('o nome do arquivo sai do Content-Disposition, nao da rota',
      "filename\\*=UTF-8''" in js, True)
# O mapa local e o que o I18nManager nao alcanca: ele traduz no load, e este
# texto nasce depois (§2).
_MARCAS = ['    en: { show:', '    br: { show:', '    es: { show:', '  };']
for i, lang in enumerate(('en', 'br', 'es')):
    bloco = js.split(_MARCAS[i], 1)[1].split(_MARCAS[i + 1], 1)[0]
    check('   %s tem exporting e exportFail no _TRANS' % lang,
          ('exporting:' in bloco, 'exportFail:' in bloco), (True, True))

print('\n== 8. o documento se le como documento ==')
# Alinhado a direita, a coluna de valores mistura texto com numero e cada linha
# comeca num ponto diferente — o nome da contraparte descola do rotulo. Numa
# planilha de LEITURA, a coluna unica e o que a vista espera.
desalinhadas = [ln for ln in range(4, ws.max_row + 1)
                if ws.cell(row=ln, column=2).value is not None
                and ws.cell(row=ln, column=2).alignment.horizontal != 'left']
check('todo valor alinhado a esquerda', desalinhadas, [])

# Ancorado rente ao canto de A1 o wordmark encostava na moldura da celula e
# ficava escondido. A ancora com deslocamento (EMU) e a unica forma de dar
# margem e centra-lo no bloco de duas linhas do cabecalho.
img = (ws._images or [None])[0]
check('o timbre entra na planilha', img is not None, True)
if img is not None:
    check('   com folga em volta, nao colado no canto da celula',
          (img.anchor._from.colOff > 0, img.anchor._from.rowOff > 0), (True, True))
    check('   e o cabecalho tem altura para ele',
          (ws.row_dimensions[1].height or 0) + (ws.row_dimensions[2].height or 0) >= 50, True)

# O documento vai para o cliente e para a auditoria: nada de linguagem de
# conversa nem de aula. O leitor sabe o que e um swap.
texto = []
for aba in wb.sheetnames:
    for linha in wb[aba].iter_rows():
        for cel in linha:
            if isinstance(cel.value, str) and not cel.value.startswith('='):
                texto.append(cel.value)
tudo = ' | '.join(texto).lower()
for frase in ('troca de mãos', 'quem paga', 'quem recebe', 'nao troca', 'o que a ponta',
              'acontece no', 'so o que a taxa', 'da 15,5031'):
    check('nada de %r no documento' % frase, frase in tudo, False)
check('a parte devedora e nomeada',
      ws[por_rotulo(ws, 'Parte devedora')].value,
      'Banco J.P. Morgan' if r.banco_paga else FORM['counterparty'])

# ─────────────────────────────────────────────────────────────────────────────
print('\n== 9. o numero CHEGA: toda formula leva o valor gravado ==')
# O openpyxl escreve a formula com o cache VAZIO, e quem abre o arquivo sem
# recalcular — o Modo de Exibicao Protegido (todo arquivo baixado pelo
# navegador entra nele), o painel de visualizacao, o Excel Online, o preview do
# anexo no Outlook — mostra a celula EM BRANCO. Era a memoria inteira chegando
# a mesa sem um numero: os rotulos e as entradas apareciam, e tau, os fatores,
# os juros, o IR e o ajuste liquido ficavam vazios. O `fullCalcOnLoad` nao
# alcanca esse caso: ele manda recalcular na ABERTURA, e nenhum desses leitores
# calcula.
#
# O que se prende aqui e que o arquivo leva as DUAS coisas — a formula (secao 2
# ja provou que ela e a conta do motor) e o resultado dela gravado — e que o
# valor gravado e o da PROPRIA formula: um cache escrito a mao seria a planilha
# afirmando um numero que a formula nao da, que e pior do que a celula vazia.
for indice, bytes_ in (('CDI', conteudo), ('SOFR', conteudo_us),
                       ('cambial', conteudo_fx), ('IPCA', conteudo_infl),
                       ('equity', conteudo_acao), ('sem aba diaria', conteudo_cru)):
    formulas = openpyxl.load_workbook(io.BytesIO(bytes_))
    valores = openpyxl.load_workbook(io.BytesIO(bytes_), data_only=True)
    vazias, divergentes, total = [], [], 0
    for aba in formulas.sheetnames:
        for linha in formulas[aba].iter_rows():
            for cel in linha:
                if not (isinstance(cel.value, str) and cel.value.startswith('=')):
                    continue
                total += 1
                gravado = valores[aba][cel.coordinate].value
                if gravado is None:
                    vazias.append('%s!%s' % (aba, cel.coordinate))
                    continue
                esperado = avaliar(formulas, aba, cel.value[1:])
                if abs(gravado - esperado) > max(1e-9, abs(esperado) * 1e-12):
                    divergentes.append('%s!%s' % (aba, cel.coordinate))
    check('%s: nenhuma formula sem valor gravado (de %d)' % (indice, total), vazias[:5], [])
    check('   e o valor gravado e o da formula', divergentes[:5], [])

# ── 9b. e chega em QUALQUER maquina, nao so nesta ──────────────────────────
# A secao acima passa ou falha conforme o que esta instalado no ambiente que a
# roda, e nao percebe: o openpyxl troca de serializador de XML conforme tenha
# ou nao o `lxml`. Com ele, o escritor incremental abre e fecha a tag do cache
# (`<v></v>`); sem ele, o ElementTree serializa o elemento vazio como `<v />`.
# Um injetor que case so com uma das grafias devolve, na maquina que tem a
# outra, o arquivo EXATAMENTE como saia antes — sem um valor e sem erro nenhum.
# `lxml` nao esta no requirements, entao a maquina do desenvolvedor e a da
# instancia do time caem em lados diferentes disso. Aqui as duas grafias sao
# escritas a mao e cobradas, para o guarda nao depender do que o ambiente tem.
print('\n== 9b. o valor entra nas DUAS grafias do cache vazio ==')
MOLDE = ('<row r="5"><c r="A5" s="4" t="inlineStr"><is><t>rotulo</t></is></c>'
         '<c r="B5" s="6"><f>A1*2</f>%s</c></row>')
for apelido, vazio in (('lxml (<v></v>)', '<v></v>'),
                       ('ElementTree (<v />)', '<v />'),
                       ('ElementTree sem espaco (<v/>)', '<v/>'),
                       ('sem cache nenhum', '')):
    saida = mx._injetar_cache(MOLDE % vazio, {'B5': 42.5})
    check('%s: a formula sai com o valor' % apelido,
          '<f>A1*2</f><v>42.5</v>' in saida, True)
    check('   e o rotulo ao lado fica intacto', '<t>rotulo</t>' in saida, True)
check('inteiro sai sem casa decimal',
      '<v>7</v>' in mx._injetar_cache(MOLDE % '<v />', {'B5': 7.0}), True)
check('celula que o cache nao conhece fica como estava',
      mx._injetar_cache(MOLDE % '<v />', {}), MOLDE % '<v />')

cdi.serie = _serie_real
print('\n' + ('TUDO OK' if not falhas else 'FALHAS: %d' % len(falhas)))
for f in falhas:
    print('  - ' + f)
sys.exit(1 if falhas else 0)
