"""`scripts/fix_asian_dates_xlsx.py`: reordena, ajusta a janela e para sem feriado.

Planilha montada como o Excel do Live Position Option sai hoje (datas como
TEXTO `dd/mm/aaaa`, bloco da Média Asiática começando na BI) e feriados de
teste por `--feriados`, sem tocar dado real.

    OTC_SHARED_DRIVE_ROOT=/tmp/otc-share python scripts/tests/check_fix_asian_dates.py
"""
import json
import os
import sys
import tempfile
from datetime import date, datetime, timedelta

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
sys.path.insert(0, os.path.join(RAIZ, 'scripts'))
from openpyxl import Workbook, load_workbook   # noqa: E402
import fix_asian_dates_xlsx as F                # noqa: E402

falhas = 0


def ok(cond, msg):
    global falhas
    print(('  ok    ' if cond else '  FAIL  ') + msg)
    if not cond:
        falhas += 1


FERIADOS = [date(2026, 1, 1), date(2026, 4, 3), date(2026, 4, 6), date(2026, 12, 25)]


def uteis(a, m):
    return F.dias_uteis_do_mes(a, m, set(FERIADOS))


def dmy(d):
    return d.strftime('%d/%m/%Y')


tmp = tempfile.mkdtemp()
arq = os.path.join(tmp, 'Live Position Option OTC Tracker - Sistema de Gestão OTC.xlsx')
fer = os.path.join(tmp, 'ipe.json')
json.dump([{'date': d.isoformat()} for d in FERIADOS], open(fer, 'w'))

wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'
cab = ['Código Identificador'] + ['Col %d' % i for i in range(2, 61)] \
    + ['Média Asiática (data) %d' % i for i in range(1, 24)]
ws.append(cab)
abr = uteis(2026, 4)                                          # 20 dias úteis (Páscoa)
fora_ordem = abr[5:] + abr[:5]
jan_fev = [d for d in uteis(2026, 1) if d.day >= 26] + uteis(2026, 2)[:17]   # 5 jan + 17 fev
empate = uteis(2026, 5)[-3:] + uteis(2026, 6)[:3]
certa = uteis(2026, 3)
for idt, datas in (('OPC-1', fora_ordem), ('OPC-2', jan_fev), ('OPC-3', empate), ('OPC-4', certa)):
    ws.append([idt] + ['x'] * 59 + [dmy(d) for d in datas])
ws.append(['OPC-5'] + ['x'] * 59)                             # sem datas
wb.save(arq)

ok(F.achar_bloco(ws)[1][0] == 61, 'bloco achado pelo rótulo na BI (coluna 61)')
ok(F.para_data('03/04/2026') == date(2026, 4, 3), 'texto dd/mm/aaaa vira data (dia primeiro)')
ok(F.para_data(46115) == date(2026, 4, 3), 'serial do Excel vira data')

# sem feriado no ano: PARA antes de gravar
vazio = os.path.join(tmp, 'vazio.json')
json.dump([{'date': '2025-12-25'}], open(vazio, 'w'))
try:
    F.main([arq, '--feriados', vazio])
    ok(False, 'calendário sem 2026 deveria parar')
except SystemExit as e:
    ok('PARADO' in str(e), 'calendário sem feriado no ano da janela PARA')
ok('Ajustado' not in load_workbook(arq).sheetnames, 'parado não grava a aba')

ok(F.main([arq, '--feriados', fer]) == 0, 'roda')
wb2 = load_workbook(arq)
ok(wb2.sheetnames == ['Sheet1', 'Ajustado', 'Log Ajuste Datas'], 'abas: origem, Ajustado, Log')
aj, orig = wb2['Ajustado'], wb2['Sheet1']


def linha(r):
    return [aj.cell(r, c).value for c in range(61, 84)]


def datas(r):
    return [v.date() for v in linha(r) if v is not None]


ok(all(aj.cell(r, c).value == orig.cell(r, c).value for r in range(1, 7) for c in range(1, 61)),
   'A:BH copiadas iguais')
ok(datas(2) == abr, 'fora de ordem: mesmas datas, reordenadas')
ok(isinstance(aj.cell(2, 61).value, datetime) and aj.cell(2, 61).number_format == 'dd/mm/yyyy',
   'data sai como DATA dd/mm/yyyy, não texto')
ok(datas(3) == uteis(2026, 2) and len(datas(3)) == 20,
   'janela jan×fev: fevereiro inteiro (mês com mais datas), primeiro ao último dia útil')
ok(linha(3)[20:] == [None, None, None], 'janela com menos dias úteis limpa as colunas que sobram')
ok(datas(4) == empate, 'empate: não ajustada')
ok(datas(5) == certa, 'já certa: igual')
ok(aj.cell(2, 61).fill.fgColor.rgb.endswith('FFF2CC') and not aj.cell(5, 61).fill.fgColor.rgb.endswith('FFF2CC'),
   'célula mexida em amarelo, intocada sem cor')
ok(date(2026, 4, 3) not in datas(2) and date(2026, 4, 6) not in datas(2), 'fora de ordem não inventa feriado')
log = {r[0]: r for r in wb2['Log Ajuste Datas'].iter_rows(min_row=2, values_only=True)}
ok(log[2][2] == 'Reordenada' and log[3][2] == 'Janela ajustada' and log[4][2].startswith('Não ajustada')
   and log[5][2] == 'Sem alteração', 'log diz a ação de cada linha')
ok(log[3][1] == 'OPC-2' and 'QUANTIDADE %d → 20' % len(jan_fev) in log[3][5], 'log identifica o trade e a mudança de quantidade')
ok(6 not in log, 'linha sem datas fica fora do log')
ok(os.path.exists(arq.replace('.xlsx', ' - original.xlsx')), 'cópia do original ao lado')

# rodar de novo: substitui a aba, lê sempre a origem
ok(F.main([arq, '--feriados', fer]) == 0 and load_workbook(arq).sheetnames.count('Ajustado') == 1,
   'rodar de novo substitui a aba')

print('\nall ok' if not falhas else '\nFALHAS: %d' % falhas)
sys.exit(1 if falhas else 0)
