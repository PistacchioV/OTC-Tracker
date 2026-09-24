# -*- coding: utf-8 -*-
"""Regressão da página New Deals › Swap › Bullet (§480).

O que este script prende:

1. o PARSER do Deal Ticket em xlsx (uma aba por operação: o DT contra o
   cliente e o B2B Banco × Atacama) acha os rótulos pela coluna dos blocos
   (Curva VCP / Curva Vanilla / Informações) e monta o deal com Parte A = a
   nossa perna, Parte B = a contraparte, decidido pelo `Ativo VCP`;
2. o mesmo DT em PDF (o texto que o pypdf devolve) chega ao MESMO deal;
3. a DENOMINAÇÃO da curva VCP é a da fórmula do Excel da mesa, byte a byte;
4. os TRÊS registros do 0301 (Cliente, Banco, Atacama) são byte a byte os
   exemplos que a mesa mandou (Meu Número e data do header à parte), salvo
   a posição onde a regra da página diverge do exemplo DE PROPÓSITO e o
   teste diz qual (Cap só na perna VCP);
5. os três arquivos de PRÊMIO (0897) idem — Papel e Titular pela conta menor;
6. o import grava no arquivo-dia da Trade Date, o re-import preserva a
   esteira e os Meu Número; o Send escreve os arquivos e vira Sent; deal com
   lacuna recusa o LOTE inteiro;
7. o template `swap-registro-premio` está na biblioteca ligado à página, o
   `swap-pagamento-final-v3` também, e os cadastros novos existem.

Roda em tmp: cache e pasta Conecta apontam para diretórios temporários.
"""
import io
import re
import os
import sys
import tempfile
from datetime import datetime, date

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..')))
from apps.config import Config as _Cfg                      # noqa: E402
_Cfg.FOUR_EYES = True   # prova a trava da PROD; na dev o maker/checker e desligado (§553)
os.environ.setdefault('OTC_DISABLE_SCHEDULERS', '1')
if os.name != 'nt' and not os.environ.get('OTC_SHARED_DRIVE_ROOT'):
    os.environ['OTC_SHARED_DRIVE_ROOT'] = tempfile.mkdtemp(prefix='otc-share-')

FALHAS = []


def check(nome, cond):
    print(('  ok  ' if cond else '  FAIL ') + nome)
    if not cond:
        FALHAS.append(nome)


# ── Os exemplos da mesa (registros tipo 1) ───────────────────────────────────
# Os registros tipo 1 que a mesa mandou (VERBATIM — 1927 caracteres cada).
EX = {
 'client': 'SWAP 10301898571069873760009                        74220005                        2026091520270607010000000034600000060001                                                                                                                                                                                                                                                                                                                                                                                                                                                                                                                        10000C99        000000000                                10000C00                                 0000011700000000                                                                                                                                                                                                                                                                                                                                                                                              000000000000010000000010615                                                                                                                                                                                                                                                                                                                                                 00000100000000001                                                                                                                                                                                                                                                                                                        010000000000000000                                                                           EDG                              ',
 'bank': 'SWAP 10301523903246573760009                        85398005                        2026091520270607010000000034600000060001                                                                                                                                                                                                                                                                                                                                                                                                                                                                                                                        10000C00                                 000001170000000010000C99        000000000                0000011700000000                                   000000000000010000000010615MSFT US : Indices Internacionais Codigo 10615 – Descricao: Microsoft Corporation – Preco Inicial: 100.00% Spot – Fonte de informacao: Bloomberg – Data de cotacao: 04-jun-2027 - Cupom limpo = Strike - Preco in ativo Close 15-Sep-26 - Denominacao: Quanto (ausencia variacao cambial)                                                                                                                                                                                                                                                                                                                                                                                                   00000100000000001                                                                                                                                                                                                                                                                                                                         000000000000000000                                                                           EDG                              ',
 'atacama': 'SWAP 10301386363168785398005                        73760009                        2026091520270607010000000034600000060001                                                                                                                                                                                                                                                                                                                                                                                                                                                                                                                        10000C99        000000000                000001170000000010000C00                                 0000011700000000                                                                                                                                                                                                                                                                                                                                                                                              000000000000010000000010615                                                                                                                                                                                                                                                                                                                                                 00000100000000001                                                                                                                                                                                                                                                                                                        010000000000000000                                                                           EDG                              ',
}
MYNUM = {'client': '8985710698', 'bank': '5239032465', 'atacama': '3863631687'}
PREM = {
 'client':  ("SWAP 00897JPMORGANBM          20260916              ", "SWAP 10897000189857106987041612026737600097422000500", "2026091600000233550001" + " " * 30),
 'bank':    ("SWAP 00897JPMORGANBM          20260916              ", "SWAP 10897000152390324654368513572737600098539800500", "2026091600000233550000" + " " * 30),
 'atacama': ("SWAP 00897INTRAGATACAMAFDO    20260916              ", "SWAP 10897000138636316873024019654853980057376000901", "2026091600000233550000" + " " * 30),
}
PREMNUM = {'client': '7041612026', 'bank': '4368513572', 'atacama': '3024019654'}


def _dt_sheet(ws, client, spn, vcp_holder, van_holder, payer, categoria):
    ws['A1'] = 'DEAL TICKET SWAP VCP'
    ws['A4'] = 'Swap Com Opção de Arrependimento - Call'
    ws['A7'] = 'Cliente'; ws['B7'] = client
    ws['A8'] = 'SPN'; ws['B8'] = spn
    ws['A9'] = 'Inicio'; ws['B9'] = datetime(2026, 9, 15); ws['D9'] = 'Funcionalidades'; ws['E9'] = 'Opção de Arrependimento'; ws['G9'] = 'Agenda de premios'; ws['H9'] = 'Sim'
    ws['A10'] = 'Vencimento'; ws['B10'] = datetime(2027, 6, 7); ws['D10'] = 'Ativo VCP'; ws['E10'] = vcp_holder; ws['G10'] = 'Data de Pagamento do Premio'; ws['H10'] = datetime(2026, 9, 16)
    ws['A11'] = 'Valor Base'; ws['B11'] = 'BRL 346.000,00'; ws['D11'] = 'Ativo Curva Vanilla'; ws['E11'] = van_holder; ws['G11'] = 'Pagador do Premio'; ws['H11'] = payer
    ws['G12'] = 'Premio'; ws['H12'] = 23355
    ws['A15'] = 'Curva VCP'; ws['D15'] = 'Curva Vanilla'; ws['G15'] = 'Informações Curvas VCP'
    ws['A16'] = 'Percentual'; ws['B16'] = 1.0; ws['B16'].number_format = '0.00%'; ws['D16'] = 'Percentual'; ws['E16'] = 1.0; ws['E16'].number_format = '0.00%'
    ws['G16'] = 'Preço Inicial(Cupom Limpo)'; ws['H16'] = '100.00% Spot'
    ws['A17'] = 'Categoria'; ws['B17'] = 'VCP'; ws['D17'] = 'Categoria'; ws['E17'] = 'JUROS'; ws['G17'] = 'Fonte de Informação'; ws['H17'] = 'Bloomberg'
    ws['A18'] = 'Curva'; ws['B18'] = 'MSFT US'; ws['D18'] = 'Curva'; ws['E18'] = 'PRÉ FIXADO BRL'; ws['G18'] = 'Data de Cotação'; ws['H18'] = datetime(2027, 6, 4)
    ws['A20'] = 'Sinal +/-'; ws['B20'] = '+'; ws['D20'] = 'Sinal +/-'; ws['E20'] = '+'; ws['G20'] = 'Denominação'; ws['H20'] = 'Quanto (ausencia variação cambial)'
    ws['A21'] = 'Juros'; ws['B21'] = 0; ws['B21'].number_format = '0.00%'; ws['D21'] = 'Juros'; ws['E21'] = 0; ws['E21'].number_format = '0.00%'; ws['H21'] = 'Cupom Limpo inicial em percentual'
    ws['A22'] = 'Lim Superior'; ws['B22'] = '117% Spot'; ws['H22'] = 'Preco in ativo Close 15-Sep-26'
    ws['A23'] = 'Lim Inferior'
    ws['A25'] = 'OBS: Caso aplicável, colocar também no campo denominação as informações dos limitadores'
    ws['A27'] = 'Informações do Indicador'
    ws['A29'] = 'Categoria da Curva VCP'; ws['B29'] = categoria
    ws['A30'] = 'Código'; ws['B30'] = 10615
    ws['A31'] = 'Curva VCP'; ws['B31'] = 'MSFT US'
    ws['A32'] = 'Descrição da Curva VCP'; ws['B32'] = 'Microsoft Corporation'
    ws['A34'] = 'Informações necessárias para registro'; ws['B34'] = 'Consultar o Manual de Operações de Swap'
    ws['A35'] = 'Link para o Manual de Operações'; ws['B35'] = 'www.b3.com.br'


def build_xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active; ws.title = 'Deal Ticket Cliente (CETIP)'
    _dt_sheet(ws, 'Safra', 281808, 'Cliente', 'Banco JP Morgan', 'Cliente', 'ACOES INTERNACIONAIS')
    ws2 = wb.create_sheet('B2B - Atacama')
    _dt_sheet(ws2, 'Atacama', 9632845, 'Banco JP Morgan', 'Atacama', 'Banco JP Morgan', 'INDICES INTERNACIONAIS')
    ws3 = wb.create_sheet('Recap'); ws3['A1'] = 'Recap'; ws3['A2'] = 'nada aqui'
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def build_pdf_text():
    """O texto que o pypdf devolve de um DT impresso: colunas coladas na mesma
    linha, uma operação por página."""
    return (
        "DEAL TICKET SWAP VCP\n"
        "Swap Com Opção de Arrependimento - Call\n"
        "Cliente Safra\n"
        "SPN 281808\n"
        "Inicio 15-Sep-26 Funcionalidades Opção de Arrependimento Agenda de premios Sim\n"
        "Vencimento 7-Jun-27 Ativo VCP Cliente Data de Pagamento do Premio 16-Sep-26\n"
        "Valor Base BRL 346.000,00 Ativo Curva Vanilla Banco JP Morgan Pagador do Premio Cliente\n"
        "Premio 23.355,00\n"
        "Curva VCP Curva Vanilla Informações Curvas VCP\n"
        "Percentual 100,00% Percentual 100,00% Preço Inicial(Cupom Limpo) 100.00% Spot\n"
        "Categoria VCP Categoria JUROS Fonte de Informação Bloomberg\n"
        "Curva MSFT US Curva PRÉ FIXADO BRL Data de Cotação 04/jun/27\n"
        "Sinal +/- + Sinal +/- + Denominação Quanto (ausencia variação cambial)\n"
        "Juros 0,00% Juros 0,00% Cupom Limpo inicial em percentual\n"
        "Lim Superior 117% Spot Preco in ativo Close 15-Sep-26\n"
        "Lim Inferior\n"
        "OBS: Caso aplicável, colocar também no campo denominação as informações dos limitadores\n"
        "Informações do Indicador\n"
        "Categoria da Curva VCP ACOES INTERNACIONAIS\n"
        "Código 10615\n"
        "Curva VCP MSFT US\n"
        "Descrição da Curva VCP Microsoft Corporation\n"
        "Informações necessárias para registro Consultar o Manual\n"
    )


def _diff(a, b, mask_ranges=()):
    """Primeira posição (1-based) em que a≠b fora das faixas mascaradas."""
    for i, (x, y) in enumerate(zip(a, b), 1):
        if x != y and not any(lo <= i <= hi for lo, hi in mask_ranges):
            return i
    return 0 if len(a) == len(b) else min(len(a), len(b)) + 1


def main():
    from run import app  # noqa: F401
    from apps.pages import routes as R
    from apps.pages.features.swap_bullet import commands, domain, queries
    from apps.pages.features.swap_bullet.infra import dt_reader, persistence

    tmp = tempfile.mkdtemp(prefix='otc-swb-')
    persistence.cache_dir = lambda: os.path.join(tmp, 'cache')
    R.CONECTA_NEW_PATH = os.path.join(tmp, 'conecta')

    print('== 1. o DT em xlsx: duas abas, dois deals ==')
    xlsx = build_xlsx()
    sheets = dt_reader.sheets_from_xlsx(xlsx)
    check('três abas lidas, duas são DT', [domain.is_dt_grid(g) for _t, g in sheets] == [True, True, False])
    raw = domain.parse_dt_grid(sheets[0][1])
    check('rótulos do topo', raw.get('Client') == 'Safra' and raw.get('SPN') == '281808'
          and raw.get('Notional') == 'BRL 346.000,00' and raw.get('Functionality') == 'Opção de Arrependimento')
    check('pernas pelo bloco: VCP e Vanilla', raw.get('vcp_Curve') == 'MSFT US' and raw.get('vanilla_Curve') == 'PRÉ FIXADO BRL'
          and raw.get('vcp_Cap') == '117% Spot' and raw.get('vanilla_Pct') == '100%')
    check('denominação em três linhas', raw.get('Denomination') == ['Quanto (ausencia variação cambial)', 'Cupom Limpo inicial em percentual', 'Preco in ativo Close 15-Sep-26'])
    check('indicador', raw.get('VcpCode') == '10615' and raw.get('VcpCurve') == 'MSFT US' and raw.get('VcpDescription') == 'Microsoft Corporation'
          and raw.get('VcpCategory') == 'ACOES INTERNACIONAIS')
    cli = domain.deal_from_raw(raw, '2026-09-16')
    check('Parte A = nossa perna (JUROS), Parte B = cliente (VCP)',
          cli['CurveACategory'] == 'JUROS' and cli['CurveA'] == 'PRÉ FIXADO BRL' and cli['CurveBCategory'] == 'VCP'
          and cli['CurveB'] == 'MSFT US' and cli['CurveBCap'] == '117' and cli['CurveACap'] == '')
    check('datas ISO, notional, prêmio', cli['StartDate'] == '2026-09-15' and cli['MaturityDate'] == '2027-06-07'
          and cli['Notional'] == '346000.00' and cli['PremiumAmount'] == '23355.00' and cli['PremiumDate'] == '2026-09-16'
          and cli['PremiumSchedule'] == 'Sim' and cli['PremiumPayer'] == 'Cliente' and cli['Currency'] == 'BRL')
    check('par JPM x CLI, LE JPM, id interno determinístico, Deal e B3 ID em BRANCO', cli['Pair'] == 'JPM x CLI' and cli['LE'] == 'JPM' and cli['_id'].startswith('SWB-')
          and domain.deal_from_raw(raw, '2026-09-16')['_id'] == cli['_id'] and cli['Deal'] == '' and cli['B3ID'] == '')
    b2b = domain.deal_from_raw(domain.parse_dt_grid(sheets[1][1]), '2026-09-16')
    check('B2B: par JPM x ATACAMA, LE ATACAMA, Parte A = VCP', b2b['Pair'] == 'JPM x ATACAMA' and b2b['LE'] == 'ATACAMA' and b2b['CurveACategory'] == 'VCP'
          and b2b['CurveACap'] == '117' and b2b['CurveBCategory'] == 'JUROS' and b2b['PremiumPayer'] == 'Banco JP Morgan')

    print('== 1b. o DT do GLD: sem Cap/Floor, preço em % do Spot, ano de dois dígitos ==')
    from openpyxl import Workbook as _WB
    wb2 = _WB(); w = wb2.active; w.title = 'DT'
    _dt_sheet(w, 'Safra', 281808, 'Cliente', 'Banco JP Morgan', 'Cliente', 'INDICES INTERNACIONAIS')
    w['B9'] = datetime(2026, 8, 14); w['B10'] = datetime(1931, 8, 18)         # o Excel leu '18-Aug-31' como 1931
    w['B11'] = 'BRL 1.100.000,00'; w['H10'] = datetime(2026, 8, 17); w['H12'] = 184140
    w['B18'] = 'GLD UP EQUITY'; w['G16'] = 'Preço Inicial(Cupom Limpo) - % do Spot'; w['H16'] = 1.5; w['H16'].number_format = '0.00%'
    w['H18'] = datetime(1931, 8, 15); w['B22'] = None; w['H22'] = 'Preco in ativo - 1.5 Close 14-Aug-26'
    w['B30'] = 10473; w['B31'] = 'GLD UP EQUITY'; w['B32'] = 'SPDR GOLD SHARES'
    bb = io.BytesIO(); wb2.save(bb)
    g2 = dt_reader.sheets_from_xlsx(bb.getvalue())[0][1]
    gld = domain.deal_from_raw(domain.parse_dt_grid(g2), '2026-08-14')
    check('Lim Superior/Inferior em branco não pegam a Denominação do bloco vizinho', gld['CurveBCap'] == '' and gld['CurveBFloor'] == '')
    check("'Preço Inicial(Cupom Limpo) - % do Spot' é o Preço Inicial, em % na grade", gld['InitialPrice'] == '150%')
    check('ano de dois dígitos lido pelo Excel como 1931 vira 2031', gld['MaturityDate'] == '2031-08-18' and gld['QuoteDate'] == '2031-08-15')
    check('a 3ª linha da denominação fica na denominação', gld['Denomination3'] == 'Preco in ativo - 1.5 Close 14-Aug-26')
    gld.update(ClientAccount='74220005', ClientRefData='ok', QuoteDateCode='01')
    vg = domain.swap_record_values(gld, 'client', queries.own_accounts(), queries.codes_for(gld), '1111111111')
    check('no arquivo, 150% é 150 no Cupom Limpo (9(08)v9(07)) e o PU segue 1.0; Cap/Floor em branco',
          vg['55'] == '000001500000000' and vg['50'] == '0000000000000100000000' and vg['38'].strip() == '' and vg['39'].strip() == '')
    print('== 2. o DT em PDF chega ao mesmo deal ==')
    raw_pdf = domain.parse_dt_text(build_pdf_text())
    pdf = domain.deal_from_raw(raw_pdf, '2026-09-16')
    campos = [k for k in domain.SWB_FIELDS if k not in ('Maker', 'Checker')]
    dif = [k for k in campos if pdf.get(k) != cli.get(k)]
    check('mesmos campos do xlsx (%s)' % (', '.join(dif) or 'nenhuma diferença'), not dif)
    if dif:
        for k in dif:
            print('      %s: pdf=%r xlsx=%r' % (k, pdf.get(k), cli.get(k)))

    print('== 2b. um PDF de verdade (reportlab → pypdf) chega ao mesmo deal ==')
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4, landscape
    except ImportError:
        canvas = None
    if canvas is None:
        print('  skip  reportlab ausente')
    else:
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=landscape(A4))
        c.setFont('Helvetica', 8)
        y = 560
        for ln in build_pdf_text().split('\n'):
            if ln.strip():
                c.drawString(30, y, ln)
            y -= 14
        c.showPage(); c.save()
        pages = dt_reader.pages_from_pdf(buf.getvalue())
        pdf2 = domain.deal_from_raw(domain.parse_dt_text(pages[0]), '2026-09-16')
        dif2 = [k for k in campos if pdf2.get(k) != cli.get(k)]
        check('pypdf: mesmos campos do xlsx (%s)' % (', '.join(dif2) or 'nenhuma diferença'), not dif2)
        for k in dif2:
            print('      %s: pdf=%r xlsx=%r' % (k, pdf2.get(k), cli.get(k)))
        kind, itens = dt_reader.read_upload('dt.pdf', buf.getvalue())
        check('read_upload reconhece .pdf', kind == 'text' and len(itens) == 1)

    print('== 3. a denominação da curva VCP é a da fórmula da mesa ==')
    esperado = ('MSFT US : Indices Internacionais Codigo 10615 – Descricao: Microsoft Corporation – Preco Inicial: 100.00% Spot '
                '– Fonte de informacao: Bloomberg – Data de cotacao: 04-jun-2027 - Cupom limpo = Strike - Preco in ativo Close 15-Sep-26 '
                '- Denominacao: Quanto (ausencia variacao cambial)')
    check('texto byte a byte (B2B: Indices Internacionais)', b2b['VcpText'] == esperado)
    if b2b['VcpText'] != esperado:
        print('      got:', repr(b2b['VcpText']))
    check('cliente: Acoes Internacionais', cli['VcpText'].startswith('MSFT US : Acoes Internacionais Codigo 10615'))

    print('== 4. os três registros 0301 ==')
    accounts = queries.own_accounts()
    check('contas próprias do b3-accounts', accounts.get('JPM') == '73760009' and accounts.get('ATACAMA') == '85398005')
    cli['ClientAccount'], cli['ClientTaxId'], cli['ClientRefData'] = '74220005', '', 'ok'
    for d in (cli, b2b):
        commands.enrich(d)
    cli['ClientRefData'] = 'ok'          # o RefData da dev não tem a SPN 281808; a regra em si é testada abaixo
    check('Data de Cotação D-1 (04/06 → 07/06/2027)', cli['QuoteDateCode'] == '01' and b2b['QuoteDateCode'] == '01')
    codes = queries.codes_for(cli)
    check('códigos pelos cadastros', codes['functionality'] == '06' and codes['adhesion'] == '01' and codes['curveA'] == 'C99'
          and codes['curveB'] == 'C00' and codes['signA'] == '00' and codes['premium_schedule'] == '00' and codes['reset'] == '01')
    check('sem lacunas', domain.missing_for_send(cli, codes, accounts) == [] and domain.missing_for_send(b2b, queries.codes_for(b2b), accounts) == [])
    # Onde a regra da página diverge dos exemplos DE PROPÓSITO: o Cap vai só na
    # perna em que o DT o declara (a VCP); o exemplo do Banco e o da Atacama
    # traziam 117 também na perna JUROS (campos 32/39, pos 670-685 / 727-742).
    # A Descrição segue os exemplos à letra: só na curva da PARTE quando é VCP.
    MASK = {'client': [], 'bank': [(727, 742)], 'atacama': [(670, 685)]}
    for view, deal in (('client', cli), ('bank', b2b), ('atacama', b2b)):
        vals = domain.swap_record_values(deal, view, accounts, queries.codes_for(deal), MYNUM[view])
        linha = commands._build_blocks(commands.SWAP_FI_KEY, vals, view, deal)
        check('%s: 1927 caracteres' % view, len(linha) == 1927)
        pos = _diff(linha, EX[view], MASK[view])
        check('%s: byte a byte com o exemplo da mesa (fora das divergências declaradas)' % view, pos == 0)
        if pos:
            print('      1ª diferença na posição %d: got=%r exp=%r' % (pos, linha[pos-1:pos+30], EX[view][pos-1:pos+30]))
    # As divergências declaradas são as que a regra manda:
    vals = domain.swap_record_values(b2b, 'atacama', accounts, queries.codes_for(b2b), MYNUM['atacama'])
    check('atacama: perna JUROS (Parte) sem Cap e SEM descrição; contraparte VCP com PU 1.0 e sem descrição',
          vals['32'].strip() == '' and vals['39'] == '0000011700000000' and vals['49'].strip() == ''
          and vals['52'].strip() == '' and vals['50'] == '0000000000000100000000')
    vals = domain.swap_record_values(cli, 'client', accounts, codes, MYNUM['client'])
    check('cliente: contraparte VCP sem descrição (só a ponta ativa a leva) e Titular = contraparte', vals['52'].strip() == '' and vals['107'] == '01')
    vals = domain.swap_record_values(b2b, 'bank', accounts, queries.codes_for(b2b), MYNUM['bank'])
    check('banco: JPM (Parte) na VCP leva a descrição (49) e PU 1.0', vals['49'].startswith('MSFT US : Indices') and vals['47'] == '0000000000000100000000')
    hdr = R._fi_build_line(commands.SWAP_FI_KEY, 'header', domain.swap_header_values('JPMORGANBM', '20260916'))
    check('header 0301', hdr == 'SWAP 00301JPMORGANBM          2026091600003')

    print('== 5. os três arquivos de prêmio 0897 ==')
    for view, deal, part in (('client', cli, 'JPMORGANBM'), ('bank', b2b, 'JPMORGANBM'), ('atacama', b2b, 'INTRAGATACAMAFDO')):
        h, reg, flow = domain.premium_values(deal, view, accounts, MYNUM[view], PREMNUM[view], part, '20260916')
        lh = R._fi_build_line(commands.PREMIUM_FI_KEY, 'header', h)
        lr = R._fi_build_line(commands.PREMIUM_FI_KEY, 'registro', reg)
        lf = R._fi_build_line(commands.PREMIUM_FI_KEY, 'fluxo', flow)
        ok = (lh, lr, lf) == PREM[view]
        check('%s: header + registro + fluxo iguais ao exemplo' % view, ok)
        if not ok:
            for a, b in zip((lh, lr, lf), PREM[view]):
                if a != b:
                    print('      got=%r\n      exp=%r' % (a, b))
        check('%s: 52 caracteres' % view, len(lh) == 52 and len(lr) == 52 and len(lf) == 52)

    print('== 6. import → arquivo-dia; re-import preserva; send grava e vira Sent ==')
    res = commands.import_upload('dt.xlsx', xlsx, datetime(2026, 9, 16), sid='A111111')
    check('dois deals importados, Recap ignorada', res['imported'] == 2 and res['ignored'] == ['Recap'])
    fp = persistence.day_path(datetime(2026, 9, 16))
    fp2, lst, idx = queries.find(cli['_id'], '2026-09-16')
    check('finder acha o deal no dia', idx is not None and os.path.normpath(fp2) == os.path.normpath(fp))
    e = lst[idx]
    check('nasce New com os quatro Meu Número', e['Status'] == 'New' and all(len(e[k]) == 10 for k in ('MyNumber', 'MyNumberMirror', 'PremiumMyNumber', 'PremiumMyNumberMirror')))
    lacuna_cli = res['missing'].get(cli['_id'], [])
    check('cliente sem conta B3 no RefData da dev cai no omnibus (sem lacuna de conta)',
          e['ClientAccount'] in ('74220005', '73760205') and not any('Client B3 Account' in x for x in lacuna_cli))
    check('SPN fora do Reference Data é LACUNA; o nome do DT fica em ClientDT',
          (e.get('ClientRefData') == 'ok' or any('Reference Data' in x for x in lacuna_cli)) and e['ClientDT'] == 'Safra')
    # A contraparte é a do Reference Data pela SPN: com cadastro, o nome vem de lá.
    ref = R._refdata_records()
    if ref:
        rec = ref[0]
        d2 = dict(cli, SPN=str(rec.get('SPN', '')), ClientAccount='', ClientTaxId='', ClientRefData='')
        commands.enrich(d2)
        check('enrich: Client = COUNTERPARTY do Reference Data pela SPN',
              d2['Client'] == str(rec.get('COUNTERPARTY', '')).strip() and d2['ClientRefData'] == 'ok')
    nums = {k: e[k] for k in ('MyNumber', 'PremiumMyNumber')}
    e['Status'] = 'Approved'; e['Maker'] = 'A111111'
    R._atomic_write_json(fp2, lst); R._daycache_forget(fp2)
    commands.import_upload('dt.xlsx', xlsx, datetime(2026, 9, 16), sid='B222222')
    _f, lst2, i2 = queries.find(cli['_id'], '2026-09-16')
    check('re-import preserva Status e Meu Número', lst2[i2]['Status'] == 'Approved' and lst2[i2]['MyNumber'] == nums['MyNumber']
          and lst2[i2]['PremiumMyNumber'] == nums['PremiumMyNumber'])
    check('sem duplicar', sum(1 for x in lst2 if (x.get('_id') or x.get('Deal')) == cli['_id']) == 1)
    if ref:
        rec = ref[0]
        d8 = commands.edit(cli['_id'], '2026-09-16', {'SPN': str(rec.get('SPN', ''))}, sid='E555555')
        check('edit com SPN nova RE-PUXA o Reference Data (nome, conta, CNPJ)',
              d8 is not None and d8['Client'] == str(rec.get('COUNTERPARTY', '')).strip() and d8['ClientRefData'] == 'ok' and d8['Status'] == 'Pending')
        commands.edit(cli['_id'], '2026-09-16', {'SPN': '281808'}, sid='E555555')
        _f, l9, i9 = queries.find(cli['_id'], '2026-09-16'); l9[i9].update(Status='Approved', Maker='A111111', ClientAccount='74220005'); R._atomic_write_json(_f, l9); R._daycache_forget(_f)
    print('== 6a. SPN de entidade nossa (le-spn) ==')
    check('le_for_spn casa por dígitos, ignorando zeros e .0',
          domain.le_for_spn([{'LE': 'ATACAMA', 'NAME': 'ATACAMA FUNDO', 'SPN': '9632845.0'}], '09632845') == {'LE': 'ATACAMA', 'NAME': 'ATACAMA FUNDO', 'SPN': '9632845.0'}
          and domain.le_for_spn([{'LE': 'ATACAMA', 'SPN': '1'}], '') is None)
    _rows_orig = R._mapping_rows
    def _rows_fake(key):
        if key == 'le-spn':
            return [{'LE': 'ATACAMA', 'NAME': 'ATACAMA MULTIMERCADO FI', 'SPN': '9632845', 'NOTES': ''}]
        return _rows_orig(key)
    R._mapping_rows = _rows_fake
    try:
        d9 = commands.edit(cli['_id'], '2026-09-16', {'SPN': '9632845'}, sid='E555555')
        check('SPN da Atacama pelo le-spn: vira LE ATACAMA / B2B, nome e conta do cadastro',
              d9 is not None and d9['LE'] == 'ATACAMA' and d9['Pair'] == 'JPM x ATACAMA' and d9['Client'] == 'ATACAMA MULTIMERCADO FI'
              and d9['ClientAccount'] == '85398005' and d9['ClientRefData'] == 'ok')
        check('B2B pelo le-spn não tem lacuna de SPN', not any('Reference Data' in x for x in domain.missing_for_send(d9, queries.codes_for(d9), accounts)))
        commands.edit(cli['_id'], '2026-09-16', {'SPN': '281808'}, sid='E555555')
        d10, _ = queries.find(cli['_id'], '2026-09-16')[1:], None
        _f, l10, i10 = queries.find(cli['_id'], '2026-09-16')
        check('SPN de cliente de volta: LE JPM, par cliente', l10[i10]['LE'] == 'JPM' and l10[i10]['Pair'] == 'JPM x CLI')
        l10[i10].update(Status='Approved', Maker='A111111', ClientAccount='74220005'); R._atomic_write_json(_f, l10); R._daycache_forget(_f)
    finally:
        R._mapping_rows = _rows_orig
    print('== 6b. dry-run + batch (duplicata → Amend) + search ==')
    dry = commands.import_upload('dt.xlsx', xlsx, datetime(2026, 9, 16), sid='A111111', dry_run=True)
    check('dry-run parseia e não grava', dry['dry_run'] and dry['imported'] == 0 and len(dry['deals']) == 2)
    dup = dict(dry['deals'][0]); dup['_replace'] = True
    n = commands.persist_deals([dup], sid='B222222')
    _f, l4, i4 = queries.find(cli['_id'], '2026-09-16')
    check('batch com _replace: Approved vira Amend, Meu Número preservado',
          n == 1 and l4[i4]['Status'] == 'Amend' and l4[i4]['MyNumber'] == nums['MyNumber'])
    d5, msg = commands.set_status(cli['_id'], '2026-09-16', 'Approved', sid='C333333')
    # Amend → Approved direto, como o New (mesa, 23/09/2026): Pending é só do Edit.
    check('Confirm em Amend vai para Approved (maker = quem confirmou)',
          d5 is not None and d5['Status'] == 'Approved' and d5['Maker'] == 'C333333' and d5['Checker'] == '')
    commands.set_status(cli['_id'], '2026-09-16', 'Pending', sid='C333333')
    d6, msg6 = commands.set_status(cli['_id'], '2026-09-16', 'Approved', sid='C333333')
    check('Pending: maker não aprova o próprio', d6 is None and 'Maker' in msg6)
    d7, _m = commands.set_status(cli['_id'], '2026-09-16', 'Approved', sid='D444444')
    check('Pending → Approved por outro usuário', d7 is not None and d7['Status'] == 'Approved' and d7['Checker'] == 'D444444')
    todos = queries.entries()
    achou = [d for d in todos if R._deal_matches(d, [{'field': 'TradeDate', 'type': 'date', 'value': '16/09/2026', 'mode': 'exact'},
                                                    {'field': 'Status', 'type': 'text', 'value': 'Success', 'mode': 'not'},
                                                    {'field': 'Pair', 'type': 'text', 'value': 'atacama'}])]
    check('search pelo contrato das irmãs (_deal_matches): Trade Date + Status ≠ Success + Pair', [d['_id'] for d in achou] == [b2b['_id']])
    _f, l5, i5 = queries.find(cli['_id'], '2026-09-16'); l5[i5]['Status'] = 'Approved'; l5[i5]['Maker'] = 'A111111'
    R._atomic_write_json(_f, l5); R._daycache_forget(_f)
    todos = queries.entries('2026-09-16')
    check('leitura do dia devolve os dois', sorted((x.get('_id') or x.get('Deal')) for x in todos) == sorted([cli['_id'], b2b['_id']]))
    # Preview do B2B: 4 arquivos (swap + prêmio × Banco e Atacama).
    _f, lstb, ib = queries.find(b2b['_id'], '2026-09-16')
    files = commands.preview(lstb[ib])
    check('preview do B2B: 4 arquivos', [(f['kind'], f['view']) for f in files] ==
          [('swap', 'bank'), ('premium', 'bank'), ('swap', 'atacama'), ('premium', 'atacama')])
    check('preview traz campo a campo com rótulo do template', files[0]['fields'][3]['field'] == 'Participante que gerou o arquivo'
          and any(x['field'] == 'Meu Número' and x['value'] == lstb[ib]['MyNumber'] for x in files[0]['fields']))
    # Send do lote com um deal sem lacuna e outro com: recusa tudo.
    lstb[ib]['Notional'] = ''
    R._atomic_write_json(_f, lstb); R._daycache_forget(_f)
    try:
        commands.send([{'deal_id': cli['_id'], 'trade_date': '2026-09-16'}, {'deal_id': b2b['_id'], 'trade_date': '2026-09-16'}])
        check('lote com lacuna recusa tudo', False)
    except ValueError as exc:
        check('lote com lacuna recusa tudo dizendo qual', 'Notional' in str(exc) and b2b['_id'] in str(exc))
    check('nada foi escrito', not os.path.isdir(R.CONECTA_NEW_PATH) or not os.listdir(R.CONECTA_NEW_PATH))
    lstb[ib]['Notional'] = '346000.00'
    for x in lstb:
        x['ClientRefData'] = 'ok'          # idem: a SPN do DT não está no RefData da dev
    R._atomic_write_json(_f, lstb); R._daycache_forget(_f)
    out = commands.send([{'deal_id': cli['_id'], 'trade_date': '2026-09-16'}, {'deal_id': b2b['_id'], 'trade_date': '2026-09-16'}], sid='C333333')
    nomes = sorted(x['filename'] for x in out['files'])
    check('seis arquivos: SWAP e PREMIO × Cliente/Banco/Atacama, com a LOB do deal no nome',
          nomes == ['PREMIO_SWAP_EDG_ATACAMA.txt', 'PREMIO_SWAP_EDG_BANCO.txt', 'PREMIO_SWAP_EDG_CLIENTE.txt',
                    'SWAP_EDG_ATACAMA.txt', 'SWAP_EDG_BANCO.txt', 'SWAP_EDG_CLIENTE.txt'])
    txt = io.open(os.path.join(R.CONECTA_NEW_PATH, 'SWAP_EDG_CLIENTE.txt'), encoding='cp1252').read().split('\n')
    check('SWAP_CLIENTE: header + 1 registro de 1927', len(txt) == 2 and txt[0].startswith('SWAP 00301JPMORGANBM') and len(txt[1]) == 1927)
    # O Conecta conta BYTES: o travessão da denominação VCP é 1 byte em
    # cp1252 e 3 em utf-8 — gravado em utf-8 o registro do Banco (o que leva
    # a Descrição) chegava à B3 com 1935 e era recusado.
    raw = io.open(os.path.join(R.CONECTA_NEW_PATH, 'SWAP_EDG_BANCO.txt'), 'rb').read().split(b'\n')
    check('SWAP_BANCO: registro de 1927 BYTES (cp1252) com o travessão em um byte',
          len(raw) == 2 and len(raw[1]) == 1927 and b'\x96' in raw[1] and b'\xe2\x80\x93' not in raw[1])
    ptxt = io.open(os.path.join(R.CONECTA_NEW_PATH, 'PREMIO_SWAP_EDG_ATACAMA.txt'), encoding='cp1252').read().split('\n')
    check('PREMIO_ATACAMA: header INTRAGATACAMAFDO + registro + fluxo', len(ptxt) == 3 and 'INTRAGATACAMAFDO' in ptxt[0] and ptxt[1][50:52] == '01')
    _f, lst3, i3 = queries.find(cli['_id'], '2026-09-16')
    check('vira Sent com os arquivos anotados', lst3[i3]['Status'] == 'Sent' and 'SWAP_EDG_CLIENTE.txt' in lst3[i3]['SentFiles'])
    try:
        commands.send([{'deal_id': cli['_id'], 'trade_date': '2026-09-16'}])
        check('Sent não reenvia', False)
    except ValueError as exc:
        check('Sent não reenvia', 'status Sent' in str(exc))
    print('== 6b2. Mapping B3 ID pelo arquivo de RETORNO (§554) ==')
    # O retorno ecoa o 0301 que foi: monta-se a partir dos arquivos que o Send
    # acabou de gravar. Cliente e B2B no mesmo arquivo; o B2B volta DUAS vezes
    # (Banco e Atacama lançam o mesmo contrato).
    ret_dir = os.path.join(tmp, 'return'); os.makedirs(ret_dir, exist_ok=True)
    R.RETURN_PATH = ret_dir
    def _rec(nome):
        return io.open(os.path.join(R.CONECTA_NEW_PATH, nome), encoding='cp1252').read().split('\n')[1]
    ret_linhas = ['HEADER DO RETORNO',
                  '00000000000001;26I05774449;2026091738383230;EXECUCAO OK;' + _rec('SWAP_EDG_CLIENTE.txt'),
                  '00000000000002;26I05760675;2026091738383231;EXECUCAO OK;' + _rec('SWAP_EDG_BANCO.txt'),
                  '00000000000003;26I05760675;2026091738383232;EXECUCAO OK;' + _rec('SWAP_EDG_ATACAMA.txt')]
    io.open(os.path.join(ret_dir, 'RET_SWAP.txt'), 'w', encoding='cp1252').write('\n'.join(ret_linhas))
    io.open(os.path.join(ret_dir, 'RET_TER.txt'), 'w', encoding='cp1252').write('x;y;z;EXECUCAO OK;TER outro sistema')
    p0 = domain.parse_return_line(ret_linhas[2])
    check('o eco do 0301 dá Meu Número, contas, datas e valor base pelas posições do layout',
          p0 and p0['my_number'] == lstb[ib]['MyNumber'] and p0['parte'] == '73760009' and p0['contraparte'] == '85398005'
          and p0['start'] == '20260915' and p0['maturity'] == '20270607' and p0['notional_cents'] == 34600000)
    check('header, linha de outro sistema e 0897 não são eco de 0301',
          domain.parse_return_line(ret_linhas[0]) is None and domain.parse_return_line('a;b;c;EXECUCAO OK;TER x') is None
          and domain.parse_return_line('a;b;c;EXECUCAO OK;SWAP 10897' + '0' * 200) is None)
    _disparos = []
    _bm = commands.b3_mapped
    commands.b3_mapped = lambda d: _disparos.append(d.get('B3ID'))
    try:
        res_map = commands.map_b3('2026-09-16', sid='G777777')
    finally:
        commands.b3_mapped = _bm
    _f, lstm, _ = queries.find(cli['_id'], '2026-09-16')
    por_id = {e['_id']: e for e in lstm}
    check('contra cliente: o B3 ID do eco com o Meu Número da visão do cliente, Success',
          por_id[cli['_id']]['B3ID'] == '26I05774449' and por_id[cli['_id']]['Status'] == 'Success'
          and por_id[cli['_id']]['MappedBy'] == 'G777777')
    check('B2B: o B3 ID do contrato com o Atacama, Success',
          por_id[b2b['_id']]['B3ID'] == '26I05760675' and por_id[b2b['_id']]['Status'] == 'Success')
    check('o resultado traz os dois, gravados, e o disparo pós-mapeamento roda para os dois',
          sorted(r['b3_id'] for r in res_map) == ['26I05760675', '26I05774449'] and all(r['saved'] for r in res_map)
          and sorted(_disparos) == ['26I05760675', '26I05774449'])
    check('o retorno todo consumido é apagado; o arquivo de outro sistema fica',
          not os.path.exists(os.path.join(ret_dir, 'RET_SWAP.txt')) and os.path.exists(os.path.join(ret_dir, 'RET_TER.txt')))
    commands.b3_mapped = lambda d: None
    try:
        d_ed = commands.edit(cli['_id'], '2026-09-16', {'Observations': 'ajuste depois do registro'}, sid='H888888')
    finally:
        commands.b3_mapped = _bm
    check('EDITAR a linha já mapeada não a tira de Success (era o Approved com B3 ID)',
          d_ed['Status'] == 'Success' and d_ed['B3ID'] == '26I05774449')
    _dconf, _merr = commands.set_status(cli['_id'], '2026-09-16', 'Approved', sid='H888888')
    check('e o Confirm não leva a linha mapeada a Approved', _dconf is None)
    # Linha presa antes da correção (B3 ID + Approved) → o Mapping cura.
    _f, lsts, isx = queries.find(cli['_id'], '2026-09-16')
    lsts[isx]['Status'] = 'Approved'; R._atomic_write_json(_f, lsts); R._daycache_forget(_f)
    commands.b3_mapped = lambda d: _disparos.append('cura:' + d.get('B3ID'))
    try:
        res_cura = commands.map_b3('2026-09-16', sid='G777777')
    finally:
        commands.b3_mapped = _bm
    check('B3 ID com Approved (a linha presa) volta a Success no Mapping, e o disparo roda de novo',
          [r['status'] for r in res_cura] == ['Success'] and queries.find(cli['_id'], '2026-09-16')[1][isx]['Status'] == 'Success'
          and 'cura:26I05774449' in _disparos)
    # Eco com outro status → Error com o texto da B3; enviado sem retorno → Error.
    cli_err = dict(cli, B3ID='', Status='Sent')
    io.open(os.path.join(ret_dir, 'RET_ERR.txt'), 'w', encoding='cp1252').write(
        '00000000000009;;2026091738383299;CAMPO 126 CONTEUDO INVALIDO;' + _rec('SWAP_EDG_CLIENTE.txt'))
    _f, lste, iex = queries.find(cli['_id'], '2026-09-16')
    lste[iex].update(B3ID='', Status='Sent'); R._atomic_write_json(_f, lste); R._daycache_forget(_f)
    res_err = commands.map_b3('2026-09-16', sid='G777777')
    e_err = queries.find(cli['_id'], '2026-09-16')[1][iex]
    check('eco com status da B3 ≠ EXECUCAO OK → Error com o texto da B3, sem B3 ID',
          e_err['Status'] == 'Error' and e_err['MappingError'] == 'CAMPO 126 CONTEUDO INVALIDO' and not e_err['B3ID'])
    check('   e o retorno processado (mesmo com erro) sai da pasta', not os.path.exists(os.path.join(ret_dir, 'RET_ERR.txt')))
    lste = queries.find(cli['_id'], '2026-09-16')[1]; lste[iex].update(Status='Sent'); lste[iex].pop('MappingError', None)
    R._atomic_write_json(_f, lste); R._daycache_forget(_f)
    commands.map_b3('2026-09-16', sid='G777777')
    e_nf = queries.find(cli['_id'], '2026-09-16')[1][iex]
    check('enviado e ausente do retorno → Error dizendo que não achou',
          e_nf['Status'] == 'Error' and 'not found' in e_nf['MappingError'])
    # Sem Meu Número: pelas características, só com candidato único.
    alheio = dict(queries.find(cli['_id'], '2026-09-16')[1][iex], MyNumber='0000000000')
    p_cli = domain.parse_return_line('1;26I05774449;x;EXECUCAO OK;' + _rec('SWAP_EDG_CLIENTE.txt'))
    check('sem o Meu Número, conta do cliente + datas + valor casam (candidato único)',
          domain.match_return(alheio, [p_cli], [alheio]) is p_cli)
    check('   e dois deals com a mesma assinatura não chutam',
          domain.match_return(alheio, [p_cli], [alheio, dict(alheio)]) is None)
    # Reimport com "substituir" de uma linha mapeada não a rebaixa a Amend.
    _f, lstr, irx = queries.find(cli['_id'], '2026-09-16')
    lstr[irx].update(B3ID='26I05774449', Status='Success'); R._atomic_write_json(_f, lstr); R._daycache_forget(_f)
    commands.persist_deals([dict(cli, _replace=True)], sid='B222222')
    e_re = queries.find(cli['_id'], '2026-09-16')[1][irx]
    check('reimport SUBSTITUINDO a linha mapeada mantém Success e o B3 ID', e_re['Status'] == 'Success' and e_re['B3ID'] == '26I05774449')

    apagados, nao = commands.delete([{'deal_id': cli['_id'], 'trade_date': '2026-09-16'}])
    check('delete apaga do arquivo', apagados == 1 and not nao and queries.find(cli['_id'], '2026-09-16')[2] is None)

    print('== 6c. Economic Affirmation (IF, D0) ==')
    from apps.pages import otc_emails
    hoje = datetime.now().strftime('%d/%m/%Y')
    fi = dict(cli, TradeDate=datetime.now().strftime('%Y-%m-%d'), ClientAccount='74220005', ClientRefData='ok', Client='BANCO SAFRA S.A.')
    omni = dict(fi, ClientAccount='73760205', ClientAccountNote='omnibus')
    drafts = otc_emails.build_swap_bullet_affirmation_emails([fi, omni, dict(b2b, TradeDate=fi['TradeDate']), dict(fi, TradeDate='2026-01-02')])
    check('um rascunho: só a IF com conta própria, na data; omnibus, B2B e outro dia ficam fora', len(drafts) == 1)
    dr = drafts[0]
    check('assunto no molde do e-mail da mesa', dr['subject'] == 'Confirmação da(s) Operação(ões) Fechada(s) em %s - BANCO SAFRA S.A. - SWAP' % hoje)
    h = dr['html']
    check('corpo: contas CETIP das duas pontas e o EDG no código identificador',
          'Conta CETIP BANCO SAFRA S.A.' in h and '74220.00-5' in h and '73760.00-9' in h and 'incluir &quot;EDG&quot;' in h.replace('"', '&quot;'))
    check('corpo: o Deal Ticket (blocos e valores do DT)',
          'DEAL TICKET SWAP VCP' in h and 'Curva Vanilla' in h and 'MSFT US' in h and 'BRL 346.000,00' in h
          and '15-set-2026' in h and '07-jun-2027' in h and '117,00% Spot' in h and '23.355,00' in h and 'Microsoft Corporation' in h)
    check('sem logo do OTC Tracker (só a marca do banco por CID)', 'otc' not in h.lower().replace('otc_derivatives', '').replace('brazil.otc', '') and 'cid:jpmwordmark' in h)
    check('casca larga (960) e nome do Reference Data dentro do DT', 'width:960px' in h and '>BANCO SAFRA S.A.<' in h)
    check('a SPN (identificador interno) não sai no e-mail', '>SPN<' not in h and '281808' not in h)
    _idx_orig = otc_emails._build_cpdetails_index
    otc_emails._build_cpdetails_index = lambda: {'281808': {'CONTACTS': [
        {'email': 'conf@safra.com.br', 'rules': ['Confirmation Letter']},
        {'email': 'liq@safra.com.br', 'rules': ['Settlement']},
        {'email': 'conf2@safra.com.br', 'rules': ['Only for Confirmation']}]}}
    try:
        dr2 = otc_emails.build_swap_bullet_affirmation_emails([dict(fi, MaturityDate='1931-08-18', QuoteDate='1931-08-15')])[0]
        check('destinatários = contatos de Confirmation da contraparte (Counterparty Details pela SPN)',
              dr2['to'] == 'conf@safra.com.br; conf2@safra.com.br')
        check('data 1931 de linha antiga sai como 2031 no e-mail', '18-ago-2031' in dr2['html'] and '15-ago-2031' in dr2['html'])
    finally:
        otc_emails._build_cpdetails_index = _idx_orig
    print('== 6d. B3 ID mapeado (§481): Intrag Swap no B2B; Pending Confirmation + esteira contra cliente ==')
    # A linha da Intrag Swap é a visão da ATACAMA (Parte Atacama com a perna
    # dela, Contraparte Banco) — os 36 campos no exemplo da mesa.
    codes_b = queries.codes_for(b2b)
    ent = domain.intrag_swap_entry(dict(b2b, B3ID='26F04329902'), codes_b)
    linha = ';'.join(ent[k] for k in domain.INTRAG_SWAP_FIELDS)
    esperado_intrag = ('INTRAGJP633;26F04329902;2026-09-15;2027-06-07;346000;Opção de Arrependimento;Atacama;Vanilla;'
                       'Banco JP Morgan;VCP;Sim;2026-09-16;Banco JP Morgan;23355;100;C99;PRÉ FIXADO BRL;;Positivo;0;;;;;;'
                       '100;C00;VCP;10615;Positivo;0;117;;2027-06-04;Quanto (ausencia variação cambial); MSFT US;100.00% Spot')
    check('a linha da Intrag Swap (36 campos, visão da Atacama) é a do exemplo da mesa', linha == esperado_intrag)
    if linha != esperado_intrag:
        print('      got=%r' % linha)
    _tpl = io.open(os.path.join(os.path.dirname(__file__), '..', '..', 'apps', 'templates', 'pages', 'intrag-swap.html'), encoding='utf-8').read()
    _m = re.search(r'var ENTRY_FIELDS = \[(.*?)\];', _tpl, re.S)
    _page_fields = tuple(re.findall(r"'([a-z0-9_]+)'", _m.group(1))) if _m else ()
    check('os 36 campos são os da página /intrag-swap, na mesma ordem', _page_fields == domain.INTRAG_SWAP_FIELDS)
    check('SWAP com Opção de Arrependimento; SWAP CORPORATE sem',
          domain.confirmation_source(cli) == 'SWAP' and domain.confirmation_source(dict(cli, Functionality='N/A')) == 'SWAP CORPORATE')
    from apps.pages.features.intrag.infra import persistence as IP
    from apps.pages.features.intrag import queries as IQ
    IP.INTRAG_SWAP_CACHE_DIR = os.path.join(tmp, 'intrag-swap')
    d11 = commands.edit(b2b['_id'], '2026-09-16', {'B3ID': '26F04329902'}, sid='F666666')
    check('B3 ID novo é MAPEAMENTO: Status Success (não Pending), com quem mapeou anotado',
          d11 is not None and d11['Status'] == 'Success' and d11['MappedBy'] == 'F666666')
    fp_i, ents_i, ii = IQ._find_intrag_swap_entry(b2b['_id'], '2026-09-15')
    check('B2B com B3 ID → linha da Intrag Swap no arquivo-dia da Data Início, New, carteira da Atacama',
          ii is not None and ents_i[ii]['b3_id'] == '26F04329902' and ents_i[ii]['carteira'] == 'INTRAGJP633'
          and ents_i[ii]['status'] == 'New' and fp_i.endswith('20260915_intrag_swap.json') and ents_i[ii]['_client'] == b2b['Client'])
    # Contra cliente: a mesma porta das outras páginas (o `_pc_save_from_deal`,
    # que pula perna interna e chama a esteira).
    calls = []
    _pc_orig = R._pc_save_from_deal
    R._pc_save_from_deal = lambda deal, pt, pending_status=None, trade_number=None, source=None: calls.append((deal, pt, pending_status, trade_number, source))
    try:
        commands.persist_deals([dict(cli, Status='Sent', ClientTaxId='58160789000128', ClientAccount='74220005')], sid='A111111')
        d12 = commands.edit(cli['_id'], '2026-09-16', {'B3ID': '26F04329901'}, sid='F666666')
    finally:
        R._pc_save_from_deal = _pc_orig
    check('cliente com B3 ID → Pending Confirmation: Product Type/source SWAP, Pending OTC, chave = B3 ID, LOB EDG',
          d12 is not None and d12['Status'] == 'Success' and len(calls) == 1 and calls[0][1] == 'SWAP' and calls[0][2] == 'Pending OTC'
          and calls[0][3] == '26F04329901' and calls[0][4] == 'SWAP' and calls[0][0]['Deal'] == '26F04329901'
          and calls[0][0]['SettlementDate'] == '2027-06-07' and calls[0][0]['LOB'] == 'EDG' and calls[0][0]['TaxID'] == '58160789000128')
    cdeal = calls[0][0]
    from apps.pages import manual_conf as MC
    _fr, _up = MC.find_row, MC.upsert_row
    rows_mc = []
    MC.find_row = lambda k: None
    MC.upsert_row = lambda row: rows_mc.append(row)
    try:
        R._mc_save_from_deal(cdeal, 'SWAP', trade_number='26F04329901')
    finally:
        MC.find_row, MC.upsert_row = _fr, _up
    row_mc = rows_mc[0] if rows_mc else {}
    check('esteira: Produto SWAP, LOB do deal (EDG), Trade ID/Cetip ID = B3 ID, ativo = curva VCP, notional BRL, vencimento',
          {'SWAP', 'SWAP CORPORATE'} <= R._MC_CONFIRMATION_SOURCES and row_mc.get('Produto') == 'SWAP' and row_mc.get('LOB') == 'EDG'
          and row_mc.get('Trade ID') == '26F04329901' and row_mc.get('Cetip ID') == '26F04329901' and row_mc.get('Moeda') == 'MSFT US'
          and row_mc.get('Notional Amount CCY') == 'BRL 346000.00' and row_mc.get('Data de vencimento') == '07/06/2027'
          and R._lob_for_source('NDF COMM') == 'COMMODITY' and R._mc_conf_trade_keys([(cdeal, None)], 'swap-edg') == ['26F04329901'])

    print('== 6e. a confirmação do swap: segregação, Cupom Limpo (Close → cotação; Spot → em branco), XML só em BRL ==')
    ref16 = datetime(2026, 9, 16)
    grupos, _st, _tot = R._conf_swap_groups(ref16)
    check('um grupo contraparte × curva VCP × opcao-arrependimento, com o B3 ID; o B2B fica fora (é Intrag)',
          len(grupos) == 1 and grupos[0]['mercadoria'] == 'MSFT US' and grupos[0]['family'] == 'opcao-arrependimento'
          and '26F04329901' in grupos[0]['trades'] and grupos[0]['eligible'] == 1)
    g = grupos[0]
    picked = R._conf_pick_swap(ref16, g['acronym'], g['mercadoria'], 'opcao-arrependimento')
    check('elegível = a operação com B3 ID', len(picked) == 1 and picked[0][0]['B3_ID'] == '26F04329901')
    fake_close = lambda label, on, w: 500.25 if (label == 'MSFT US' and on == datetime(2026, 9, 15).date()) else None
    w1 = []
    conf = R._conf_swap_conf(picked[0][0], ref16, g['acronym'], g['mercadoria'], 'opcao-arrependimento', w1, fetch_close=fake_close)
    check('Close 15-Sep-26: Preço Inicial = fechamento do dia, Strike = 100% dele, Cap 117% vira preço; Fator Equities na Parte B; Parte A Exponencial 252',
          conf['b_preco_inicial'] == '500,2500' and conf['b_strike'] == '500,2500' and conf['b_lim_alta'] == '585,2925'
          and conf['b_fator_equities'] == 'Aplicável' and conf['b_ativo'] == 'MSFT US - Microsoft Corporation'
          and conf['b_dt_obs'] == '04/06/2027' and conf['b_bolsa'] == 'Bloomberg' and conf['b_repasse'] == 'Não'
          and conf['a_forma_juros'] == 'Exponencial 252' and conf['a_fator_equities'] == 'Não aplicável' and conf['equities_side'] == 'B')
    du = R._anbima_biz_diff(datetime(2026, 9, 15), datetime(2027, 6, 7))
    check('3.1 e 3.2: nº = B3 ID, Valor Base, datas, prêmio (Parte B paga e tem o direito), dias corridos/úteis, juros pré 0% × Não Aplicável',
          conf['num_conf'] == '26F04329901' and conf['valor_base'] == '346.000,00' and conf['dt_efetiva'] == '15/09/2026'
          and conf['dt_venc'] == '07/06/2027' and conf['premio'] == 'R$ 23.355,00' and conf['devedor_premio'] == 'Parte B'
          and conf['dt_premio'] == '16/09/2026' and conf['arrependimento'] == 'Sim' and conf['parte_arrependimento'] == 'Parte B'
          and conf['dias_corridos'] == '265' and conf['dias_uteis'] == str(du) and conf['taxa_amort'] == '100%'
          and conf['juros_pre_a'] == '0%' and conf['juros_pre_b'] == 'Não Aplicável'
          and conf['parteb_cnpj'] == '58.160.789/0001-28' and conf['partea_nome'] == 'BANCO J.P. MORGAN S.A.')
    spot = dict(picked[0][0]); spot['_conf_cupom'] = {'mode': 'spot', 'date': '', 'pct': 100.0}
    w2 = []
    c2 = R._conf_swap_conf(spot, ref16, g['acronym'], g['mercadoria'], 'opcao-arrependimento', w2, fetch_close=fake_close)
    check('Spot: Preço Inicial e Strike em BRANCO, avisando que são obrigatórios',
          c2['b_preco_inicial'] == '' and c2['b_strike'] == '' and any('obrigat' in x for x in w2))
    w3 = []
    c3 = R._conf_swap_conf(picked[0][0], ref16, g['acronym'], g['mercadoria'], 'opcao-arrependimento', w3, fetch_close=lambda *a: None)
    check('Close sem cotação: em branco e avisando', c3['b_preco_inicial'] == '' and c3['b_strike'] == '' and any('em branco' in x for x in w3))
    from flask import render_template as _rt
    with app.test_request_context():
        html_doc = _rt('confirmations/swap-edg-opcao-arrependimento.html', conf=conf, doc_only=True)
        html_edit = _rt('confirmations/swap-edg-opcao-arrependimento.html', conf=conf)
    check('documento (doc_only): nº, Partes, Valor Base, preços e datas no lugar do exemplo do Word; sem painel; utf-8',
          '26F04329901' in html_doc and '346.000,00' in html_doc and html_doc.count('500,2500') == 2 and 'Microsoft Corporation' in html_doc
          and '58.160.789/0001-28' in html_doc and 'id="editor-panel"' not in html_doc and 'BANCO SAFRA' not in html_doc
          and '26I04036287' not in html_doc and '6.180.800,00' not in html_doc and '7.591,7000' not in html_doc and 'charset=utf-8' in html_doc)
    check('editor: painel + script, obrigatórios marcados na perna do Fator Equities, save na rota do swap',
          'id="editor-panel"' in html_edit and 'inp_b_strike' in html_edit and 'Strike (obrigatório)' in html_edit
          and '/api/confirmation/swap-edg/save' in html_edit)
    check('as fórmulas do Word (OMML) têm texto linear fora do Word — a imagem que não veio não fica quebrada',
          'v:imagedata' not in html_doc and 'Fator Equities= 1+' in html_doc)
    numero, xml, xw = R._conf_swap_xml(picked, 'MSFT US', ref16)
    check('XML: tipoOperacao SWAP, numeroContrato = B3 ID, valor = notional em BRL, moeda e valor estrangeiro VAZIOS',
          numero == '26F04329901' and '<tipoOperacao>SWAP</tipoOperacao>' in xml and '<valor>346000.00</valor>' in xml
          and '<moedaEstrangeira></moedaEstrangeira>' in xml and '<valorEstrangeiro></valorEstrangeiro>' in xml
          and '<cnpjCliente>58160789000128</cnpjCliente>' in xml and '<dataVencimento>20270607</dataVencimento>' in xml)
    url, motivo = R._mc_generate_url({'Produto': 'SWAP', 'LOB': 'EDG', 'Data Operação': '16/09/2026',
                                      'Trade ID': '26F04329901', 'Cliente': g['acronym'], 'Legal Entity': ''}, ['26F04329901'])
    check('Generate do Monitor abre o editor de swap do grupo', url.startswith('/confirmation/swap-edg/opcao-arrependimento?date=2026-09-16'))
    check('a família sem Opção de Arrependimento (corporate) não tem template — o Generate diz isso',
          R._conf_swap_family({'_conf_kind': 'SWAP CORPORATE'}, None) == 'corporate' and 'corporate' not in R._CONF_SWAP_FAMILY_TEMPLATES)
    rules_c = {str(r.rule) for r in app.url_map.iter_rules()}
    check('rotas do editor de swap registradas', {'/confirmation/swap-edg/<family>', '/api/confirmation/swap-edg/save',
                                                 '/api/confirmation/swap-edg/pdf', '/confirmation/swap-edg/validate',
                                                 '/api/confirmation/swap-edg/validate'} <= rules_c)

    print('== 6f. o editor e o save pela API (test client): Word + PDF + XML na pasta SWAP ==')
    import json as _json
    from urllib.parse import quote as _q
    from datetime import timezone as _tz, timedelta as _td
    from apps.pages.platform import confirmations as PC
    _real = (PC._conf_swap_close, PC._conf_cgd_lookup, PC._conf_state_path, R._ei_resolve_client_dir,
             R._conf_pc_set_fepweb, R._mc_stamp_generated, R._create_notification)
    EI = os.path.join(tmp, 'ei'); stamped = []; fep = []
    PC._conf_swap_close = fake_close
    PC._conf_cgd_lookup = lambda first: '08 de agosto de 2008'
    PC._conf_state_path = lambda ref, product='ndf-comm': os.path.join(tmp, 'conf-state-%s-%s.json' % (product, ref.strftime('%Y%m%d')))
    R._ei_resolve_client_dir = lambda name, create=False: os.path.join(EI, name)
    R._conf_pc_set_fepweb = lambda deals, num: fep.append((list(deals), num)) or 1
    R._mc_stamp_generated = lambda picked, product, link='': stamped.append((product, len(picked), link))
    R._create_notification = lambda *a, **k: None
    try:
        cl = app.test_client()
        with cl.session_transaction() as ss:
            ss['authenticated'] = True; ss['user_sid'] = 'T000000'; ss['user_name'] = 'T'
            ss['session_expires_at'] = (datetime.now(tz=_tz.utc) + _td(hours=8)).isoformat()
        page = cl.get('/confirmation/swap-edg/opcao-arrependimento?date=2026-09-16&acronym=%s&mercadoria=MSFT%%20US' % _q(g['acronym']))
        hv = page.data.decode('utf-8')
        check('GET do editor: 200 com o painel e o fechamento do Close', page.status_code == 200 and 'id="editor-panel"' in hv and '500,2500' in hv)
        confj = _json.loads(re.search(r'id="conf-data">(.*?)</script>', hv, re.S).group(1))
        fields = {k: confj.get(k, '') for k in R._CONF_SWAP_FIELDS}
        body = {'family': 'opcao-arrependimento', 'date': '2026-09-16', 'acronym': g['acronym'], 'mercadoria': 'MSFT US', 'equities_side': 'b'}
        r0 = cl.post('/api/confirmation/swap-edg/save', json=dict(body, fields=dict(fields, b_preco_inicial='', b_strike=''))).get_json()
        check('save recusa Preço Inicial/Strike em branco na perna do Fator Equities (missing_price)',
              not r0.get('success') and r0.get('error') == 'missing_price')
        res = cl.post('/api/confirmation/swap-edg/save', json=dict(body, fields=fields)).get_json()
        files = res.get('files') or []
        check('salvou .doc + .pdf + .xml na pasta SWAP do cliente', res.get('success') and len(files) == 3
              and all(os.sep + 'SWAP' + os.sep in f for f in files) and all(os.path.isfile(f) for f in files)
              and files[0].endswith('.doc') and files[1].endswith('.pdf') and files[2].endswith('.xml'))
        if not res.get('success'):
            print('      ', res)
        xml_txt = io.open(files[2], encoding='utf-8').read() if len(files) == 3 else ''
        check('o XML gravado: SWAP, B3 ID, valor em BRL, estrangeiro vazio; FepWeb ID e esteira carimbados pelo B3 ID',
              '<tipoOperacao>SWAP</tipoOperacao>' in xml_txt and '<valor>346000.00</valor>' in xml_txt and '<valorEstrangeiro></valorEstrangeiro>' in xml_txt
              and fep == [(['26F04329901'], '26F04329901')] and stamped and stamped[-1][0] == 'swap-edg' and stamped[-1][1] == 1)
        doc = io.open(files[0], encoding='utf-8').read() if files else ''
        check('o .doc é o documento sem o painel, com os valores e o nome da contraparte na assinatura',
              'id="editor-panel"' not in doc and '500,2500' in doc and '346.000,00' in doc and 'out_parteb_nome_assin">' + g['acronym'] in doc)
        check('o PDF saiu do mesmo HTML (não vazio)', len(files) == 3 and os.path.getsize(files[1]) > 20000)
        # Os quadrados das Barreiras são FORMCHECKBOX do Word (sem glifo fora
        # dele): o documento mostra ☒/☐ a quem não desenha campos, e o PDF os
        # escreve na fonte de símbolos da máquina (ou [X]/[  ] sem ela).
        check('o .doc traz os 16 quadrados: 8 marcados (os Não Aplicável) e 8 vazios, escondidos do Word',
              doc.count('\u2612') == 8 and doc.count('\u2610') == 8 and doc.count('<![if !supportFields]>') == 16)
        try:
            from pypdf import PdfReader
            pdf_txt = ''.join((pg.extract_text() or '') for pg in PdfReader(files[1]).pages)
        except Exception as exc:                             # noqa: BLE001
            pdf_txt = ''
            print('      (pypdf indisponível: %s)' % exc)
        check('o PDF desenha os quadrados (fonte de símbolos) ou escreve [X]/[  ]',
              ('\u2612' in pdf_txt and '\u2610' in pdf_txt) or ('[X]' in pdf_txt and '[  ]' in pdf_txt))
        val = cl.get(res.get('validate_url') or '/x')
        check('a janela de validação abre', val.status_code == 200)
    finally:
        (PC._conf_swap_close, PC._conf_cgd_lookup, PC._conf_state_path, R._ei_resolve_client_dir,
         R._conf_pc_set_fepweb, R._mc_stamp_generated, R._create_notification) = _real
    commands.delete([{'deal_id': cli['_id'], 'trade_date': '2026-09-16'}])

    print('== 7. templates e cadastros ==')
    tpl = R._fi_tpl_cached('swap-registro-premio')
    check('swap-registro-premio na biblioteca, 3 blocos (6+9+4), ligado à página',
          bool(tpl) and tpl.get('status') == 'library' and [len(b['fields']) for b in tpl['blocks']] == [6, 9, 4]
          and any(p.get('url') == '/new_deals-swap-bullet' for p in tpl.get('linked_pages', [])))
    tpl3 = R._fi_tpl_cached('swap-pagamento-final-v3')
    check('swap-pagamento-final-v3 ligado à página com as colunas',
          bool(tpl3) and any(p.get('url') == '/new_deals-swap-bullet' and 'VCP Text' in (p.get('columns') or []) for p in tpl3.get('linked_pages', [])))
    check('cadastro swap-bullet-curve existe', 'swap-bullet-curve' in R._MAPPING_DEFS and any(r.get('B3 CODE') == 'C99' for r in R._mapping_rows('swap-bullet-curve')))
    check('swap-code-labels ganhou a Adesão (upgrade)', any(str(r.get('FIELD')) == 'Adesão' and r.get('LABEL') == 'CGD' for r in R._mapping_rows('swap-code-labels')))
    from apps.pages.features.deals_monitor import domain as ND
    # O card `Swap Bullet` deixou de existir (21/09/2026): a pasta gravada
    # alimenta Swap Equities e Swap CEM, e quem decide e a LOB da linha.
    check('a pasta gravada alimenta os dois cards de swap do Monitor, pela LOB',
          all(any(c['key'] == k and 'Swap/Bullet' in c['dirs'] and c.get('lob') == lob for c in ND._NDM_CARDS)
              for k, lob in (('swap-equities', 'EDG'), ('swap-cem', 'CEM')))
          and ND._ndm_bucket('Swap/Bullet', {'LOB': 'EDG'}) in ND.card_buckets(
              next(c for c in ND._NDM_CARDS if c['key'] == 'swap-equities'))
          and not any(c['key'] == 'swap-bullet' for c in ND._NDM_CARDS))
    check('B2B do Swap Bullet conta como ATA no Monitor', ND._ndm_deal_le('Swap/Bullet', {'Client': 'Atacama'}) == 'ATA'
          and ND._ndm_deal_le('Swap/Bullet', {'Client': 'Safra'}) == 'JPM')
    rules = {str(r.rule) for r in app.url_map.iter_rules()}
    check('rotas registradas', {'/api/new-deals/swap-bullet', '/api/new-deals/swap-bullet/import-file',
                                '/api/new-deals/swap-bullet/preview', '/api/new-deals/swap-bullet/send-conecta'} <= rules)

    print()
    if FALHAS:
        print('FALHOU: ' + '; '.join(FALHAS))
        return 1
    print('TUDO OK')
    return 0


if __name__ == '__main__':
    sys.exit(main())
