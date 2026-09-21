# -*- coding: utf-8 -*-
"""check_tools_calculators.py — NDF Calculator, Unwind NDF Calculator e Option
Calculator (Tools), e a ordem ALFABETICA do menu.

O que este teste prende:

  1. o menu de Tools esta em ordem alfabetica pelo rotulo em ingles, com as
     tres paginas novas, e cada link tem rota (link sem rota e 404 calado);
  2. NDF no vencimento: o sinal e a POSICAO do banco, o IR de 0,005% so quando
     o banco PAGA, isento nao paga, e o nocional fixo em reais divide pela taxa;
  3. recompra de NDF: a MESMA formula que a pagina de Unwinds confere — as tres
     operacoes REAIS do check_unwind_notification fecham aqui tambem, e o motor
     puro bate com `unwinds.domain.conferir_apuracao` (duas copias da conta que
     divergem sao dois numeros para a mesma recompra);
  4. o saldo da recompra sao TRES parcelas, e o total diz Zero;
  5. opcao: call/put, titular/lancador, a media da asiatica, o premio separado
     do exercicio, fora do dinheiro nao exerce;
  6. as tres paginas de ponta a ponta: abrem, calculam, o erro sai em frase (e
     nao em 500), e o fixing em branco busca a PTAX dizendo de que dia e.
"""
import io
import os
import re
import sys
import tempfile
from datetime import date, datetime, timedelta, timezone

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
sys.path.insert(0, ROOT)
os.environ.setdefault('OTC_DISABLE_SCHEDULERS', '1')
if os.name != 'nt' and not os.environ.get('OTC_SHARED_DRIVE_ROOT'):
    os.environ['OTC_SHARED_DRIVE_ROOT'] = tempfile.mkdtemp(prefix='otc-share-')

FALHAS = []


def check(nome, got, exp=True):
    ok = got == exp
    print(('  ok  ' if ok else ' FAIL ') + nome + ('' if ok else '\n        got=%r\n        exp=%r' % (got, exp)))
    if not ok:
        FALHAS.append(nome)


def main():
    from run import app
    from apps.pages.precificador import derivativos as D
    from apps.pages.features.tools import queries
    from apps.pages.features.unwinds import domain as unw

    print('== 1. o menu de Tools: ordem alfabetica e link com rota ==')
    menu = io.open(os.path.join(ROOT, 'apps', 'templates', 'partials', 'sidenav.html'),
                   encoding='utf-8').read()
    # Tools tem DOIS niveis: o grupo **Calculators** (as calculadoras) e as
    # irmas dele. Cada nivel e alfabetico por conta propria.
    tools = menu.split('id="toolsSection"', 1)[1].split('<!-- Estrutura de "Daily Settlement"', 1)[0]
    grupo = tools.split('id="toolsCalculators"', 1)[1].split('</ul>', 1)[0]
    fora = tools.split('id="toolsCalculators"', 1)[0] + tools.split('id="toolsCalculators"', 1)[1].split('</ul>', 1)[1]
    _RE_ITEM = r'href="([^"#][^"]*)"[^>]*>\s*<span class="menu-text"[^>]*>\s*([^<]+?)\s*</span>'
    no_grupo = re.findall(_RE_ITEM, grupo)
    irmas = re.findall(_RE_ITEM, fora)
    itens = no_grupo + irmas
    check('as calculadoras estao no grupo Calculators, em ordem alfabetica',
          [r for _h, r in no_grupo],
          ['NDF Calculator', 'Option Calculator', 'Swap Calculator', 'Unwind NDF Calculator'])
    check('nenhuma calculadora ficou solta fora do grupo',
          [r for _h, r in irmas if 'calculator' in r.lower()], [])
    nivel1 = ['Calculators'] + [r for _h, r in irmas]
    check('o primeiro nivel de Tools em ordem alfabetica', nivel1, sorted(nivel1, key=lambda x: x.lower()))
    rotulos = [r for _h, r in itens]
    check('com as tres paginas novas',
          [r for r in rotulos if r in ('NDF Calculator', 'Option Calculator', 'Unwind NDF Calculator')],
          ['NDF Calculator', 'Option Calculator', 'Unwind NDF Calculator'])
    regras = set(str(r) for r in app.url_map.iter_rules())
    check('todo link do menu de Tools tem rota', [h for h, _r in itens if h not in regras], [])

    print('\n== 2. NDF no vencimento ==')
    r = D.liquidar_ndf(1000000.0, 5.20, 5.35, D.COMPRADO)
    check('comprado e o fixing ACIMA do termo: o banco recebe',
          (r.liquidacao, r.direcao, r.ir, r.liquido_cliente), (150000.0, 'RECEIVE', 0.0, None))
    r = D.liquidar_ndf(1000000.0, 5.20, 5.35, D.VENDIDO)
    check('vendido: o mesmo numero, o banco PAGA — e ai ha IR de 0,005%',
          (r.liquidacao, r.direcao, r.ir, r.liquido_cliente), (-150000.0, 'PAY', 7.5, 149992.5))
    check('contraparte isenta nao paga IR',
          D.liquidar_ndf(1000000.0, 5.20, 5.35, D.VENDIDO, isento_ir=True).ir, 0.0)
    r = D.liquidar_ndf(5200000.0, 5.20, 5.35, D.COMPRADO, fixo_em_reais=True)
    check('fixo em reais: o nocional divide pela taxa a termo',
          (round(r.nocional_me, 2), r.liquidacao), (1000000.0, 150000.0))
    check('fixing igual ao termo: nada liquida, sem direcao',
          (D.liquidar_ndf(1e6, 5.2, 5.2, D.COMPRADO).liquidacao,
           D.liquidar_ndf(1e6, 5.2, 5.2, D.COMPRADO).direcao), (0.0, ''))
    for nome, args in (('posicao em branco', (1e6, 5.2, 5.3, '')), ('fixing zero', (1e6, 5.2, 0, D.COMPRADO)),
                       ('nocional negativo', (-1, 5.2, 5.3, D.COMPRADO))):
        try:
            D.liquidar_ndf(*args)
            check(nome + ' e recusado', 'passou')
        except D.ErroDerivativo:
            check(nome + ' e recusado', True)

    print('\n== 3. recompra: a MESMA formula da pagina de Unwinds, nas tres operacoes reais ==')
    CASOS = [('USD 10/09/2026', 42227.42, 5.3748, 5.109, 13.75, 14, False, False, 11144.00, 42227.42),
             ('BRL fixed #1', 59999.98, 5.2039, 5.2077, 13.93, 45, True, True, 42.80, 11529.81),
             ('BRL fixed #2', 750000.01, 5.2039, 5.2067, 13.81, 23, True, True, 398.81, 144122.68)]
    for rot, noc, strike, term, pre, du, fixo, comprado, esperado, me in CASOS:
        r = D.recomprar_ndf(noc, strike, term, pre / 100.0, du,
                            D.COMPRADO if comprado else D.VENDIDO, fixo_em_reais=fixo)
        check('%s: resultado e nocional ME' % rot, (r.resultado, round(r.nocional_me, 2)), (esperado, me))
        ref = unw.conferir_apuracao({'Strike': str(strike)},
                                    {'Unwound Amount': str(noc), 'Termination Rate': str(term),
                                     'Pre FWD Rate': str(pre), 'DU': str(du),
                                     'Input Termination Fee': str(esperado)},
                                    brl_fixed=fixo, comprado=comprado)
        check('%s: bate com o conferir_apuracao da vertical' % rot,
              (r.resultado, r.direcao), (round(ref['resultado_calc'], 2), ref['direcao_calc']))
    check('os dias uteis sao ANBIMA, da liquidacao ao vencimento',
          D.dias_uteis_ate(date(2026, 9, 21), date(2026, 9, 28)), 5)
    try:
        D.dias_uteis_ate(date(2026, 9, 28), date(2026, 9, 21))
        check('vencimento antes da liquidacao e recusado', 'passou')
    except D.ErroDerivativo:
        check('vencimento antes da liquidacao e recusado', True)

    print('\n== 4. o saldo sao TRES parcelas ==')
    r = D.recomprar_ndf(300000.0, 5.2, 5.3, 0.13, 10, D.COMPRADO,
                        nocional_original=1000000.0, ja_recomprado=200000.0)
    check('original - ja recomprado - recomprado agora', (r.saldo, r.total), (500000.0, False))
    r = D.recomprar_ndf(800000.0, 5.2, 5.3, 0.13, 10, D.COMPRADO,
                        nocional_original=1000000.0, ja_recomprado=200000.0)
    check('zerou: recompra TOTAL', (r.saldo, r.total), (0.0, True))
    check('sem o original nao se afirma nem total nem parcial',
          (D.recomprar_ndf(1e5, 5.2, 5.3, 0.13, 10, D.COMPRADO).saldo,
           D.recomprar_ndf(1e5, 5.2, 5.3, 0.13, 10, D.COMPRADO).total), (None, None))
    try:
        D.recomprar_ndf(900000.0, 5.2, 5.3, 0.13, 10, D.COMPRADO,
                        nocional_original=1000000.0, ja_recomprado=200000.0)
        check('recomprar mais do que sobra e recusado', 'passou')
    except D.ErroDerivativo:
        check('recomprar mais do que sobra e recusado', True)

    print('\n== 5. opcao no exercicio ==')
    r = D.liquidar_opcao(D.CALL, D.TITULAR, 5.20, 1000000.0, [5.35], premio_unitario=0.05)
    check('call titular dentro do dinheiro: recebe o exercicio e PAGOU o premio',
          (r.payoff, r.direcao_exercicio, r.premio, r.direcao_premio, r.resultado),
          (150000.0, 'RECEIVE', 50000.0, 'PAY', 100000.0))
    r = D.liquidar_opcao(D.CALL, D.LANCADOR, 5.20, 1000000.0, [5.35], premio_unitario=0.05)
    check('o lancador e o espelho', (r.direcao_exercicio, r.direcao_premio, r.resultado),
          ('PAY', 'RECEIVE', -100000.0))
    r = D.liquidar_opcao(D.PUT, D.TITULAR, 5.20, 1000000.0, [5.35])
    check('put com o preco acima do strike: fora do dinheiro, nao exerce',
          (r.exercida, r.payoff, r.direcao_exercicio), (False, 0.0, ''))
    r = D.liquidar_opcao(D.PUT, D.TITULAR, 80.0, 1000.0, [70.0, 72.0, 74.0], paridade=5.0)
    check('asiatica: a media dos precos, e a paridade leva a reais',
          (r.fixing, r.intrinseco, r.payoff), (72.0, 8.0, 40000.0))
    r = D.liquidar_opcao(D.CALL, D.TITULAR, 80.0, 1000.0, [90.0], paridade=5.0,
                         premio_unitario=2.0, paridade_premio=4.0)
    check('o premio usa a paridade do dia DELE', (r.payoff, r.premio), (50000.0, 8000.0))
    try:
        D.liquidar_opcao(D.CALL, D.TITULAR, 80.0, 1000.0, [])
        check('sem preco de verificacao e recusado', 'passou')
    except D.ErroDerivativo:
        check('sem preco de verificacao e recusado', True)

    print('\n== 6. as tres paginas, de ponta a ponta ==')
    cl = app.test_client()
    with cl.session_transaction() as ss:
        ss['authenticated'] = True
        ss['user_sid'] = 'T000000'
        ss['user_name'] = 'Teste'
        ss['session_expires_at'] = (datetime.now(tz=timezone.utc) + timedelta(hours=8)).isoformat()
    for url in ('/tools/ndf-calculator', '/tools/unwind-ndf-calculator', '/tools/option-calculator'):
        check(url + ' abre', cl.get(url).status_code, 200)
        check(url + ' sem sessao redireciona', app.test_client().get(url).status_code in (301, 302))
    h = cl.post('/tools/ndf-calculator', data={
        'posicao': 'vendido', 'moeda': 'USD', 'nocional': '1,000,000.00', 'taxa_termo': '5.2000',
        'fixing': '5.3500', 'vencimento': '2026-09-21', 'ptax_offset': '1'}).data.decode('utf-8')
    check('NDF: o numero e a direcao na tela', ('150,000.00' in h, '>PAY<' in h, '149,992.50' in h),
          (True, True, True))
    # Fixing em branco: a PTAX do BCB, pela MESMA funcao do Swap Calculator.
    real = queries._ptax_do_fixing
    queries._ptax_do_fixing = lambda moeda, fim, n: (5.4321, date(2026, 9, 18), '')
    try:
        h = cl.post('/tools/ndf-calculator', data={
            'posicao': 'comprado', 'moeda': 'USD', 'nocional': '1000000', 'taxa_termo': '5.2',
            'fixing': '', 'vencimento': '2026-09-21', 'ptax_offset': '1'}).data.decode('utf-8')
        check('fixing em branco busca a PTAX e DIZ de que dia', 'PTAX USD 18/09/2026' in h)
        # A taxa USADA volta para o CAMPO da esquerda (mesa, 21/09/2026): e ele
        # que a mesa confere, e so no quadro do resultado ela ficava do outro
        # lado da tela. Vai marcada como automatica, com a procedencia embaixo.
        check('e a taxa usada VOLTA para o campo do formulario, marcada como automatica',
              ('id="fixing" name="fixing" value="5.4321"' in h,
               'id="fixing_auto" name="fixing_auto" value="1"' in h,
               'id="fixing_auto_nota">PTAX USD 18/09/2026' in h), (True, True, True))
        # Com a marca, o Calculate REBUSCA: trocar o vencimento nao pode deixar a
        # PTAX do vencimento anterior no campo, calada.
        queries._ptax_do_fixing = lambda moeda, fim, n: (5.5555, date(2026, 9, 17), '')
        h = cl.post('/tools/ndf-calculator', data={
            'posicao': 'comprado', 'moeda': 'USD', 'nocional': '1000000', 'taxa_termo': '5.2',
            'fixing': '5.4321', 'fixing_auto': '1', 'vencimento': '2026-09-18',
            'ptax_offset': '1'}).data.decode('utf-8')
        check('campo marcado como automatico e REBUSCADO no Calculate',
              'id="fixing" name="fixing" value="5.5555"' in h)
        h = cl.post('/tools/ndf-calculator', data={
            'posicao': 'comprado', 'moeda': 'USD', 'nocional': '1000000', 'taxa_termo': '5.2',
            'fixing': '5.4000', 'fixing_auto': '', 'vencimento': '2026-09-18',
            'ptax_offset': '1'}).data.decode('utf-8')
        check('o DIGITADO manda: sem a marca nao ha busca nem procedencia',
              ('id="fixing" name="fixing" value="5.4000"' in h,
               'id="fixing_auto" name="fixing_auto" value=""' in h), (True, True))
        h = cl.post('/tools/ndf-calculator', data={
            'posicao': 'comprado', 'moeda': 'USD', 'nocional': '1000000', 'taxa_termo': '5.2',
            'fixing': '', 'vencimento': '2030-01-15', 'ptax_offset': '1'}).data.decode('utf-8')
        check('fixing no FUTURO nao tem PTAX: erro em frase, nunca a ultima cotacao no lugar',
              ('still in the future' in h, 'Settlement at maturity' in h), (True, False))
        queries._ptax_do_fixing = lambda moeda, fim, n: (None, None, 'sem rede')
        h = cl.post('/tools/ndf-calculator', data={
            'posicao': 'comprado', 'moeda': 'USD', 'nocional': '1000000', 'taxa_termo': '5.2',
            'fixing': '', 'vencimento': '2026-09-21'}).data.decode('utf-8')
        check('sem PTAX e sem numero: erro em FRASE com o motivo, nunca um fixing chutado',
              ('sem rede' in h, 'alert-danger' in h, 'Settlement at maturity' in h), (True, True, False))
    finally:
        queries._ptax_do_fixing = real
    h = cl.post('/tools/unwind-ndf-calculator', data={
        'posicao': 'vendido', 'moeda': 'USD', 'nocional': '42,227.42', 'strike': '5.3748',
        'taxa_recompra': '5.109', 'taxa_pre': '13.75', 'du': '14',
        'liquidacao': '2026-09-10', 'vencimento': '2026-09-30'}).data.decode('utf-8')
    check('recompra: a operacao real fecha na tela (R$ 11.144,00, RECEIVE)',
          ('11,144.00' in h, '>RECEIVE<' in h), (True, True))
    h = cl.post('/tools/unwind-ndf-calculator', data={
        'posicao': 'vendido', 'moeda': 'USD', 'nocional': '1000', 'strike': '5.3', 'taxa_recompra': '5.1',
        'taxa_pre': '13.75', 'du': '', 'liquidacao': '2026-09-21', 'vencimento': '2026-09-28'}).data.decode('utf-8')
    check('DU em branco e CONTADO (ANBIMA), e a tela diz isso', 'counted, ANBIMA' in h)
    # O numero aparece NO CAMPO (mesa, 21/09/2026), marcado como automatico: e o
    # campo que a mesa le, e o Calculate devolve a ele o DU que usou.
    check('o DU contado VOLTA para o campo, marcado como automatico',
          ('id="du" name="du" value="5"' in h, 'id="du_auto" name="du_auto" value="1"' in h), (True, True))
    h = cl.post('/tools/unwind-ndf-calculator', data={
        'posicao': 'vendido', 'moeda': 'USD', 'nocional': '1000', 'strike': '5.3', 'taxa_recompra': '5.1',
        'taxa_pre': '13.75', 'du': '5', 'du_auto': '1', 'liquidacao': '2026-09-21',
        'vencimento': '2026-10-05'}).data.decode('utf-8')
    check('DU automatico e RECONTADO quando a data muda (nao sobra o numero da data anterior)',
          'id="du" name="du" value="10"' in h)
    h = cl.post('/tools/unwind-ndf-calculator', data={
        'posicao': 'vendido', 'moeda': 'USD', 'nocional': '1000', 'strike': '5.3', 'taxa_recompra': '5.1',
        'taxa_pre': '13.75', 'du': '14', 'du_auto': '', 'liquidacao': '2026-09-21',
        'vencimento': '2026-10-05'}).data.decode('utf-8')
    check('o DU DIGITADO manda (e como se reproduz o numero do aviso), sem a marca',
          ('id="du" name="du" value="14"' in h, 'id="du_auto" name="du_auto" value=""' in h, '>typed<' in h),
          (True, True, True))
    # A tela conta na hora, pelo SERVIDOR: a regra do calendario e uma so (§509).
    bd = cl.get('/api/tools/business-days?start=2026-09-21&end=2026-09-28').get_json()
    check('o endpoint dos dias uteis responde a MESMA conta do motor',
          (bd.get('success'), bd.get('days'), bd.get('calendar')),
          (True, D.dias_uteis_ate(date(2026, 9, 21), date(2026, 9, 28)), 'ANBIMA'))
    check('vencimento antes da liquidacao: 400 com o motivo',
          cl.get('/api/tools/business-days?start=2026-09-28&end=2026-09-21').status_code, 400)
    check('data invalida: 400', cl.get('/api/tools/business-days?start=x&end=y').status_code, 400)
    _js = io.open(os.path.join(ROOT, 'apps', 'static', 'js', 'pages', 'tools.js'), encoding='utf-8').read()
    check('e a busca pelo B3 ID dispara a contagem', 'contarDU();' in _js and "/api/tools/business-days" in _js, True)
    h = cl.post('/tools/option-calculator', data={
        'tipo': 'put', 'lado': 'titular', 'moeda': 'BRL', 'strike': '80', 'quantidade': '1000',
        'fixings': '70\n72;74', 'exercicio': '2026-09-21', 'premio_unitario': '2'}).data.decode('utf-8')
    check('opcao asiatica: media de 3, payoff 8.000 e premio 2.000',
          ('average of' in h, '8,000.00' in h, '2,000.00' in h, '6,000.00' in h), (True, True, True, True))
    h = cl.post('/tools/option-calculator', data={
        'tipo': 'call', 'lado': 'titular', 'moeda': 'BRL', 'strike': '', 'quantidade': '1000',
        'fixings': '70', 'exercicio': '2026-09-21'}).data.decode('utf-8')
    check('campo obrigatorio vazio vira frase, nao 500', 'alert-danger' in h)

    print('\n== 7. o B3 ID puxa a POSICAO (o esquema do Swap Calculator) ==')
    # A fonte e o MESMO coletor da tela de Live Position, e a leitura e a MESMA
    # da vertical de Unwinds: numero de TELA (o ultimo separador manda), saldo =
    # Valor Base no registro - Valor Antecipado, e na conta GUARDA-CHUVA o
    # cliente e a coluna do CPF/CNPJ, nunca o titular (que e o banco).
    from apps.pages import routes as R
    from apps.pages import quotes as Q
    NDF_COLS = list(R._LPNDF_COLUMNS)

    def ndf_row(**kw):
        return [kw.get(c, '') for c in NDF_COLS]
    LINHAS_NDF = [
        ndf_row(**{'Contrato': '26C03202688', 'Codigo Identificador': 'XE-10G5U5X-0-0',
                   'Codigo da Contraparte': '73760.10-2', 'Nome da Contraparte': 'BANCO J.P. MORGAN S/A',
                   'CPF/CNPJ da Contraparte': 'USINA ALTO ALEGRE SA', 'Simbolo da Moeda': 'USD',
                   'Classe do Ativo Subjacente': 'TAXAS DE CAMBIO', 'Data de Emissao': '04/11/2025',
                   'Data de Vencimento': '30/09/2026', 'Data de Fixing da Moeda': '29/09/2026',
                   'Valor Base no registro': '587,224.31', 'Valor Antecipado': '155,652.49',
                   'Taxa Forward': '5.374', 'Descricao da posicao do Participante': 'VENDEDOR'}),
        ndf_row(**{'Contrato': '26C09999999', 'Codigo da Contraparte': '73760.10-2',
                   'Nome da Contraparte': 'BANCO J.P. MORGAN S/A',
                   'CPF/CNPJ da Contraparte': '12.345.678/0001-90', 'Simbolo da Moeda': 'XXX',
                   'Data de Vencimento': '30/09/2026', 'Valor Base no registro': '1.000.000,00',
                   'Taxa Forward': '5,2000', 'Descricao da posicao do Participante': 'COMPRADOR'}),
        ndf_row(**{'Contrato': '26C01111111', 'Codigo da Contraparte': '04880.00-6',
                   'Nome da Contraparte': 'LAWTON MULTIMERCADO', 'Simbolo da Moeda': 'EUR',
                   'Data de Vencimento': '20260930', 'Valor Base no registro': '250,000.00',
                   'Taxa Forward': '6.1', 'Descricao da posicao do Participante': ''}),
    ]
    _o = (R._lpndf_collect, R._lpopt_collect, R._b3_is_omnibus, R._ndfc_ir_exempt,
          R._mapping_rows, Q.fetch_ohlc, Q.fetch_ptax)
    _subj_real = R._subjacente_by_code
    # HERMETICO: o prefill busca a PTAX quando a data do fixing ja passou, e sem
    # isto o teste dependeria da rede (e do BCB) para passar.
    _ptax_base = queries._ptax_do_fixing
    queries._ptax_do_fixing = lambda moeda, fim, n: (None, None, 'sem rede no teste')
    SUBJ = {}                       # o Index B3 do teste: vazio = ninguem cotado em centavos
    R._subjacente_by_code = lambda: SUBJ
    R._lpndf_collect = lambda ref, exact=False: {'columns': NDF_COLS, 'rows': LINHAS_NDF,
                                                 'source_date': '2026-09-18'}
    # por DIGITOS, como o `_b3_is_omnibus` de verdade: a posicao de NDF escreve
    # `73760.10-2` e a de opcao `73760102`
    R._b3_is_omnibus = lambda conta: ''.join(ch for ch in str(conta) if ch.isdigit()) == '73760102'
    R._ndfc_ir_exempt = lambda nome: 'ALTO ALEGRE' in str(nome)
    try:
        d = queries.ndf_prefill('26c03202688')
        f = d['fields']
        check('NDF: acha pelo Contrato, cego a caixa', (d['found'], d['b3_id'], d['source_date']),
              (True, '26C03202688', '2026-09-18'))
        check('guarda-chuva: a contraparte e a coluna do CPF/CNPJ, NAO o titular (o banco)',
              d['counterparty'], 'USINA ALTO ALEGRE SA')
        check('o nocional e o SALDO: valor base - ja recomprado', f['nocional'], '431571.82')
        check('a taxa `5.374` e 5,374 — tres casas NAO sao milhar numa taxa', f['taxa_termo'], '5.3740')
        check('moeda, vencimento, lado e o D-1 contado pelas duas datas',
              (f['moeda'], f['vencimento'], f['posicao'], f['ptax_offset']),
              ('USD', '2026-09-30', D.VENDIDO, '1'))
        check('emissao e classe do ativo vao para a tela', (f['data_emissao'], f['classe']),
              ('2025-11-04', 'TAXAS DE CAMBIO'))
        check('a isencao de IR sai do cadastro, pelo nome do CLIENTE', f['isento_ir'], True)
        check('e a nota diz quanto ja foi recomprado',
              [n['code'] for n in d['notes']], ['unwound_before'])
        # O vencimento da fixture (30/09/2026) ainda nao chegou: o fixing fica em
        # BRANCO. Com a data do fixing ja passada, a PTAX vem no CAMPO.
        check('fixing no futuro: campo em branco, sem marca', (f['fixing'], f['fixing_auto']), ('', ''))
        _venc = LINHAS_NDF[0][NDF_COLS.index('Data de Vencimento')]
        _fixd = LINHAS_NDF[0][NDF_COLS.index('Data de Fixing da Moeda')]
        LINHAS_NDF[0][NDF_COLS.index('Data de Vencimento')] = '18/09/2026'
        LINHAS_NDF[0][NDF_COLS.index('Data de Fixing da Moeda')] = '17/09/2026'
        _ptax_real = queries._ptax_do_fixing
        queries._ptax_do_fixing = lambda moeda, fim, n: (5.4321, date(2026, 9, 17), '')
        try:
            dp = queries.ndf_prefill('26C03202688')
            check('fixing ja passado: a PTAX vem no CAMPO, marcada, com a nota da data',
                  (dp['fields']['fixing'], dp['fields']['fixing_auto'],
                   [n for n in dp['notes'] if n['code'] == 'fixing_ptax'][0]['params']),
                  ('5.4321', '1', {'moeda': 'USD', 'data': '17/09/2026'}))
        finally:
            queries._ptax_do_fixing = _ptax_real
            LINHAS_NDF[0][NDF_COLS.index('Data de Vencimento')] = _venc
            LINHAS_NDF[0][NDF_COLS.index('Data de Fixing da Moeda')] = _fixd
        check('o Codigo Identificador tambem acha', queries.ndf_prefill('XE-10G5U5X-0-0')['b3_id'],
              '26C03202688')
        d2 = queries.ndf_prefill('26C09999999')
        check('guarda-chuva com DOCUMENTO sem cadastro: nome VAZIO avisando, nunca o banco',
              (d2['counterparty'], [n['code'] for n in d2['notes']]), ('', ['cpty_not_registered']))
        check('numero brasileiro na mesma posicao (o ultimo separador manda)',
              (d2['fields']['nocional'], d2['fields']['taxa_termo']), ('1000000.00', '5.2000'))
        check('moeda que a tela nao conhece fica em branco e SINALIZADA',
              (d2['fields']['moeda'], 'moeda' in d2['missing']), ('', True))
        d3 = queries.ndf_prefill('26C01111111')
        check('conta direta: o Nome da Contraparte E a contraparte', d3['counterparty'], 'LAWTON MULTIMERCADO')
        check('lado ilegivel NAO se chuta — o sinal depende dele',
              (d3['fields']['posicao'], 'posicao' in d3['missing']), ('', True))
        check('data aaaammdd tambem e lida', d3['fields']['vencimento'], '2026-09-30')
        check('contrato fora da posicao: nao achou, com a data do arquivo',
              (queries.ndf_prefill('NADA')['found'], queries.ndf_prefill('NADA')['source_date']),
              (False, '2026-09-18'))

        u = queries.unwind_ndf_prefill('26C03202688')
        uf = u['fields']
        check('recompra: o contrato ORIGINAL — strike, base e o ja recomprado',
              (uf['strike'], uf['nocional_original'], uf['ja_recomprado'], uf['vencimento'], uf['posicao']),
              ('5.3740', '587224.31', '155652.49', '2026-09-30', D.VENDIDO))
        check('o nocional nasce com o SALDO (recompra total), marcado como aproximacao',
              (uf['nocional'], u['assumed']), ('431571.82', ['nocional']))
        check('as taxas do NEGOCIO ficam para a mesa, sinalizadas',
              [m for m in u['missing'] if m.startswith('taxa_')], ['taxa_recompra', 'taxa_pre'])

        print('\n== 8. a opcao: a posicao e os precos de verificacao do QUOTES ==')
        OPT_COLS = list(R._LPOPT_COLUMNS) + ['Média Asiática (data) 1', 'Média Asiática (data) 2',
                                             'Média Asiática (data) 3']

        def opt_row(**kw):
            return [kw.get(c, '') for c in OPT_COLS]
        LINHAS_OPT = [
            opt_row(**{'Código IF': 'CHASM26081V', 'Combinação de operações': 'D5YJ-QD5YF',
                       'Tipo de Opção': 'PUT', 'Posição da Parte': 'TITULAR',
                       'Parte (Nome simplificado)': 'JPMORGAN', 'Contraparte (Nome simplificado)': 'BASF SA',
                       'Data Registro': '10/03/2026', 'Data Vencimento': '18/09/2026',
                       'Classe do ativo subjacente': 'COMMODITIES', 'Ativo subjacente / Moeda base': 'CTZ6',
                       'Moeda do ativo / Moeda cotada': 'USD', 'Quantidade': '1,000.00',
                       'Quantidade Antecipada': '250.00', 'Strike (valor)': '80.5',
                       'Prêmio Unitário': '2.125', 'Quantidade de datas de verificação': '3',
                       'Média Asiática (data) 1': '14/09/2026', 'Média Asiática (data) 2': '15/09/2026',
                       'Média Asiática (data) 3': '16/09/2026'}),
            opt_row(**{'Código IF': 'FXO0000001', 'Tipo de Opção': 'CALL', 'Posição da Parte': 'LANÇADOR',
                       'Contraparte (Nome simplificado)': 'CLIENTE FX', 'Data Vencimento': '15/09/2026',
                       'Classe do ativo subjacente': 'TAXA DE CAMBIO', 'Ativo subjacente / Moeda base': 'USD',
                       'Moeda do ativo / Moeda cotada': 'BRL', 'Quantidade': '500000', 'Strike (valor)': '5.4',
                       'Data de fixing do ativo subjacente': '14/09/2026', 'Barreira de KO': '6.2'}),
        ]
        R._lpopt_collect = lambda ref, exact=False: {'columns': OPT_COLS, 'rows': LINHAS_OPT,
                                                     'source_date': '2026-09-18'}
        R._mapping_rows = lambda k: ([{'LABEL': 'CT"MY"', 'SYMBOL': 'CT"MY".NYB'}]
                                     if k == 'quotes-commodity' else [])
        pedidos = []

        def _ohlc(sym, ini, fim):
            pedidos.append(sym)
            return (list(Q.OHLC_COLUMNS),
                    [['16/09/2026', '74.10', '74.000000', '', '', '', ''],
                     ['15/09/2026', '72.10', '72.000000', '', '', '', ''],
                     ['14/09/2026', '70.10', '70.000000', '', '', '', '']])
        Q.fetch_ohlc = _ohlc
        o = queries.opcao_prefill('chasm26081v')
        of = o['fields']
        check('opcao: tipo, lado, strike, exercicio e o premio unitario',
              (of['tipo'], of['lado'], of['strike'], of['exercicio'], of['premio_unitario'], of['moeda']),
              (D.PUT, D.TITULAR, '80.5000', '2026-09-18', '2.1250', 'USD'))
        check('a quantidade e a ABERTA (quantidade - antecipada)', of['quantidade'], '750')
        check('asiatica: os TRES precos vem do Quotes, o Close (nao o Adj Close), numa chamada so',
              (of['fixings'], len(pedidos)), ('70.0000\n72.0000\n74.0000', 1))
        check('o simbolo sai do cadastro quotes-commodity (o "MY" do vencimento)', pedidos[0], 'CTZ26.NYB')
        # COTADO EM CENTAVOS (mesa, 21/09/2026): o Index B3 diz `Fator Conversao`
        # 0,01 para o CTZ6 — a bolsa cota 81,15 c/lb e o strike da B3 e 0,6960
        # US$/lb. Sem o fator a opcao sairia CEM vezes dentro do dinheiro.
        SUBJ['CTZ6'] = {'Fator Conversao': 0.01, 'Moeda': 'DOLAR DOS EUA'}
        oc = queries.opcao_prefill('CHASM26081V')
        check('cotado em centavos: os precos do Quotes saem x 0,01',
              oc['fields']['fixings'], '0.7000\n0.7200\n0.7400')
        check('e a nota DIZ que multiplicou, com o ativo',
              [n['params'] for n in oc['notes'] if n['code'] == 'quoted_in_cents'], [{'ativo': 'CTZ6'}])
        SUBJ['CTZ6'] = {'Fator Conversao': '0,01'}
        check('o fator escrito como texto brasileiro tambem e centavos',
              queries.opcao_prefill('CHASM26081V')['fields']['fixings'].split('\n')[0], '0.7000')
        SUBJ['CTZ6'] = {'Fator Conversao': 5000.0}
        check('so 0,01 e centavos: outro fator NAO mexe no preco (a regra da casa)',
              queries.opcao_prefill('CHASM26081V')['fields']['fixings'].split('\n')[0], '70.0000')
        # a moeda vem por NOME na posicao: quem traduz e o cadastro currency-base
        i_moeda = OPT_COLS.index('Moeda do ativo / Moeda cotada')
        LINHAS_OPT[0][i_moeda] = 'DOLAR DOS EUA'
        _map_prev = R._mapping_rows
        R._mapping_rows = lambda k: ([{'DESCRICAO DO CAMPO': 'DOLAR DOS EUA', 'CODIGO DE CADASTRO': '220',
                                       'SIMBOLO': 'USD'}] if k == 'currency-base' else _map_prev(k))
        mo = queries.opcao_prefill('CHASM26081V')
        check('`DOLAR DOS EUA` vira USD pelo cadastro — e sai da lista do que faltou',
              (mo['fields']['moeda'], 'moeda' in mo['missing']), ('USD', False))
        LINHAS_OPT[0][i_moeda] = ''
        SUBJ['CTZ6'] = {'Fator Conversao': 0.01, 'Moeda': 'DOLAR DOS EUA'}
        check('coluna vazia: responde a Moeda do proprio ativo no Index B3',
              queries.opcao_prefill('CHASM26081V')['fields']['moeda'], 'USD')
        check('o codigo Sisbacen tambem traduz', queries._moeda_iso('220'), 'USD')
        LINHAS_OPT[0][i_moeda] = 'USD'
        R._mapping_rows = _map_prev
        SUBJ.clear()
        check('o lado vai marcado como aproximacao — e o da PARTE do registro',
              ('lado' in o['assumed'], [n['code'] for n in o['notes']][:1]), (True, ['side_of_party']))
        check('e as notas dizem de onde os precos vieram',
              'fixings_from_quotes' in [n['code'] for n in o['notes']], True)
        check('registro e classe · ativo na tela', (of['data_emissao'], of['classe']),
              ('2026-03-10', 'COMMODITIES \u00b7 CTZ6'))
        # A contraparte da OPCAO (mesa, 21/09/2026): o `Nome simplificado` e
        # apelido de conta (JPMORGANBM, INTRAGLAWTONFDO) e nao identifica ninguem.
        # Guarda-chuva -> o CNPJ pelo Reference Data; outra conta -> a CONTA CETIP
        # pelo Reference Data. Quem diz que a conta e guarda-chuva e o cadastro
        # `b3-accounts`, nunca o numero no codigo.
        _conta_real = R._lp_cpty_by_account
        R._lp_cpty_by_account = lambda c: {'00041007': 'LAWTON MULTIMERCADO EXCLUSIVO'}.get(
            ''.join(ch for ch in str(c) if ch.isdigit()), '')
        try:
            i_conta, i_doc = OPT_COLS.index('Contraparte (Conta)'), OPT_COLS.index('CPF/CNPJ Cliente Contraparte')
            i_ap = OPT_COLS.index('Contraparte (Nome simplificado)')
            LINHAS_OPT[0][i_conta], LINHAS_OPT[0][i_doc], LINHAS_OPT[0][i_ap] = '73760102', 'SUZANO SA', 'JPMORGANBM'
            check('guarda-chuva (73760102): o nome sai do CNPJ resolvido, NAO do apelido',
                  queries.opcao_prefill('CHASM26081V')['counterparty'], 'SUZANO SA')
            LINHAS_OPT[0][i_doc] = '12.345.678/0001-90'
            g = queries.opcao_prefill('CHASM26081V')
            check('guarda-chuva com CNPJ sem cadastro: nome VAZIO avisando, nunca o banco',
                  (g['counterparty'], 'cpty_not_registered' in [n['code'] for n in g['notes']]), ('', True))
            LINHAS_OPT[0][i_conta], LINHAS_OPT[0][i_doc], LINHAS_OPT[0][i_ap] = '00041007', '', 'INTRAGLAWTONFDO'
            check('outra conta (00041007): o nome sai do Reference Data pela CONTA CETIP',
                  queries.opcao_prefill('CHASM26081V')['counterparty'], 'LAWTON MULTIMERCADO EXCLUSIVO')
            LINHAS_OPT[0][i_conta] = '99999000'
            s2 = queries.opcao_prefill('CHASM26081V')
            check('conta fora do Reference Data: fica o apelido, AVISANDO que e ele',
                  (s2['counterparty'], 'cpty_short_name' in [n['code'] for n in s2['notes']]),
                  ('INTRAGLAWTONFDO', True))
        finally:
            R._lp_cpty_by_account = _conta_real
        # `Strike/Limitador/Barreiras em Reais` = S: o preco ja esta em reais.
        i_sr = OPT_COLS.index('Strike/Limitador/Barreiras em Reais')
        LINHAS_OPT[0][i_sr] = 'S'
        check('strike em reais (S): sem conversao, mesmo com moeda cotada estrangeira',
              queries.opcao_prefill('CHASM26081V')['fields']['moeda'], 'BRL')
        LINHAS_OPT[0][i_sr] = 'N'
        check('strike fora de reais (N): a moeda cotada, e a paridade leva a reais',
              queries.opcao_prefill('CHASM26081V')['fields']['moeda'], 'USD')
        check('o numero da confirmacao tambem acha', queries.opcao_prefill('D5YJ-QD5YF')['b3_id'],
              'CHASM26081V')
        # A PARIDADE vem no campo ja na busca (mesa, 21/09/2026): a PTAX de venda
        # da moeda do preco, D-n do exercicio, com o `n` contado pela data de
        # fixing da MOEDA que a posicao traz.
        _ptax_opt = queries._ptax_do_fixing
        pedidos_ptax = []

        def _ptax_stub(moeda, fim, n):
            pedidos_ptax.append((moeda, fim, n))
            return 5.4321, date(2026, 9, 16), ''
        queries._ptax_do_fixing = _ptax_stub
        i_fm = OPT_COLS.index('Data de fixing da moeda do ativo subjacente')
        LINHAS_OPT[0][i_fm] = '16/09/2026'                   # exercicio 18/09 -> D-2
        try:
            pr = queries.opcao_prefill('CHASM26081V')
            check('a paridade vem PREENCHIDA e marcada como automatica, com o D-n da posicao',
                  (pr['fields']['paridade'], pr['fields']['paridade_auto'], pr['fields']['ptax_offset'],
                   pedidos_ptax[-1]), ('5.4321', '1', '2', ('USD', '2026-09-18', 2)))
            check('e a nota diz de que dia e a PTAX',
                  [n['params'] for n in pr['notes'] if n['code'] == 'parity_ptax'],
                  [{'moeda': 'USD', 'data': '16/09/2026'}])
            LINHAS_OPT[0][i_fm] = ''
            check('sem a data de fixing da moeda: o D-1 de costume',
                  queries.opcao_prefill('CHASM26081V')['fields']['ptax_offset'], '1')
            i_ex = OPT_COLS.index('Data Vencimento')
            _ex = LINHAS_OPT[0][i_ex]
            LINHAS_OPT[0][i_ex] = '15/01/2030'
            fu = queries.opcao_prefill('CHASM26081V')
            check('exercicio no FUTURO: paridade em branco, sem marca — a PTAX ainda nao existe',
                  (fu['fields']['paridade'], fu['fields']['paridade_auto']), ('', ''))
            LINHAS_OPT[0][i_ex] = _ex
            queries._ptax_do_fixing = lambda moeda, fim, n: (None, None, 'sem rede')
            fa = queries.opcao_prefill('CHASM26081V')
            check('PTAX fora do ar: campo em branco, SINALIZADO, com o motivo',
                  (fa['fields']['paridade'], 'paridade' in fa['missing'],
                   [n['params']['motivo'] for n in fa['notes'] if n['code'] == 'parity_failed']),
                  ('', True, ['sem rede']))
            i_sr2 = OPT_COLS.index('Strike/Limitador/Barreiras em Reais')
            LINHAS_OPT[0][i_sr2] = 'S'
            queries._ptax_do_fixing = _ptax_stub
            n_antes = len(pedidos_ptax)
            br = queries.opcao_prefill('CHASM26081V')
            check('preco em reais: nao ha paridade a buscar',
                  (br['fields']['paridade'], len(pedidos_ptax) == n_antes), ('', True))
            LINHAS_OPT[0][i_sr2] = ''
        finally:
            queries._ptax_do_fixing = _ptax_opt
        # cambio: a PTAX de VENDA da moeda base, pela data de fixing do ativo
        Q.fetch_ptax = lambda ccy, ini, fim: (list(Q.PTAX_COLUMNS),
                                              [['14/09/2026', ccy, '5.4300', '5.4321', '', '']])
        fx = queries.opcao_prefill('FXO0000001')
        check('opcao de cambio: o preco e a PTAX de venda do dia do fixing; preco em reais',
              (fx['fields']['fixings'], fx['fields']['moeda'], fx['fields']['lado']),
              ('5.4321', 'BRL', D.LANCADOR))
        check('barreira no contrato e AVISADA (nao e apurada aqui)',
              'has_barrier' in [n['code'] for n in fx['notes']], True)
        # serie pela metade NAO e media
        Q.fetch_ohlc = lambda sym, ini, fim: (list(Q.OHLC_COLUMNS), [])
        m = queries.opcao_prefill('CHASM26081V')
        check('sem preco no Quotes: fixings em branco e SINALIZADO, com a nota',
              (m['fields']['fixings'], 'fixings' in m['missing'],
               'fixings_missing' in [n['code'] for n in m['notes']]), ('', True, True))
        R._mapping_rows = lambda k: []
        s = queries.opcao_prefill('CHASM26081V')
        check('simbolo fora do cadastro: o MOTIVO vai na nota (pede cadastro)',
              'quotes-commodity' in [n for n in s['notes'] if n['code'] == 'fixings_missing'][0]['params']['motivo'],
              True)

        print('\n== 8b. o termo de MERCADORIA no NDF Calculator ==')
        # O nocional e QUANTIDADE, o fixing e o PRECO do ativo — do Quotes, x 0,01
        # quando o Index B3 diz centavos — e a paridade (a PTAX) leva a reais.
        LINHAS_NDF.append(ndf_row(**{
            'Contrato': '26C05550000', 'Codigo da Contraparte': '04880.00-6',
            'Nome da Contraparte': 'COFCO INTERNATIONAL BRASIL SA', 'Simbolo da Moeda': 'USD',
            'Classe do Ativo Subjacente': 'COMMODITIES', 'Codigo do Ativo Subjacente': 'CTZ6',
            'Data de Vencimento': '18/09/2026', 'Data de Fixing da Moeda': '17/09/2026',
            'Data de Fixing do Ativo Subjacente': '16/09/2026',
            'Valor Base no registro': '100,000.00', 'Taxa Forward': '0.696',
            'Descricao da posicao do Participante': 'COMPRADOR'}))
        R._mapping_rows = lambda k: ([{'LABEL': 'CT"MY"', 'SYMBOL': 'CT"MY".NYB'}]
                                     if k == 'quotes-commodity' else [])
        Q.fetch_ohlc = _ohlc
        SUBJ['CTZ6'] = {'Fator Conversao': 0.01}
        _ptax_real2 = queries._ptax_do_fixing
        queries._ptax_do_fixing = lambda moeda, fim, n: (5.4, date(2026, 9, 17), '')
        try:
            c = queries.ndf_prefill('26C05550000')
            cf = c['fields']
            check('mercadoria: o fixing e o PRECO do Quotes x 0,01 (74 c -> 0,74), nao a PTAX',
                  (cf['fixing'], cf['fixing_auto']), ('0.7400', ''))
            check('a paridade e a PTAX, marcada como automatica; e o ativo vai para a tela',
                  (cf['paridade'], cf['paridade_auto'], cf['ativo']), ('5.4000', '1', 'CTZ6'))
            check('as notas dizem Quotes, centavos e a PTAX da paridade',
                  [n['code'] for n in c['notes']], ['fixings_from_quotes', 'quoted_in_cents', 'parity_ptax'])
            h = cl.post('/tools/ndf-calculator', data={
                'posicao': 'comprado', 'moeda': 'USD', 'nocional': '100000', 'taxa_termo': '0.696',
                'fixing': '0.74', 'vencimento': '2026-09-18', 'ptax_offset': '1',
                'classe': 'COMMODITIES', 'ativo': 'CTZ6', 'paridade': '', 'paridade_auto': ''}).data.decode('utf-8')
            # 100.000 x (0,74 - 0,696) x 5,4 = 23.760,00
            check('a conta: quantidade x (preco - termo) x paridade, em reais',
                  ('23,760.00' in h, 'id="paridade" name="paridade" value="5.4000"' in h,
                   'id="tl-ndf-commodity" >' in h or 'id="tl-ndf-commodity"  >' in h), (True, True, True))
            h = cl.post('/tools/ndf-calculator', data={
                'posicao': 'comprado', 'moeda': 'USD', 'nocional': '100000', 'taxa_termo': '0.696',
                'fixing': '', 'vencimento': '2026-09-18', 'classe': 'COMMODITIES'}).data.decode('utf-8')
            check('mercadoria com o fixing em branco NAO vira PTAX: erro dizendo de onde ele vem',
                  ('PRICE of the underlying' in h, 'Settlement at maturity' in h), (True, False))
            check('no termo de MOEDA o bloco da mercadoria nao aparece',
                  'id="tl-ndf-commodity" hidden' in cl.get('/tools/ndf-calculator').data.decode('utf-8'), True)
        finally:
            queries._ptax_do_fixing = _ptax_real2
            SUBJ.clear()
            LINHAS_NDF.pop()

        print('\n== 9. o endpoint e as tres telas ==')
        R._mapping_rows = _o[4]
        for tool in ('ndf-calculator', 'unwind-ndf-calculator', 'option-calculator'):
            h = cl.get('/tools/' + tool).data.decode('utf-8')
            check(tool + ': o bloco do B3 ID aponta para o endpoint DELA',
                  ('data-tl-prefill="/api/tools/%s/prefill"' % tool in h, 'id="tl-pos-lookup"' in h),
                  (True, True))
        j = cl.get('/api/tools/ndf-calculator/prefill?id=26C03202688').get_json()
        check('GET: os campos pelo id do formulario', (j['success'], j['fields']['nocional']), (True, '431571.82'))
        r404 = cl.get('/api/tools/ndf-calculator/prefill?id=NADA')
        check('fora da posicao: 404 com CODIGO e os parametros da frase',
              (r404.status_code, r404.get_json().get('code'), r404.get_json()['params']['data']),
              (404, 'not_in_position', '2026-09-18'))
        check('sem id: 400', cl.get('/api/tools/ndf-calculator/prefill').status_code, 400)
        check('ferramenta que nao existe: 404', cl.get('/api/tools/nada/prefill?id=1').status_code, 404)
        sw = cl.get('/api/tools/swap-calculator/prefill?id=NAOEXISTE').get_json()
        check('a rota do Swap Calculator NAO foi sombreada', 'swap position' in (sw.get('error') or ''), True)
        R._lpndf_collect = lambda ref, exact=False: (_ for _ in ()).throw(IOError('banco ocupado'))
        r500 = cl.get('/api/tools/ndf-calculator/prefill?id=26C03202688')
        check('falha de leitura volta com o MOTIVO, em JSON',
              (r500.status_code, 'IOError' in (r500.get_json().get('error') or 'OSError')
               or 'OSError' in r500.get_json().get('error', '')), (500, True))
    finally:
        (R._lpndf_collect, R._lpopt_collect, R._b3_is_omnibus, R._ndfc_ir_exempt,
         R._mapping_rows, Q.fetch_ohlc, Q.fetch_ptax) = _o
        R._subjacente_by_code = _subj_real
        queries._ptax_do_fixing = _ptax_base

    print('\n== 10. Export: a memoria de calculo em Excel, na logica da do swap ==')
    # FORMULA, nao valor — e com o VALOR gravado junto, porque o Modo de Exibicao
    # Protegido e o preview do anexo nao recalculam (a memoria chegaria em branco).
    # O cache sai da PROPRIA formula: o que se cobra aqui e que ele seja o numero
    # do MOTOR. Planilha afirmando um numero que a tela nao deu e o defeito.
    import io as _io
    import zipfile
    import openpyxl
    from apps.pages.features.tools.infra import memoria_derivativos as M

    def abrir(conteudo):
        return (openpyxl.load_workbook(_io.BytesIO(conteudo)).active,
                openpyxl.load_workbook(_io.BytesIO(conteudo), data_only=True).active)

    def celula(ws, rotulo):
        for row in ws.iter_rows(min_col=1, max_col=2):
            if row[0].value == rotulo:
                return row[1].value
        return None

    rn = D.liquidar_ndf(1000000.0, 5.20, 5.35, D.VENDIDO)
    wf, wv = abrir(M.ndf(rn, 'USD', date(2026, 9, 21), cetip_id='26C03202688',
                         contraparte='USINA ALTO ALEGRE SA', nocional_informado=1000000.0,
                         fixing_nota='PTAX USD 18/09/2026'))
    check('NDF: a liquidacao e FORMULA', str(celula(wf, 'Liquidação (resultado do banco)')).startswith('=ROUND('))
    check('NDF: o valor em cache e o do motor (liquidacao, IR e liquido ao cliente)',
          (celula(wv, 'Liquidação (resultado do banco)'), celula(wv, 'IR da operação'),
           celula(wv, 'Líquido ao cliente')), (rn.liquidacao, rn.ir, rn.liquido_cliente))
    check('NDF: a parte devedora e NOMEADA (o banco paga)', celula(wv, 'Parte devedora'), 'Banco J.P. Morgan')
    rm = D.liquidar_ndf(100000.0, 0.696, 0.74, D.COMPRADO, paridade=5.4)
    _f, wv = abrir(M.ndf(rm, 'USD', date(2026, 9, 18), classe='COMMODITIES', ativo='CTZ6',
                         contraparte='COFCO', nocional_informado=100000.0))
    check('NDF de mercadoria: a paridade entra na formula e o cache e o do motor',
          (celula(wv, 'Liquidação (resultado do banco)'), celula(wv, 'Paridade para reais'),
           celula(wv, 'Parte devedora')), (rm.liquidacao, 5.4, 'COFCO'))
    rf = D.liquidar_ndf(5200000.0, 5.20, 5.35, D.COMPRADO, fixo_em_reais=True)
    _f, wv = abrir(M.ndf(rf, 'USD', date(2026, 9, 21), nocional_informado=5200000.0, fixo_em_reais=True))
    check('fixo em reais: a planilha divide pela taxa POR FORMULA',
          (round(celula(wv, 'Nocional em moeda estrangeira'), 2), celula(wv, 'Liquidação (resultado do banco)')),
          (1000000.0, rf.liquidacao))

    for rot, noc, strike, term, pre, du, fixo, comprado, esperado, _me in CASOS:
        ru = D.recomprar_ndf(noc, strike, term, pre / 100.0, du, D.COMPRADO if comprado else D.VENDIDO,
                             fixo_em_reais=fixo)
        wf, wv = abrir(M.unwind_ndf(ru, 'USD', date(2026, 9, 10), date(2026, 9, 30),
                                    nocional_informado=noc, fixo_em_reais=fixo))
        check('recompra %s: o cache da planilha e o resultado REAL (%s)' % (rot, esperado),
              celula(wv, 'Resultado da recompra (valor presente, do banco)'), esperado)
    check('recompra: o fator de desconto e FORMULA — (1 + pre) ^ (DU / 252)',
          bool(re.match(r'^=\(1\+\$B\$\d+\)\^\(\$B\$\d+/252\)$', str(celula(wf, 'Fator de desconto')))), True)
    rs = D.recomprar_ndf(300000.0, 5.2, 5.3, 0.13, 10, D.COMPRADO, nocional_original=1000000.0,
                         ja_recomprado=200000.0)
    _f, wv = abrir(M.unwind_ndf(rs, 'USD', date(2026, 9, 10), date(2026, 9, 30), nocional_informado=300000.0,
                                original_informado=1000000.0, ja_recomprado=200000.0))
    check('recompra: o novo valor base sai das TRES parcelas, por formula',
          (celula(wv, 'Novo valor base'), celula(wv, 'Recompra')), (500000.0, 'Parcial'))

    ro = D.liquidar_opcao(D.PUT, D.TITULAR, 80.0, 1000.0, [70.0, 72.0, 74.0], paridade=5.0,
                          premio_unitario=2.0, paridade_premio=4.0)
    wf, wv = abrir(M.opcao(ro, 'USD', date(2026, 9, 18), cetip_id='CHASM26081V', contraparte='SUZANO SA',
                           premio_unitario=2.0, paridade_premio=4.0))
    check('opcao asiatica: a media e a soma das TRES celulas / 3 (o avaliador nao conhece AVERAGE)',
          str(celula(wf, 'Preço de exercício apurado')).count('+'), 2)
    check('opcao: exercicio, premio e resultado do banco em cache = o motor',
          (celula(wv, 'Liquidação do exercício'), celula(wv, 'Prêmio'), celula(wv, 'Resultado do banco'),
           celula(wv, 'Preço de exercício apurado')), (ro.payoff, ro.premio, ro.resultado, 72.0))
    rl = D.liquidar_opcao(D.CALL, D.LANCADOR, 5.20, 1000000.0, [5.35], premio_unitario=0.05)
    _f, wv = abrir(M.opcao(rl, 'BRL', date(2026, 9, 18), contraparte='CLIENTE', premio_unitario=0.05))
    check('opcao, banco LANCADOR: o sinal vira por formula e o banco e o devedor',
          (celula(wv, 'Resultado do banco'), celula(wv, 'Parte devedora')), (rl.resultado, 'Banco J.P. Morgan'))

    # o arquivo: o cache esta ESCRITO no xml (nas duas grafias do <v> vazio) e o
    # documento e do banco — nada do sistema que o gerou
    bruto = M.opcao(ro, 'USD', date(2026, 9, 18))
    z = zipfile.ZipFile(_io.BytesIO(bruto))
    folha = z.read('xl/worksheets/sheet1.xml').decode('utf-8')
    check('nenhuma formula ficou com o valor em branco no xml',
          ('<v></v>' in folha, '<v />' in folha, '<v/>' in folha), (False, False, False))
    app_xml = z.read('docProps/app.xml').decode('utf-8') if 'docProps/app.xml' in z.namelist() else ''
    check('nada do sistema que gerou (openpyxl / OTC Tracker) no documento',
          ('openpyxl' in app_xml.lower(), 'otc tracker' in (app_xml + folha).lower()), (False, False))

    print('\n== 11. o Export pelas tres telas ==')
    for tool, dados in (
            ('ndf-calculator', {'posicao': 'vendido', 'moeda': 'USD', 'nocional': '1,000,000.00',
                                'taxa_termo': '5.2', 'fixing': '5.35', 'vencimento': '2026-09-21',
                                'b3_id': '26C03202688', 'counterparty': 'USINA ALTO ALEGRE SA'}),
            ('unwind-ndf-calculator', {'posicao': 'vendido', 'moeda': 'USD', 'nocional': '42,227.42',
                                       'strike': '5.3748', 'taxa_recompra': '5.109', 'taxa_pre': '13.75',
                                       'du': '14', 'liquidacao': '2026-09-10', 'vencimento': '2026-09-30'}),
            ('option-calculator', {'tipo': 'put', 'lado': 'titular', 'moeda': 'BRL', 'strike': '80',
                                   'quantidade': '1000', 'fixings': '70\n72\n74', 'exercicio': '2026-09-18',
                                   'premio_unitario': '2'})):
        h = cl.get('/tools/' + tool).data.decode('utf-8')
        check(tool + ': o botao Export e o action EXPLICITO do formulario',
              ('id="tl-export"' in h and ('formaction="/tools/%s/extract"' % tool) in h,
               ('<form method="post" action="/tools/%s"' % tool) in h), (True, True))
        r = cl.post('/tools/%s/extract' % tool, data=dados)
        cd = r.headers.get('Content-Disposition', '')
        check(tool + ': baixa um .xlsx com o nome do DOCUMENTO',
              (r.status_code, r.data[:2], 'Mem' in cd and '.xlsx' in cd), (200, b'PK', True))
    nome = cl.post('/tools/ndf-calculator/extract', data={
        'posicao': 'vendido', 'moeda': 'USD', 'nocional': '1000000', 'taxa_termo': '5.2', 'fixing': '5.35',
        'vencimento': '2026-09-21', 'b3_id': '26C03202688',
        'counterparty': 'USINA ALTO ALEGRE SA'}).headers.get('Content-Disposition', '')
    check('o nome leva CETIP ID, contraparte e a data', ('26C03202688' in nome, '21-09-2026' in nome), (True, True))
    # O PRODUTO no nome (mesa, 21/09/2026): e o que distingue os arquivos na pasta
    # de quem baixou. O cabecalho manda o nome em RFC 5987 (percent-encoded).
    from urllib.parse import unquote
    nomes = {}
    for tool, dados in (
            ('ndf-calculator', {'posicao': 'vendido', 'moeda': 'USD', 'nocional': '1000000',
                                'taxa_termo': '5.2', 'fixing': '5.35', 'vencimento': '2026-09-21'}),
            ('unwind-ndf-calculator', {'posicao': 'vendido', 'moeda': 'USD', 'nocional': '42227.42',
                                       'strike': '5.3748', 'taxa_recompra': '5.109', 'taxa_pre': '13.75',
                                       'du': '14', 'liquidacao': '2026-09-10', 'vencimento': '2026-09-30'}),
            ('option-calculator', {'tipo': 'put', 'lado': 'titular', 'moeda': 'BRL', 'strike': '80',
                                   'quantidade': '1000', 'fixings': '70', 'exercicio': '2026-09-18'})):
        nomes[tool] = unquote(cl.post('/tools/%s/extract' % tool, data=dados).headers.get('Content-Disposition', ''))
    check('cada calculadora diz o SEU produto no nome do arquivo',
          ('Memória de Cálculo NDF - 21-09-2026.xlsx' in nomes['ndf-calculator'],
           'Memória de Cálculo Recompra NDF - 10-09-2026.xlsx' in nomes['unwind-ndf-calculator'],
           'Memória de Cálculo Opção - 18-09-2026.xlsx' in nomes['option-calculator']),
          (True, True, True))
    rx = cl.post('/tools/option-calculator/extract', data={'tipo': 'call', 'lado': 'titular', 'strike': ''},
                 headers={'X-Requested-With': 'XMLHttpRequest'})
    check('conta que nao fecha: 422 em JSON para o botao (nunca um download quebrado)',
          (rx.status_code, bool((rx.get_json() or {}).get('error'))), (422, True))
    check('sem JavaScript, o mesmo POST devolve a TELA com a mensagem',
          'alert-danger' in cl.post('/tools/option-calculator/extract',
                                    data={'tipo': 'call', 'lado': 'titular', 'strike': ''}).data.decode('utf-8'), True)
    check('ferramenta que nao existe: 404', cl.post('/tools/nada/extract', data={}).status_code, 404)
    sw = cl.post('/tools/swap-calculator/extract', data={}, headers={'X-Requested-With': 'XMLHttpRequest'})
    check('a rota do Export do SWAP nao foi sombreada', sw.status_code in (200, 422), True)

    print('\n%s' % ('TUDO OK' if not FALHAS else '%d FALHA(S): %s' % (len(FALHAS), FALHAS)))
    return 1 if FALHAS else 0


if __name__ == '__main__':
    sys.exit(main())
